import os
import re
import shutil
import subprocess
import tempfile
from io import BytesIO
from zipfile import ZipFile

from docx import Document

from app.services.clp_editor_schema import (
    default_semantic_clp,
    get_tiptap_projection,
    semantic_to_tiptap_doc,
    set_tiptap_projection,
)


def import_docx_to_semantic(file_stream, title="", department="", owner_user_id=None, plan_id=None):
    file_bytes = file_stream.read()
    stream = BytesIO(file_bytes)
    source = Document(stream)
    semantic = default_semantic_clp(title=title, department=department, owner_user_id=owner_user_id, plan_id=plan_id)
    warnings = []
    sections = []
    current = _new_section("imported_section_1", "Imported Content", 1)

    for paragraph in source.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = (getattr(paragraph.style, "name", "") or "").lower()
        if "heading" in style_name:
            if current["blocks"]:
                sections.append(current)
            current = _new_section(f"imported_section_{len(sections) + 1}", text, len(sections) + 1)
            continue
        current["blocks"].append(_paragraph_block(text, {"style": getattr(paragraph.style, "name", "") or ""}))

    if current["blocks"] or not sections:
        sections.append(current)

    for index, table in enumerate(source.tables, start=1):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if not any(any(cell for cell in row) for row in rows):
            continue
        sections.append(
            {
                "section_id": f"imported_table_{index}",
                "section_type": "freeform",
                "title": f"Imported Table {index}",
                "order": len(sections) + 1,
                "department_scope": {"department_id": None, "required": False, "template_placeholder": ""},
                "blocks": [{"block_id": "", "block_type": "table", "content": "", "items": [], "table_ref": "", "metadata": {}, "rows": rows}],
                "provenance": [],
            }
        )

    if not any(section.get("blocks") for section in sections):
        warnings.append("No body paragraphs or tables were found through python-docx; using DOCX XML text fallback.")
        fallback_lines = _extract_docx_xml_lines(file_bytes)
        if fallback_lines:
            sections = [
                {
                    "section_id": "xml_fallback_content",
                    "section_type": "freeform",
                    "title": "Imported DOCX Text",
                    "order": 1,
                    "department_scope": {"department_id": None, "required": False, "template_placeholder": ""},
                    "blocks": [_paragraph_block(line, {"source": "docx_xml_fallback"}) for line in fallback_lines],
                    "provenance": [],
                }
            ]
        else:
            warnings.append("DOCX import found no readable text. The file may rely on unsupported shapes, images, or embedded objects.")

    semantic["sections"] = sections
    semantic.setdefault("validation", {}).setdefault("warnings", []).extend(warnings)
    semantic = set_tiptap_projection(semantic, semantic_to_tiptap_doc(semantic))
    return semantic


def export_semantic_to_docx_bytes(semantic):
    document = Document()
    course = (semantic or {}).get("course", {})
    document.add_heading(course.get("course_title") or "Course Learning Plan", level=1)
    if course.get("course_code"):
        document.add_paragraph(f"Course Code: {course['course_code']}")
    if course.get("course_description"):
        document.add_heading("Course Description", level=2)
        document.add_paragraph(course["course_description"])

    _write_institutional_context(document, semantic)
    _write_outcomes(document, semantic)
    _write_alignment(document, semantic)
    _write_weekly_outline(document, semantic)
    _write_sections(document, semantic)

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


def export_tiptap_projection_to_docx_bytes(semantic):
    document = Document()
    _write_tiptap_nodes(document, get_tiptap_projection(semantic).get("content", []))
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


def convert_docx_bytes_to_pdf(docx_bytes):
    libreoffice = shutil.which("libreoffice") or shutil.which("soffice")
    if not libreoffice:
        raise RuntimeError("LibreOffice is not installed on this server.")

    with tempfile.TemporaryDirectory() as temp_dir:
        docx_path = os.path.join(temp_dir, "clp_export.docx")
        with open(docx_path, "wb") as handle:
            handle.write(docx_bytes.getvalue())
        subprocess.run(
            [libreoffice, "--headless", "--convert-to", "pdf", "--outdir", temp_dir, docx_path],
            check=True,
            timeout=60,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )
        pdf_path = os.path.join(temp_dir, "clp_export.pdf")
        with open(pdf_path, "rb") as handle:
            return BytesIO(handle.read())


def _new_section(section_id, title, order):
    return {
        "section_id": section_id,
        "section_type": "freeform",
        "title": title,
        "order": order,
        "department_scope": {"department_id": None, "required": False, "template_placeholder": ""},
        "blocks": [],
        "provenance": [],
    }


def _paragraph_block(text, metadata=None):
    return {"block_id": "", "block_type": "paragraph", "content": text, "items": [], "table_ref": "", "metadata": metadata or {}}


def _write_institutional_context(document, semantic):
    context = (semantic or {}).get("institutional_context", {})
    for key in ("vision", "mission"):
        section = context.get(key, {})
        blocks = section.get("blocks") or []
        if blocks:
            document.add_heading(section.get("title") or key.title(), level=2)
            for block in blocks:
                _write_block(document, block)


def _write_outcomes(document, semantic):
    outcomes = (semantic or {}).get("outcomes", {})
    clos = outcomes.get("clos") or []
    if clos:
        document.add_heading("Course Learning Outcomes", level=2)
        for item in clos:
            document.add_paragraph(f"{item.get('code') or item.get('clo_id')}: {item.get('statement', '')}")
    plos = outcomes.get("plos") or []
    if plos:
        document.add_heading("Program Learning Outcomes", level=2)
        for item in plos:
            document.add_paragraph(f"{item.get('code') or item.get('plo_id')}: {item.get('description') or item.get('statement', '')}")


def _write_alignment(document, semantic):
    rows = (semantic or {}).get("alignment", {}).get("clo_to_plo") or []
    if not rows:
        return
    document.add_heading("CLO to PLO Alignment", level=2)
    table = document.add_table(rows=1, cols=4)
    table.rows[0].cells[0].text = "CLO"
    table.rows[0].cells[1].text = "PLO"
    table.rows[0].cells[2].text = "Value"
    table.rows[0].cells[3].text = "Rationale"
    for row in rows:
        cells = table.add_row().cells
        cells[0].text = str(row.get("source_clo_id", ""))
        cells[1].text = str(row.get("target_plo_id", ""))
        cells[2].text = str(row.get("mapping_value", ""))
        cells[3].text = str(row.get("rationale", ""))


def _write_weekly_outline(document, semantic):
    weeks = (semantic or {}).get("weekly_outline", {}).get("weeks") or []
    if not weeks:
        return
    document.add_heading("Weekly Outline", level=2)
    table = document.add_table(rows=1, cols=5)
    headers = ["Week", "Topics", "Outcomes", "Activities", "Assessment"]
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for week in weeks:
        cells = table.add_row().cells
        cells[0].text = str(week.get("label") or week.get("week_id") or week.get("week_numbers") or "")
        cells[1].text = _join_value(week.get("topics"))
        cells[2].text = _join_value(week.get("intended_learning_outcomes") or week.get("learning_outcomes"))
        cells[3].text = _join_value(week.get("teaching_learning_activities") or week.get("activities"))
        cells[4].text = _join_value(week.get("assessments") or week.get("assessment"))


def _write_sections(document, semantic):
    for section in (semantic or {}).get("sections") or []:
        document.add_heading(section.get("title") or "Section", level=2)
        for block in section.get("blocks") or []:
            _write_block(document, block)


def _write_block(document, block):
    block_type = block.get("block_type") or block.get("type")
    if block_type == "heading":
        document.add_heading(str(block.get("content", "")), level=3)
    elif block_type == "list":
        for item in block.get("items") or []:
            document.add_paragraph(str(item), style="List Bullet")
    elif block_type == "table":
        rows = block.get("rows") or []
        if rows:
            table = document.add_table(rows=len(rows), cols=max(len(row) for row in rows))
            for row_index, row in enumerate(rows):
                for cell_index, value in enumerate(row):
                    table.rows[row_index].cells[cell_index].text = str(value)
    else:
        text = str(block.get("content", "") or "")
        if text:
            document.add_paragraph(text)


def _write_tiptap_nodes(document, nodes):
    for node in nodes or []:
        node_type = node.get("type")
        text = _node_text(node)
        if node_type == "heading":
            document.add_heading(text, level=(node.get("attrs") or {}).get("level", 1))
        elif node_type == "bulletList":
            for item in node.get("content") or []:
                document.add_paragraph(_node_text(item), style="List Bullet")
        elif node_type == "table":
            rows = node.get("content") or []
            col_count = max((len(row.get("content") or []) for row in rows), default=1)
            table = document.add_table(rows=len(rows), cols=col_count)
            for row_index, row in enumerate(rows):
                for cell_index, cell in enumerate(row.get("content") or []):
                    table.rows[row_index].cells[cell_index].text = _node_text(cell)
        elif text:
            document.add_paragraph(text)


def _node_text(node):
    if not isinstance(node, dict):
        return ""
    if node.get("type") == "text":
        return node.get("text", "")
    return "".join(_node_text(child) for child in node.get("content") or [])


def _join_value(value):
    if isinstance(value, list):
        return "\n".join(_join_value(item) for item in value)
    if isinstance(value, dict):
        return "\n".join(f"{key}: {_join_value(val)}" for key, val in value.items() if val)
    if value is None:
        return ""
    return str(value)


def _extract_docx_xml_lines(file_bytes):
    lines = []
    with ZipFile(BytesIO(file_bytes)) as archive:
        for name in archive.namelist():
            if not name.endswith(".xml"):
                continue
            if not (
                name.startswith("word/document")
                or name.startswith("word/header")
                or name.startswith("word/footer")
            ):
                continue
            xml = archive.read(name).decode("utf-8", "ignore")
            text_nodes = re.findall(r"<w:t[^>]*>(.*?)</w:t>", xml)
            current = []
            for raw in text_nodes:
                text = re.sub(r"<[^>]+>", "", raw)
                text = (
                    text.replace("&amp;", "&")
                    .replace("&lt;", "<")
                    .replace("&gt;", ">")
                    .replace("&quot;", '"')
                    .replace("&apos;", "'")
                ).strip()
                if text:
                    current.append(text)
                if len(" ".join(current)) > 120:
                    lines.append(" ".join(current))
                    current = []
            if current:
                lines.append(" ".join(current))
    seen = set()
    unique = []
    for line in lines:
        normalized = re.sub(r"\s+", " ", line).strip()
        if normalized and normalized not in seen:
            seen.add(normalized)
            unique.append(normalized)
    return unique[:250]
