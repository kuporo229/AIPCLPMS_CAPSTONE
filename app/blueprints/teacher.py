# app/blueprints/teacher.py

import os
import json
import logging
import httpx
import re
from flask import (Blueprint, render_template, request, redirect, url_for, flash,
                   session, abort, jsonify, Response, current_app, make_response)
from app.compat_supabase import (
    build_public_storage_url,
    read_storage_bytes,
    write_storage_bytes,
    verify_password_hash,
)
from app import supabase, STORAGE_BUCKET_NAME, csrf, limiter, cache
from app.forms import (CLPUploadForm, CLPUpdateForm, ChangePasswordForm,
                       CLPGenerateForm, UserProfileForm,
                       AICopilotBetaForm, TeacherSubjectForm,
                       TemplateProfileUploadForm)
from flask_wtf import FlaskForm
from app.decorators import login_required, roles_required
from app.utils import (allowed_file, get_current_user_profile,
                        parse_supabase_timestamp, start_clp_data_generation,
                        start_clp_refinement, start_clp_finalization,
                        create_notification, generate_jwt_token,
                         attach_onlyoffice_ai_plugin,
                         get_department_signatory_settings,
                        get_department_outcomes_bundle, invalidate_user_profile_cache,
                        delete_clp_with_dependencies,
                        log_system_event, user_can_access_clp, user_can_delete_clp,
                        make_workflow_feedback, cleanup_storage_after_commit,
                         ensure_plan_document_file,
                         get_system_settings_map,
                         get_onlyoffice_base_url, stream_storage_file, storage_file_exists,
                         log_document_timing, get_department_id_by_name)
from app.services.docx_service import has_docx_placeholders, replace_placeholders
from app.services.ai_client import AIClient
from app.services.task_queue import TaskQueue
from app.services.copilot_beta_service import (
    beta_ai_call_context,
    beta_display_texts,
    build_beta_docx_replacements,
    build_initial_beta_content,
    compute_validation,
    enforce_beta_weekly_quality,
    ensure_beta_shape,
    generate_beta_alignment,
    generate_beta_clos,
    generate_beta_weekly,
    get_cached_beta_final_review,
    get_context_source_log,
    get_beta_final_review_signature,
    get_review_stage_labels,
    get_static_alignment_context,
    normalize_beta_content,
    run_beta_final_fix,
    run_beta_final_review,
    strip_runtime_copilot_cache,
    update_beta_content_from_form,
    validate_beta_content,
)
from app.services.template_generation_spec import build_generation_spec, merge_spec_into_profile
import time
import requests
import hashlib
import traceback
import jwt
from urllib.parse import urlencode
from pathlib import Path
from werkzeug.utils import secure_filename
from werkzeug.exceptions import HTTPException
from datetime import datetime, timezone
from flask import Blueprint, render_template, redirect, url_for, session, flash, request, current_app
from app.decorators import  login_required
from app import supabase
from app.forms import ChangePasswordForm
import json
from io import BytesIO
from docx import Document

teacher_bp = Blueprint('teacher', __name__)

from app.services.version_service import VersionService

ONLYOFFICE_LPMS_AI_REWRITE_PLUGIN_GUID = "asc.{7BDB4B91-DB22-4B95-8A1F-3D51479C0E8A}"


def _dump_beta_content(content):
    return json.dumps(strip_runtime_copilot_cache(content))

def _log_teacher_onlyoffice(level, event, **details):
    safe_details = {}
    for key, value in details.items():
        if isinstance(value, (str, int, float, bool)) or value is None:
            safe_details[key] = value
        elif isinstance(value, (list, tuple)):
            safe_details[key] = [str(item) for item in value]
        elif isinstance(value, dict):
            safe_details[key] = {str(k): str(v) for k, v in value.items()}
        else:
            safe_details[key] = str(value)
    current_app.logger.log(level, "ONLYOFFICE_TEACHER %s %s", event, json.dumps(safe_details, default=str, sort_keys=True))
    if level >= logging.WARNING:
        log_system_event(
            'onlyoffice',
            'error' if level >= logging.ERROR else 'warning',
            f"Teacher OnlyOffice {event}",
            details=safe_details,
            user_id=session.get('user_id'),
            plan_id=safe_details.get('plan_id'),
        )


def _coerce_json_object(value):
    """Normalizes JSONB payloads from Supabase that may arrive as dict or JSON string."""
    if isinstance(value, dict):
        return value
    if isinstance(value, str):
        try:
            parsed = json.loads(value)
            return parsed if isinstance(parsed, dict) else None
        except Exception:
            return None
    return None


def _get_beta_plan_or_404(plan_id):
    plan_res = supabase.table('course_learning_plans').select('*').eq('id', plan_id).single().execute()
    plan = plan_res.data
    if not plan:
        abort(404)
    if plan.get('user_id') != session.get('user_id'):
        abort(403)
    if plan.get('upload_type') != 'ai_copilot_beta':
        abort(404)
    return plan


def _handle_supabase_route_failure(exc, fallback_endpoint, **fallback_values):
    current_app.logger.warning(f"Supabase request failed: {exc}")
    flash('The database service is temporarily unavailable. Please try again in a moment.', 'warning')
    return redirect(url_for(fallback_endpoint, **fallback_values))


def _wants_json_response():
    return request.headers.get('X-Requested-With') == 'XMLHttpRequest' or request.accept_mimetypes.best == 'application/json'


def _json_error_response(message, status=500, exc=None):
    if exc is not None:
        current_app.logger.error("Beta async route failed: %s\n%s", exc, traceback.format_exc())
    return jsonify({'ok': False, 'error': message}), status


_ACTIVE_BETA_TASK_STATUSES = ('queued', 'processing')


def _normalize_beta_task_index(value):
    if isinstance(value, list):
        value = value[0] if value else None
    if value is None:
        return ''
    return str(value).strip()


def _beta_task_payload_matches(payload, action, clo_index=None, week_index=None):
    payload = _coerce_json_object(payload) or {}
    if payload.get('action') != action:
        return False
    if action == 'regenerate_clo_row':
        return _normalize_beta_task_index(payload.get('clo_index')) == _normalize_beta_task_index(clo_index)
    if action == 'regenerate_week_row':
        return _normalize_beta_task_index(payload.get('week_index')) == _normalize_beta_task_index(week_index)
    return True


def _find_active_beta_action_task(plan_id, user_id, action, clo_index=None, week_index=None):
    try:
        client = current_app.config.get('SUPABASE_SERVICE') or supabase
        res = (
            client.table('background_tasks')
            .select('id,payload,status,progress_percent,progress_label,created_at')
            .eq('task_name', 'beta_action')
            .eq('plan_id', plan_id)
            .eq('user_id', user_id)
            .in_('status', _ACTIVE_BETA_TASK_STATUSES)
            .order('id', desc=True)
            .limit(20)
            .execute()
        )
    except Exception as exc:
        current_app.logger.warning('Could not inspect active beta tasks for plan %s: %s', plan_id, exc)
        return None

    from datetime import datetime, timezone
    now = datetime.now(timezone.utc)
    for row in res.data or []:
        if _beta_task_payload_matches(row.get('payload'), action, clo_index=clo_index, week_index=week_index):
            status = row.get('status')
            created = row.get('created_at')
            # Allow new task if existing is queued (never started) or stale older than 5 min
            if status == 'queued':
                client.table('background_tasks').update({'status': 'failed', 'error_message': 'Replaced: new task started'}).eq('id', row['id']).execute()
                continue
            if created:
                try:
                    if isinstance(created, str):
                        created = datetime.fromisoformat(created.replace('Z', '+00:00'))
                    age_seconds = (now - created).total_seconds()
                    if age_seconds > 300:
                        client.table('background_tasks').update({'status': 'failed', 'error_message': 'Expired: task timed out'}).eq('id', row['id']).execute()
                        continue
                except Exception:
                    pass
            return row
    return None


def _beta_action_form_payload(form):
    return {key: form.getlist(key) for key in form.keys()}


def _active_beta_task_response(task, plan_id):
    payload = _coerce_json_object(task.get('payload')) or {}
    return {
        'ok': False,
        'duplicate': True,
        'error': 'That Copilot action is already running. Please wait for it to finish.',
        'task_id': task.get('id'),
        'plan_id': plan_id,
        'status': task.get('status'),
        'progress_percent': task.get('progress_percent') or 0,
        'progress_label': task.get('progress_label') or '',
        'ai_live_preview': payload.get('ai_live_preview') or '',
        'ai_live_task': payload.get('ai_live_task') or '',
    }


def _start_beta_action_task(plan_id, action, form, clo_index=None, week_index=None):
    existing_task = _find_active_beta_action_task(
        plan_id,
        session.get('user_id'),
        action,
        clo_index=clo_index,
        week_index=week_index,
    )
    if existing_task:
        return _active_beta_task_response(existing_task, plan_id), 409

    payload = {
        'action': action,
        'form_data': _beta_action_form_payload(form),
    }
    if clo_index is not None:
        payload['clo_index'] = str(clo_index)
    if week_index is not None:
        payload['week_index'] = str(week_index)

    task_id = TaskQueue.enqueue('beta_action', payload, user_id=session.get('user_id'), plan_id=plan_id)
    if not task_id:
        return {'ok': False, 'error': 'Could not start the beta task right now.'}, 500
    return {'ok': True, 'task_id': task_id, 'plan_id': plan_id, 'action': action}, 200


def _redirect_after_beta_task_start(plan_id, action, started_message, clo_index=None, week_index=None):
    data, status = _start_beta_action_task(
        plan_id,
        action,
        request.form,
        clo_index=clo_index,
        week_index=week_index,
    )
    if data.get('duplicate'):
        flash(data['error'], 'warning')
    elif not data.get('ok'):
        flash(data.get('error') or 'Could not start the beta task right now.', 'danger')
    return redirect(url_for('teacher.review_clp_ai_beta', plan_id=plan_id))


def _flash_workflow_feedback(feedback):
    flash(feedback['message'], feedback['level'])


def _get_beta_output_template_path():
    configured = current_app.config.get('BETA_OUTPUT_TEMPLATE_PATH')
    if configured:
        return Path(configured)
    return Path(current_app.root_path).parent / 'new_template.docx'


def _required_beta_template_placeholders():
    return {
        "{{course_title}}",
        "{{target_sdgs_display}}",
        "{{clo_1_statement}}",
        "{{clo_1_aligned_plos}}",
        "{{week_1_topics}}",
        "{{week_1_learning_resources}}",
    }


def _resolve_beta_output_template(plan):
    required_placeholders = _required_beta_template_placeholders()
    configured_path = _get_beta_output_template_path()
    candidates = []

    if configured_path.exists():
        candidates.append(("configured", configured_path.name, configured_path.read_bytes()))

    department_name = plan.get('department')
    department_id = get_department_id_by_name(department_name, local_supabase=supabase) if department_name else None
    template_rows = []
    if department_id:
        result = supabase.table('templates').select('filename,name').eq('department_id', department_id).order('created_at', desc=True).execute()
        template_rows.extend(result.data or [])
    default_result = supabase.table('templates').select('filename,name').eq('is_default', True).order('created_at', desc=True).execute()
    for row in default_result.data or []:
        if row not in template_rows:
            template_rows.append(row)

    for row in template_rows:
        filename = row.get('filename')
        if not filename:
            continue
        try:
            candidates.append(("library", row.get('name') or filename, read_storage_bytes(STORAGE_BUCKET_NAME, filename)))
        except Exception:
            continue

    invalid_candidates = []
    for source_type, label, template_bytes in candidates:
        valid, placeholders = has_docx_placeholders(template_bytes, required_placeholders=required_placeholders)
        if valid:
            return template_bytes, label, source_type
        invalid_candidates.append({
            'label': label,
            'placeholder_count': len(placeholders),
            'sample_placeholders': sorted(placeholders)[:8],
        })

    raise ValueError(
        "No valid beta placeholder template is configured. The current template files are static CLP documents, not placeholder templates. "
        "Ask the admin to generate or promote a beta template with placeholders before using Insert Into Template."
    )


def _validate_onlyoffice_callback(data, expected_key=None):
    secret = current_app.config.get('ONLYOFFICE_JWT_SECRET', '')
    if not secret:
        current_app.logger.error("ONLYOFFICE_JWT_SECRET not configured; rejecting callback.")
        return False

    token = None
    auth_header = request.headers.get('Authorization', '')
    if auth_header.lower().startswith('bearer '):
        token = auth_header.split(' ', 1)[1].strip()
    if not token and isinstance(data, dict):
        token = data.get('token')

    if not token:
        current_app.logger.warning("OnlyOffice callback missing token.")
        return False

    try:
        payload = jwt.decode(token, secret, algorithms=['HS256'])
        if expected_key:
            token_key = payload.get('document', {}).get('key') if isinstance(payload, dict) else None
            if token_key and token_key != expected_key:
                current_app.logger.warning("OnlyOffice callback key mismatch.")
                return False
        return True
    except Exception as e:
        current_app.logger.warning(f"OnlyOffice callback token invalid: {e}")
        return False


def _build_onlyoffice_ai_rewrite_token(plan_id, doc_key, user_id):
    secret = current_app.config.get('ONLYOFFICE_JWT_SECRET', '')
    if not secret:
        return ''
    token = jwt.encode(
        {
            'purpose': 'onlyoffice_ai_rewrite',
            'plan_id': int(plan_id),
            'doc_key': doc_key,
            'user_id': str(user_id),
        },
        secret,
        algorithm='HS256',
    )
    return token.decode('utf-8') if isinstance(token, bytes) else token


def _validate_onlyoffice_ai_rewrite_token(plan_id, doc_key=None, token_value=None):
    secret = current_app.config.get('ONLYOFFICE_JWT_SECRET', '')
    if not secret:
        return None
    auth_header = request.headers.get('Authorization', '')
    token = ''
    if token_value:
        token = str(token_value).strip()
    elif auth_header.lower().startswith('bearer '):
        token = auth_header.split(' ', 1)[1].strip()
    token = token or request.args.get('token', '').strip()
    if not token:
        return None
    try:
        payload = jwt.decode(token, secret, algorithms=['HS256'])
    except Exception as exc:
        _log_teacher_onlyoffice(logging.WARNING, 'ai_rewrite_token_invalid', plan_id=plan_id, error=exc)
        return None
    if payload.get('purpose') != 'onlyoffice_ai_rewrite':
        return None
    if str(payload.get('plan_id')) != str(plan_id):
        return None
    if doc_key is not None and str(payload.get('doc_key') or '') != str(doc_key):
        return None
    return str(payload.get('user_id') or '')


def _authorized_onlyoffice_ai_user_id(plan_id):
    if session.get('role') == 'teacher' and session.get('user_id'):
        return str(session.get('user_id'))
    return _validate_onlyoffice_ai_rewrite_token(plan_id)


def _attach_lpms_onlyoffice_rewrite_plugin(config, plan_id, doc_key, user_id):
    rewrite_token = _build_onlyoffice_ai_rewrite_token(plan_id, doc_key, user_id)
    if not rewrite_token:
        return ''

    public_base_url = get_onlyoffice_base_url(internal=False).rstrip('/')
    plugin_config_url = (
        f"{public_base_url}"
        f"{url_for('teacher.onlyoffice_ai_rewrite_plugin_config', plan_id=plan_id, doc_key=doc_key, rewrite_token=rewrite_token)}"
    )
    editor_config = config.setdefault("editorConfig", {})
    plugins = editor_config.setdefault("plugins", {})

    plugins_data = plugins.setdefault("pluginsData", [])
    if plugin_config_url not in plugins_data:
        plugins_data.append(plugin_config_url)

    autostart = plugins.setdefault("autostart", [])
    if ONLYOFFICE_LPMS_AI_REWRITE_PLUGIN_GUID not in autostart:
        autostart.append(ONLYOFFICE_LPMS_AI_REWRITE_PLUGIN_GUID)

    options = plugins.setdefault("options", {})
    options.setdefault(ONLYOFFICE_LPMS_AI_REWRITE_PLUGIN_GUID, {})["rewriteUrl"] = url_for(
        'teacher.onlyoffice_ai_rewrite_selection',
        plan_id=plan_id,
    )
    return rewrite_token


def _sanitize_onlyoffice_ai_error(exc):
    error_text = str(exc or "")
    lowered = error_text.lower()
    if (
        "gemini_api_key" in lowered
        or "api key" in lowered
        or "unauthorized" in lowered
        or "unauthenticated" in lowered
        or "permission_denied" in lowered
    ):
        return "AI service is not configured correctly. Please contact an administrator."
    return "AI rewrite failed. Please try again."


@teacher_bp.route('/onlyoffice_client_log/<int:plan_id>', methods=['POST'])
@login_required
@roles_required('teacher')
@csrf.exempt
def onlyoffice_client_log(plan_id):
    payload = request.get_json(silent=True) or {}
    try:
        plan_res = supabase.table('course_learning_plans').select('id, user_id, filename').eq('id', plan_id).single().execute()
        plan = plan_res.data
    except Exception as exc:
        _log_teacher_onlyoffice(logging.ERROR, 'client_log_plan_lookup_failed', plan_id=plan_id, error=exc)
        return jsonify({"ok": False}), 500

    if not plan:
        _log_teacher_onlyoffice(logging.WARNING, 'client_log_plan_missing', plan_id=plan_id)
        return jsonify({"ok": False}), 404
    if plan['user_id'] != session.get('user_id'):
        _log_teacher_onlyoffice(logging.WARNING, 'client_log_forbidden', plan_id=plan_id, session_user_id=session.get('user_id'))
        return jsonify({"ok": False}), 403

    _log_teacher_onlyoffice(
        logging.INFO,
        'client_event',
        plan_id=plan_id,
        session_user_id=session.get('user_id'),
        filename=plan.get('filename'),
        remote_addr=request.headers.get('X-Forwarded-For', request.remote_addr),
        user_agent=request.headers.get('User-Agent'),
        client_event=payload.get('event'),
        message=payload.get('message'),
        extra=payload.get('extra') if isinstance(payload.get('extra'), dict) else None,
    )
    return jsonify({"ok": True})


@teacher_bp.route('/onlyoffice_ai/rewrite_selection/<int:plan_id>', methods=['POST'])
@csrf.exempt
@limiter.limit("20 per hour")
def onlyoffice_ai_rewrite_selection(plan_id):
    authorized_user_id = _authorized_onlyoffice_ai_user_id(plan_id)
    if not authorized_user_id:
        return jsonify({"ok": False, "error": "Unauthorized."}), 403

    payload = request.get_json(silent=True) or {}
    selected_text = (payload.get('text') or '').strip()
    instruction = (payload.get('instruction') or '').strip()

    if not selected_text:
        return jsonify({"ok": False, "error": "Select text in the document first."}), 400
    if len(selected_text) > 6000:
        return jsonify({"ok": False, "error": "Selected text is too long for a focused rewrite."}), 400

    try:
        plan_res = supabase.table('course_learning_plans').select(
            'id, user_id, subject, department, content'
        ).eq('id', plan_id).single().execute()
        plan = plan_res.data
    except Exception as exc:
        _log_teacher_onlyoffice(logging.ERROR, 'ai_rewrite_plan_lookup_failed', plan_id=plan_id, error=exc)
        return jsonify({"ok": False, "error": "Could not load the plan."}), 500

    if not plan:
        return jsonify({"ok": False, "error": "Plan not found."}), 404
    if str(plan.get('user_id')) != authorized_user_id:
        _log_teacher_onlyoffice(logging.WARNING, 'ai_rewrite_forbidden', plan_id=plan_id, authorized_user_id=authorized_user_id)
        return jsonify({"ok": False, "error": "Unauthorized."}), 403

    content = _coerce_json_object(plan.get('content')) or {}
    metadata = content.get('metadata') if isinstance(content.get('metadata'), dict) else {}
    course_title = metadata.get('course_title') or plan.get('subject') or ''
    course_code = metadata.get('course_code') or content.get('course_code') or ''

    try:
        model_instance = AIClient.get_model()
        prompt = f"""
You are an academic curriculum editor for a Course Learning Plan.
Rewrite only the selected document text below. If the text is a Course Learning Outcome (CLO), make it measurable, concise, outcome-based, and aligned with Bloom's taxonomy. Preserve any CLO numbering, labels, and table-friendly line breaks already present.

Course context:
- Department: {plan.get('department') or ''}
- Course code: {course_code}
- Course title: {course_title}

Teacher instruction:
{instruction or 'Rewrite this selection as a stronger CLO or academic CLP statement.'}

Selected text:
{selected_text}

Return only the replacement text. Do not add explanations, markdown, or quotation marks.
"""
        response = AIClient.generate_with_retry(
            model_instance,
            [prompt],
            {"temperature": 0.35},
            task_type="onlyoffice_selection_rewrite",
            plan_id=plan_id,
            user_id=authorized_user_id,
        )
        rewritten = (response.text or '').strip().strip('"')
        if not rewritten:
            return jsonify({"ok": False, "error": "AI returned an empty rewrite."}), 502

        _log_teacher_onlyoffice(
            logging.INFO,
            'ai_rewrite_success',
            plan_id=plan_id,
            session_user_id=authorized_user_id,
            selected_length=len(selected_text),
            rewritten_length=len(rewritten),
        )
        return jsonify({"ok": True, "replacement": rewritten})
    except Exception as exc:
        _log_teacher_onlyoffice(logging.ERROR, 'ai_rewrite_failed', plan_id=plan_id, error=exc, traceback=traceback.format_exc())
        return jsonify({"ok": False, "error": _sanitize_onlyoffice_ai_error(exc)}), 500


def _onlyoffice_plugin_response(response):
    response.headers['Access-Control-Allow-Origin'] = '*'
    response.headers['Access-Control-Allow-Headers'] = 'Authorization, Content-Type'
    response.headers['Access-Control-Allow-Methods'] = 'GET, OPTIONS'
    return response


@teacher_bp.route('/onlyoffice_ai/rewrite_plugin/<int:plan_id>/<doc_key>/config.json', methods=['GET', 'OPTIONS'])
@teacher_bp.route('/onlyoffice_ai/rewrite_plugin/<int:plan_id>/<doc_key>/<path:rewrite_token>/config.json', methods=['GET', 'OPTIONS'])
@csrf.exempt
@limiter.exempt
def onlyoffice_ai_rewrite_plugin_config(plan_id, doc_key, rewrite_token=None):
    if request.method == 'OPTIONS':
        return _onlyoffice_plugin_response(Response(status=204))
    authorized_user_id = _validate_onlyoffice_ai_rewrite_token(plan_id, doc_key=doc_key, token_value=rewrite_token)
    if not authorized_user_id:
        authorized_user_id = _validate_onlyoffice_ai_rewrite_token(plan_id, doc_key=doc_key)
    if not authorized_user_id:
        return _onlyoffice_plugin_response(jsonify({"error": "Unauthorized"})), 403

    plugin_token = rewrite_token or request.args.get('token', '')
    public_base_url = get_onlyoffice_base_url(internal=False)
    plugin_base_url = (
        f"{public_base_url}/teacher/onlyoffice_ai/rewrite_plugin/"
        f"{plan_id}/{doc_key}/{plugin_token}/"
    )
    config = {
        "name": "LPMS AI Rewrite",
        "guid": ONLYOFFICE_LPMS_AI_REWRITE_PLUGIN_GUID,
        "baseUrl": plugin_base_url,
        "version": "1.0.0",
        "variations": [
            {
                "name": "LPMS AI Rewrite",
                "description": "Rewrite selected CLO text through LPMS AI.",
                "url": "index.html",
                "isViewer": False,
                "EditorsSupport": ["word"],
                "isVisual": False,
                "isSystem": False,
                "isModal": False,
                "isInsideMode": False,
                "initDataType": "none",
                "initData": "",
                "events": ["onToolbarMenuClick"],
                "buttons": [],
            }
        ],
    }
    return _onlyoffice_plugin_response(jsonify(config))


@teacher_bp.route('/onlyoffice_ai/rewrite_plugin/<int:plan_id>/<doc_key>/index.html', methods=['GET', 'OPTIONS'])
@teacher_bp.route('/onlyoffice_ai/rewrite_plugin/<int:plan_id>/<doc_key>/<path:rewrite_token>/index.html', methods=['GET', 'OPTIONS'])
@csrf.exempt
@limiter.exempt
def onlyoffice_ai_rewrite_plugin_index(plan_id, doc_key, rewrite_token=None):
    if request.method == 'OPTIONS':
        return _onlyoffice_plugin_response(Response(status=204))
    token = (rewrite_token or request.args.get('token', '')).strip()
    authorized_user_id = _validate_onlyoffice_ai_rewrite_token(plan_id, doc_key=doc_key, token_value=token)
    if not authorized_user_id or not token:
        return _onlyoffice_plugin_response(Response("Unauthorized", status=403, mimetype='text/plain'))

    rewrite_url = url_for('teacher.onlyoffice_ai_rewrite_selection', plan_id=plan_id, _external=True)
    html = f"""<!DOCTYPE html>
<html>
<head>
  <meta charset="utf-8">
  <script src="https://onlyoffice.github.io/sdkjs-plugins/v1/plugins.js"></script>
  <script>
    const LPMS_REWRITE_URL = {json.dumps(rewrite_url)};
    const LPMS_REWRITE_TOKEN = {json.dumps(token)};

    function selectedTextOptions() {{
      return {{
        Numbering: false,
        Math: false,
        TableCellSeparator: "\\n",
        TableRowSeparator: "\\n",
        ParaSeparator: "\\n",
        TabSymbol: "\\t",
        NewLineSeparator: "\\n"
      }};
    }}

    function showMessage(message) {{
      try {{
        window.Asc.plugin.executeMethod("ShowMessage", [message]);
      }} catch (err) {{
        alert(message);
      }}
    }}

    function rewriteSelection() {{
      window.Asc.plugin.executeMethod("GetSelectedText", [selectedTextOptions()], function(selectedText) {{
        const text = String(selectedText || "").trim();
        if (!text) {{
          showMessage("Select CLO text in the document first.");
          return;
        }}
        const instruction = prompt("AI rewrite instruction", "Rewrite this as a stronger CLO.");
        if (instruction === null) return;

        fetch(LPMS_REWRITE_URL, {{
          method: "POST",
          headers: {{
            "Authorization": "Bearer " + LPMS_REWRITE_TOKEN,
            "Content-Type": "application/json"
          }},
          body: JSON.stringify({{ text, instruction }})
        }})
          .then(function(response) {{
            return response.json().then(function(data) {{
              if (!response.ok || !data.ok) throw new Error(data.error || "AI rewrite failed.");
              return data.replacement || "";
            }});
          }})
          .then(function(replacement) {{
            const lines = String(replacement).split(/\\r?\\n/);
            window.Asc.plugin.executeMethod("ReplaceTextSmart", [lines], function(done) {{
              if (done === false) {{
                window.Asc.plugin.executeMethod("RemoveSelectedContent", [], function() {{
                  window.Asc.plugin.executeMethod("PasteText", [replacement], function() {{
                    showMessage("Selected text updated.");
                  }});
                }});
                return;
              }}
              showMessage("Selected text updated.");
            }});
          }})
          .catch(function(err) {{
            showMessage(err.message || "AI rewrite failed.");
          }});
      }});
    }}

    window.Asc.plugin.init = function() {{
      const pluginGuid = (window.Asc.plugin.info && window.Asc.plugin.info.guid) || window.Asc.plugin.guid;
      window.Asc.plugin.executeMethod("AddToolbarMenuItem", [{{
        guid: pluginGuid,
        tabs: [{{
          id: "lpms_ai",
          text: "LPMS AI",
          items: [{{
            id: "lpms_ai_rewrite_clo",
            type: "button",
            text: "Rewrite CLO",
            hint: "Rewrite selected CLO text with LPMS AI.",
            lockInViewMode: true,
            items: []
          }}]
        }}]
      }}]);
      if (window.Asc.plugin.attachToolbarMenuClickEvent) {{
        window.Asc.plugin.attachToolbarMenuClickEvent("lpms_ai_rewrite_clo", rewriteSelection);
      }}
    }};

    window.Asc.plugin.event_onToolbarMenuClick = function(id) {{
      if (id === "lpms_ai_rewrite_clo") rewriteSelection();
    }};

    window.Asc.plugin.button = function() {{}};
  </script>
</head>
<body></body>
</html>"""
    return _onlyoffice_plugin_response(Response(html, mimetype='text/html'))

@teacher_bp.route('/clp/<int:plan_id>/versions')
@login_required
def list_clp_versions(plan_id):
    """Lists all versions of a specific CLP."""
    res = supabase.table('course_learning_plans').select('user_id, subject').eq('id', plan_id).single().execute()
    if not res.data: abort(404)
    if res.data['user_id'] != session.get('user_id') and session.get('role') != 'admin':
        abort(403)

    versions = VersionService.get_versions(plan_id)
    return render_template('clp_versions.html', plan_id=plan_id, subject=res.data['subject'], versions=versions)

@teacher_bp.route('/clp/<int:plan_id>/versions/<int:version_number>')
@login_required
def view_clp_version(plan_id, version_number):
    """Views a specific version snapshot of a CLP."""
    res = supabase.table('course_learning_plans').select('id, user_id, subject, department, status').eq('id', plan_id).single().execute()
    if not res.data: abort(404)
    if res.data['user_id'] != session.get('user_id') and session.get('role') != 'admin':
        abort(403)

    version_res = supabase.table('clp_versions').select('*').eq('plan_id', plan_id).eq('version_number', version_number).single().execute()
    if not version_res.data: abort(404)

    content_data = version_res.data['content']

    outcomes_bundle = get_department_outcomes_bundle(res.data.get('department'))

    # Reuse view_ai_clp template but with a "Version View" flag
    return render_template('view_ai_clp.html', plan=res.data, content_data=content_data,
                           program_outcomes=outcomes_bundle['program_outcomes'],
                           course_outcomes=outcomes_bundle['course_outcomes'],
                           institutional_headers=outcomes_bundle['institutional_headers'],
                           program_headers=outcomes_bundle['program_headers'],
                           is_version_view=True,
                           version_number=version_number)

@teacher_bp.route('/clp/<int:plan_id>/versions/<int:version_number>/restore', methods=['POST'])
@login_required
def restore_clp_version(plan_id, version_number):
    """Restores a CLP to a previous version."""
    res = supabase.table('course_learning_plans').select('user_id').eq('id', plan_id).single().execute()
    if not res.data: abort(404)
    if res.data['user_id'] != session.get('user_id'): abort(403)

    success, message = VersionService.restore_version(plan_id, version_number, session['user_id'])
    if success:
        flash(f"Successfully restored to version {version_number}.", "success")
    else:
        flash(f"Restore failed: {message}", "danger")

    return redirect(url_for('teacher.view_clp', plan_id=plan_id))

@teacher_bp.route('/api/department_outcomes/<string:dept_name>')
@login_required
def get_dept_outcomes(dept_name):
    """API endpoint to fetch all outcomes for a department (cached)."""
    try:
        result = get_department_outcomes_bundle(dept_name)
        if result['department_id'] is None:
            return jsonify({'error': 'Dept not found'}), 404
        return jsonify(result)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@teacher_bp.route('/api/ai_suggest_rephrase', methods=['POST'])
@login_required
def ai_suggest_rephrase():
    """AI endpoint to rephrase a learning outcome or description."""
    data = request.get_json()
    text = data.get('text')
    context = data.get('context', 'Learning Outcome')

    if not text: return jsonify({'error': 'No text provided'}), 400

    try:
        model_instance = AIClient.get_model()
        prompt = f"""
        You are an academic curriculum expert. Rephrase the following {context} to be more professional,
        concise, and measurable (using Bloom's Taxonomy verbs where applicable).

        ORIGINAL TEXT:
        {text}

        Return ONLY the rephrased text. No conversational filler.
        """
        response = AIClient.generate_with_retry(model_instance, [prompt], {"temperature": 0.7})
        return jsonify({'suggestion': response.text.strip()})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# --- STATUS CHECK ROUTE ---
@teacher_bp.route('/clp/<int:plan_id>/status')
@login_required
@limiter.exempt
def clp_status(plan_id):
    """Returns the current status of a specific CLP as JSON for polling."""
    try:
        from app.services.task_queue import TaskQueue
        res = supabase.table('course_learning_plans').select('status, content, user_id').eq('id', plan_id).single().execute()
        if res.data:
            if res.data.get('user_id') != session.get('user_id'):
                return jsonify({'status': 'error', 'details': 'Unauthorized'}), 403

            content = res.data.get('content')
            percent = 0
            label = "Processing..."

            if content:
                import json
                if isinstance(content, str):
                    try:
                        content = json.loads(content)
                    except:
                        content = {}

                progress = content.get('progress', {})
                percent = progress.get('percent', 5)
                label = progress.get('label', 'Processing...')

            # Ensure at least 5% if it's generating
            if percent < 5: percent = 5

            # Also check background_tasks table
            queue_tasks = supabase.table('background_tasks').select('progress_percent, progress_label, status').eq('plan_id', plan_id).in_('status', ['queued', 'processing']).order('id', desc=True).limit(1).execute()
            if queue_tasks.data:
                latest_task = queue_tasks.data[0]
                task_percent = latest_task.get('progress_percent', 0) or 0
                # Only override if the task is actually reporting progress beyond initialization
                if task_percent > percent:
                    percent = task_percent
                    label = latest_task.get('progress_label', label)

            return jsonify({
                'status': res.data['status'],
                'percent': percent,
                'label': label
            })
        return jsonify({'status': 'unknown'}), 404
    except Exception as e:
        current_app.logger.warning(f"Failed to fetch CLP status for {plan_id}: {e}")
        return jsonify({'status': 'error', 'details': 'Unable to fetch the latest progress right now.'}), 500

@teacher_bp.route('/clp/<int:plan_id>/abort', methods=['POST'])
@login_required
def abort_generation(plan_id):
    """Aborts AI generation without deleting the CLP."""
    try:
        # Verify ownership
        res = supabase.table('course_learning_plans').select('user_id, status').eq('id', plan_id).single().execute()
        if not res.data or res.data['user_id'] != session.get('user_id'):
            return jsonify({'success': False, 'error': 'Unauthorized'}), 403

        current_status = res.data['status']

        # Determine revert status: If it was in finalization phase, revert to draft_review
        # Otherwise, revert to draft (which it usually is before generation)
        new_status = 'draft'
        if current_status == 'generating_doc':
            new_status = 'draft_review'
        elif current_status == 'generating':
            # Could be from draft or draft_review (apply fixes)
            # Defaulting to draft_review is usually safer as it keeps it visible in the main lists
            new_status = 'draft_review'

        # Revert status
        supabase.table('course_learning_plans').update({'status': new_status}).eq('id', plan_id).execute()

        # Mark background tasks as failed
        supabase.table('background_tasks').update({'status': 'failed', 'error_message': 'Aborted by user'}).eq('plan_id', plan_id).in_('status', ['queued', 'processing']).execute()

        from app.services.observability import StructuredLogger
        StructuredLogger.info(f"User aborted generation for plan {plan_id}", user_id=session.get('user_id'), plan_id=plan_id)
        return jsonify({'success': True})
    except Exception as e:
        current_app.logger.warning(f"Failed to abort generation for plan {plan_id}: {e}")
        return jsonify({'success': False, 'error': 'Unable to stop generation right now.'}), 500

@teacher_bp.route('/create_clp_ai/draft_description', methods=['POST'])
@login_required
@roles_required('teacher')
def generate_course_description_draft():
    """Generates a draft course description based on Subject and Department Outcomes."""
    data = request.get_json()
    department_name = data.get('department')
    subject = data.get('subject')

    if not department_name or not subject:
        return jsonify({'error': 'Please select a Department and enter a Course Title first.'}), 400

    try:
        # 1. Fetch Department Outcomes to make the description relevant
        po_context = ""
        try:
            outcomes_bundle = get_department_outcomes_bundle(department_name)
            if outcomes_bundle['program_outcomes']:
                # Take top 3-5 outcomes to contextually ground the description without overloading tokens
                outcomes = [p['description'] for p in outcomes_bundle['program_outcomes'][:5]]
                po_context = "\n".join([f"- {o}" for o in outcomes])
        except Exception as db_e:
            current_app.logger.warning(f"Could not fetch outcomes for draft generation: {db_e}")

        # 2. Generate with AI
        model_instance = AIClient.get_model()

        prompt = f"""
        You are an academic curriculum specialist.
        Write a professional, academic Course Description for a subject titled "{subject}" offered by the "{department_name}".

        The description MUST:
        1. Be approximately 80-150 words in a single cohesive paragraph.
        2. Emphasize the COURSE itself — what concepts, theory, and practice it covers.
        3. Mention key topical areas, methodologies, or technologies the course addresses.
        4. Briefly note evaluation or design approaches used (e.g., critical evaluation, hands-on methods, design guidelines).
        5. State what students will have learned or be able to do by the end of the course.
        6. Align implicitly with these Program Outcomes (do not list them, just let them influence the description):
        {po_context}

        The tone should be formal, direct, and content-focused — similar to:
        "This course provides concepts, theory and practice to the field of ... The course covers ... Students at the end of the course will have learned ..."

        Return ONLY the paragraph text. No labels, no markdown, no bullet points.
        """

        response = AIClient.generate_with_retry(model_instance, [prompt], {"temperature": 0.7})
        draft_text = response.text.strip()

        return jsonify({'description': draft_text})

    except Exception as e:
        current_app.logger.error(f"Draft Gen Error: {e}")
        return jsonify({'error': f"AI Generation failed: {str(e)}"}), 500


@teacher_bp.route('/subjects/generate_service_learning', methods=['POST'])
@login_required
@roles_required('teacher')
def generate_service_learning():
    """Generates a draft Service Learning Component for a subject."""
    data = request.get_json()
    department_name = data.get('department')
    subject = data.get('subject')
    description = data.get('description', '')

    if not department_name or not subject:
        return jsonify({'error': 'Please provide a Department and Course Title first.'}), 400

    try:
        model_instance = AIClient.get_model()
        prompt = f"""
You are an academic curriculum specialist writing for a Philippine higher education institution.
Generate a short Service Learning Component project title for the course "{subject}" from the "{department_name}".

Course context:
{description or 'No description provided.'}

Rules:
- Output a single short project title (5-10 words maximum).
- It must name a specific community engagement or service project tied to the course content.
- Format: Title Case, like "Visual Communication Enhancement for Campus Offices" or "Digital Literacy Program for Public School Teachers".
- Return ONLY the title. No punctuation at the end, no labels, no explanation.
"""
        response = AIClient.generate_with_retry(model_instance, [prompt], {"temperature": 0.7})
        return jsonify({'text': response.text.strip()})
    except Exception as e:
        current_app.logger.error(f"Service Learning Gen Error: {e}")
        return jsonify({'error': f'AI generation failed: {str(e)}'}), 500


@teacher_bp.route('/subjects/generate_target_sdg', methods=['POST'])
@login_required
@roles_required('teacher')
def generate_target_sdg():
    """Suggests relevant UN SDGs for a subject based on course content."""
    data = request.get_json()
    subject = data.get('subject')
    department_name = data.get('department')
    description = data.get('description', '')

    if not subject:
        return jsonify({'error': 'Please enter a Course Title first.'}), 400

    try:
        from app.services.copilot_beta_service import SDG_CONTEXT
        sdg_rows = []
        if department_name:
            try:
                dept_result = supabase.table('departments').select('id').eq('name', department_name).limit(1).execute()
                dept_id = (dept_result.data or [{}])[0].get('id')
                if dept_id:
                    sdg_result = supabase.table('copilot_reference_entries').select('code,title,description').eq('department_id', dept_id).eq('category', 'sdg').order('sort_order').execute()
                    sdg_rows = sdg_result.data or []
            except Exception:
                pass
        if not sdg_rows:
            sdg_rows = [{'code': k, 'title': v['title'], 'description': v['guidance']} for k, v in SDG_CONTEXT.items()]

        sdg_list_text = "\n".join(
            f"- {row['code']}: {row.get('title', row['code'])} — {row.get('description', '')}"
            for row in sdg_rows
        )

        model_instance = AIClient.get_model()
        prompt = f"""
You are an academic curriculum specialist.
From the APPROVED list below, identify the most relevant SDGs for the course "{subject}" from the "{department_name}".

Course description:
{description or 'No description provided.'}

APPROVED SDGs ONLY (you must only pick from these):
{sdg_list_text}

Rules:
- Choose 1 to 3 SDGs from the approved list above ONLY. Do not invent others.
- Output ONLY a comma-separated list, e.g.: SDG 4, SDG 9
- No explanations, labels, or extra text.
"""
        response = AIClient.generate_with_retry(model_instance, [prompt], {"temperature": 0.3})
        return jsonify({'text': response.text.strip().rstrip('.')})
    except Exception as e:
        current_app.logger.error(f"Target SDG Gen Error: {e}")
        return jsonify({'error': f'AI generation failed: {str(e)}'}), 500


def _user_has_consultation_hours(user_profile):
    """Return True if the user has at least one consultation slot with time and room."""
    cons = user_profile.get('consultation_hours')
    if not isinstance(cons, dict):
        return False
    day_order = ['monday', 'tuesday', 'wednesday', 'thursday', 'friday']
    for day in day_order:
        day_data = cons.get(day)
        if isinstance(day_data, dict) and day_data.get('time') and day_data.get('room'):
            return True
    return False


def _parse_active_term_from_current_semester(current_semester):
    text = str(current_semester or '').strip()
    if not text:
        return '', ''
    year_match = re.search(r'\b\d{4}-\d{4}\b', text)
    academic_year = year_match.group(0) if year_match else ''
    semester = text.replace(academic_year, '').strip(' -') if academic_year else text
    return semester.strip(), academic_year.strip()


def _get_active_term_settings():
    settings = get_system_settings_map(current_app.config.get('SUPABASE_SERVICE'))
    active_semester = str(settings.get('active_semester') or '').strip()
    active_academic_year = str(settings.get('active_academic_year') or '').strip()

    if not active_semester or not active_academic_year:
        fallback_semester, fallback_year = _parse_active_term_from_current_semester(settings.get('current_semester'))
        active_semester = active_semester or fallback_semester or '1st Semester'
        active_academic_year = active_academic_year or fallback_year or str(datetime.now().year)

    return active_semester, active_academic_year


def _get_visible_teacher_subjects(user_profile, active_semester, active_academic_year):
    department_name = (user_profile or {}).get('assigned_department') or ''
    user_id = session.get('user_id')
    rows = (
        supabase.table('teacher_subjects')
        .select('*')
        .eq('semester', active_semester)
        .eq('academic_year', active_academic_year)
        .order('course_code')
        .limit(50)
        .execute()
        .data
        or []
    )

    visible = []
    for row in rows:
        row_user_id = row.get('user_id')
        row_department = row.get('department') or ''
        if str(row_user_id or '') == str(user_id):
            visible.append(row)
            continue
        if row_user_id is None and department_name and row_department == department_name:
            visible.append(row)
    return visible


def _beta_subject_metadata_errors(subject):
    def word_count(value):
        return len(re.findall(r"\b[\w'-]+\b", str(value or "")))

    errors = []
    if word_count(subject.get('course_description')) < 8:
        errors.append("Course description must explain the course in at least 8 words.")

    units = str(subject.get('units') or '').strip()
    if not re.fullmatch(r"\d+(?:\.\d+)?(?:\s*(?:unit|units))?", units, flags=re.IGNORECASE):
        errors.append("Units must be a proper numeric value, for example '3' or '3 Units'.")

    contact_hours = str(subject.get('contact_hours') or '').strip().lower()
    has_number = bool(re.search(r"\d", contact_hours))
    has_time_unit = any(marker in contact_hours for marker in ("hour", "hr", "lecture", "lab", "laboratory"))
    if not has_number or not has_time_unit:
        errors.append("Contact hours must use a proper format, for example '5 hours per week' or '3 lecture, 2 lab'.")

    return errors


@teacher_bp.route('/subjects', methods=['GET'])
@login_required
@roles_required('teacher')
def manage_subjects():
    user_profile = get_current_user_profile() or {}
    active_semester, active_academic_year = _get_active_term_settings()
    subjects = _get_visible_teacher_subjects(user_profile, active_semester, active_academic_year)

    # Annotate each subject with its linked template profile info
    profile_ids = set()
    for s in subjects:
        pid = s.get('template_profile_id')
        if pid:
            profile_ids.add(pid)

    profile_map = {}
    if profile_ids:
        try:
            tp_result = supabase.table('teacher_template_profiles').select(
                'id,name,status,source_filename,department'
            ).in_('id', list(profile_ids)).execute()
            for tp in (tp_result.data or []):
                profile_map[tp['id']] = tp
        except Exception:
            pass

    for s in subjects:
        pid = s.get('template_profile_id')
        if pid and pid in profile_map:
            s['_profile'] = profile_map[pid]
        else:
            s['_profile'] = None

    # Annotate with CLP usage count
    subject_ids = [s['id'] for s in subjects]
    clp_counts = {}
    if subject_ids:
        try:
            clp_res = supabase.table('course_learning_plans').select('subject_id').eq('user_id', session.get('user_id')).in_('subject_id', subject_ids).execute()
            for row in (clp_res.data or []):
                sid = row.get('subject_id')
                clp_counts[sid] = clp_counts.get(sid, 0) + 1
        except Exception:
            pass
    for s in subjects:
        s['_clp_count'] = clp_counts.get(s['id'], 0)

    # Stats
    total_subjects = len(subjects)
    linked_count = sum(1 for s in subjects if s.get('_profile'))
    clp_subjects = sum(1 for s in subjects if s.get('_clp_count', 0) > 0)

    return render_template(
        'teacher_subjects.html',
        subjects=subjects,
        total_subjects=total_subjects,
        linked_count=linked_count,
        clp_subjects=clp_subjects,
        active_semester=active_semester,
        active_academic_year=active_academic_year,
        user_id=session.get('user_id'),
    )


@teacher_bp.route('/subjects/add', methods=['GET', 'POST'])
@login_required
@roles_required('teacher')
def add_subject():
    user_profile = get_current_user_profile() or {}
    active_semester, active_academic_year = _get_active_term_settings()

    # Load confirmed template profiles for the user
    confirmed_profiles = []
    try:
        tp_result = supabase.table('teacher_template_profiles').select('id,source_filename,department').eq('user_id', session.get('user_id')).eq('status', 'confirmed').order('created_at', desc=True).execute()
        confirmed_profiles = tp_result.data or []
    except Exception:
        pass

    form = TeacherSubjectForm(template_profiles=confirmed_profiles)

    if request.method == 'GET' and user_profile.get('assigned_department'):
        form.department.data = user_profile.get('assigned_department')

    if request.method == 'POST' and not form.validate_on_submit():
        for field_name, errors in form.errors.items():
            for error in errors:
                flash(f"{getattr(form, field_name).label.text}: {error}", 'danger')

    if form.validate_on_submit():
        department_value = form.department.data
        if user_profile.get('assigned_department'):
            department_value = user_profile.get('assigned_department')

        insert_data = {
            'user_id': session.get('user_id'),
            'department': department_value,
            'semester': active_semester,
            'academic_year': active_academic_year,
            'course_code': form.course_code.data,
            'course_title': form.course_title.data,
            'course_description': form.course_description.data,
            'type_of_course': form.type_of_course.data,
            'units': form.units.data,
            'contact_hours': form.contact_hours.data,
            'pre_requisites': form.pre_requisites.data,
            'co_requisites': form.co_requisites.data,
            'class_schedule': form.class_schedule.data,
            'room_assignment': form.room_assignment.data,
            'service_learning_component': form.service_learning_component.data,
            'target_sdg': form.target_sdg.data,
            'prepared_by_name': form.prepared_by_name.data,
            'prepared_by_position': form.prepared_by_position.data,
            'reviewed_by_name': form.reviewed_by_name.data,
            'reviewed_by_position': form.reviewed_by_position.data,
            'endorsed_by_name': form.endorsed_by_name.data,
            'endorsed_by_position': form.endorsed_by_position.data,
            'approved_by_name': form.approved_by_name.data,
            'approved_by_position': form.approved_by_position.data,
        }
        # Save template_profile_id if selected
        if form.template_profile_id.data:
            insert_data['template_profile_id'] = int(form.template_profile_id.data)

        try:
            supabase.table('teacher_subjects').insert(insert_data).execute()
            flash('Subject added.' + (' Linked to template profile.' if form.template_profile_id.data else ''), 'success')
            return redirect(url_for('teacher.manage_subjects'))
        except Exception as exc:
            from app.utils import friendly_error
            flash(friendly_error(exc, 'Could not save subject.'), 'danger')

    return render_template(
        'teacher_subject_form.html',
        form=form,
        mode='add',
        active_semester=active_semester,
        active_academic_year=active_academic_year,
        confirmed_profiles=confirmed_profiles,
    )


@teacher_bp.route('/subjects/<int:subject_id>/edit', methods=['GET', 'POST'])
@login_required
@roles_required('teacher')
def edit_subject(subject_id):
    user_id = session.get('user_id')
    subject = (
        supabase.table('teacher_subjects')
        .select('*')
        .eq('id', subject_id)
        .eq('user_id', user_id)
        .single()
        .execute()
        .data
    )
    if not subject:
        abort(404)

    # Load confirmed template profiles
    confirmed_profiles = []
    try:
        tp_result = supabase.table('teacher_template_profiles').select('id,source_filename,department').eq('user_id', user_id).eq('status', 'confirmed').order('created_at', desc=True).execute()
        confirmed_profiles = tp_result.data or []
    except Exception:
        pass

    active_semester, active_academic_year = _get_active_term_settings()
    form = TeacherSubjectForm(data={
        'department': subject.get('department'),
        'course_code': subject.get('course_code'),
        'course_title': subject.get('course_title'),
        'course_description': subject.get('course_description'),
        'type_of_course': subject.get('type_of_course'),
        'units': subject.get('units'),
        'contact_hours': subject.get('contact_hours'),
        'pre_requisites': subject.get('pre_requisites'),
        'co_requisites': subject.get('co_requisites'),
        'class_schedule': subject.get('class_schedule'),
        'room_assignment': subject.get('room_assignment'),
        'service_learning_component': subject.get('service_learning_component'),
        'target_sdg': subject.get('target_sdg'),
        'template_profile_id': str(subject.get('template_profile_id') or ''),
        'prepared_by_name': subject.get('prepared_by_name') or '',
        'prepared_by_position': subject.get('prepared_by_position') or '',
        'reviewed_by_name': subject.get('reviewed_by_name') or '',
        'reviewed_by_position': subject.get('reviewed_by_position') or '',
        'endorsed_by_name': subject.get('endorsed_by_name') or '',
        'endorsed_by_position': subject.get('endorsed_by_position') or '',
        'approved_by_name': subject.get('approved_by_name') or '',
        'approved_by_position': subject.get('approved_by_position') or '',
    }, template_profiles=confirmed_profiles)

    if form.validate_on_submit():
        update_data = {
            'department': form.department.data,
            'course_code': form.course_code.data,
            'course_title': form.course_title.data,
            'course_description': form.course_description.data,
            'type_of_course': form.type_of_course.data,
            'units': form.units.data,
            'contact_hours': form.contact_hours.data,
            'pre_requisites': form.pre_requisites.data,
            'co_requisites': form.co_requisites.data,
            'class_schedule': form.class_schedule.data,
            'room_assignment': form.room_assignment.data,
            'service_learning_component': form.service_learning_component.data,
            'target_sdg': form.target_sdg.data,
            'prepared_by_name': form.prepared_by_name.data,
            'prepared_by_position': form.prepared_by_position.data,
            'reviewed_by_name': form.reviewed_by_name.data,
            'reviewed_by_position': form.reviewed_by_position.data,
            'endorsed_by_name': form.endorsed_by_name.data,
            'endorsed_by_position': form.endorsed_by_position.data,
            'approved_by_name': form.approved_by_name.data,
            'approved_by_position': form.approved_by_position.data,
            'updated_at': datetime.utcnow().isoformat(),
        }
        if form.template_profile_id.data:
            update_data['template_profile_id'] = int(form.template_profile_id.data)
        else:
            update_data['template_profile_id'] = None

        try:
            supabase.table('teacher_subjects').update(update_data).eq('id', subject_id).eq('user_id', user_id).execute()
            flash('Subject updated.', 'success')
            return redirect(url_for('teacher.manage_subjects'))
        except Exception as exc:
            flash(f'Could not update subject: {exc}', 'danger')

    return render_template(
        'teacher_subject_form.html',
        form=form,
        mode='edit',
        subject=subject,
        active_semester=subject.get('semester') or active_semester,
        active_academic_year=subject.get('academic_year') or active_academic_year,
    )


@teacher_bp.route('/subjects/<int:subject_id>/delete', methods=['POST'])
@login_required
@roles_required('teacher')
def delete_subject(subject_id):
    try:
        supabase.table('teacher_subjects').delete().eq('id', subject_id).eq('user_id', session.get('user_id')).execute()
        flash('Subject deleted.', 'success')
    except Exception as exc:
        flash(f'Could not delete subject: {exc}', 'danger')
    return redirect(url_for('teacher.manage_subjects'))


# ---------------------------------------------------------------------------
# Template profile management
# ---------------------------------------------------------------------------

@teacher_bp.route('/template-profiles', methods=['GET'])
@login_required
@roles_required('teacher')
def list_template_profiles():
    user_id = session.get('user_id')
    result = supabase.table('teacher_template_profiles').select('*').eq('user_id', user_id).order('created_at', desc=True).execute()
    profiles = result.data or []
    # Load cloneable department defaults (not owned by this user).
    user_profile = get_current_user_profile() or {}
    department = user_profile.get('assigned_department', '')
    cloneable = []
    try:
        clone_q = supabase.table('teacher_template_profiles').select(
            'id,name,department,source_filename,user_id,is_department_default,users(first_name,last_name)'
        ).eq('status', 'confirmed').neq('user_id', user_id)
        if department:
            clone_q = clone_q.eq('department', department)
        clone_result = clone_q.order('is_department_default', desc=True).limit(20).execute()
        cloneable = clone_result.data or []
    except Exception:
        pass
    confirmed_count = sum(1 for p in profiles if p.get('status') == 'confirmed')
    return render_template('teacher_template_profiles.html', profiles=profiles, cloneable=cloneable, confirmed_count=confirmed_count)


@teacher_bp.route('/template-profiles/<int:profile_id>/clone', methods=['POST'])
@login_required
@roles_required('teacher')
def clone_template_profile(profile_id):
    """Clone a confirmed profile into the current teacher's profile set."""
    user_id = session.get('user_id')
    # Load source profile (must be confirmed, not owned by this user or owned is fine too).
    result = supabase.table('teacher_template_profiles').select('*').eq('id', profile_id).eq('status', 'confirmed').single().execute()
    if not result.data:
        flash('Source profile not found or not confirmed.', 'danger')
        return redirect(url_for('teacher.list_template_profiles'))
    source = result.data

    # Check user doesn't already have too many profiles.
    existing = supabase.table('teacher_template_profiles').select('id').eq('user_id', user_id).execute()
    if len(existing.data or []) >= 10:
        flash('You have reached the maximum number of profiles (10). Delete an old one first.', 'warning')
        return redirect(url_for('teacher.list_template_profiles'))

    try:
        new_profile = {
            'user_id': user_id,
            'department': source.get('department', ''),
            'name': f"{source.get('name', 'Profile')} (cloned)",
            'profile_data': source.get('profile_data', '{}'),
            'source_filename': source.get('source_filename'),
            'source_file_hash': source.get('source_file_hash'),
            'source_storage_path': source.get('source_storage_path'),
            'status': 'confirmed',
            'confirmed_at': datetime.utcnow().isoformat(),
        }
        supabase.table('teacher_template_profiles').insert(new_profile).execute()
        flash(f"Profile \"{source.get('name')}\" cloned into your profiles.", 'success')
    except Exception as exc:
        flash(f'Clone failed: {exc}', 'danger')

    return redirect(url_for('teacher.list_template_profiles'))


def _build_insertion_template(profile_id, user_id, source_bytes, source_filename):
    """Generate a placeholder DOCX from the source and store it.

    Runs ``generate_template_from_docx`` on *source_bytes* to produce a copy
    of the teacher's template where all variable content is replaced by
    canonical ``{{placeholder}}`` tokens.  This "insertion template" is stored
    at a deterministic storage path and the path is returned so it can be
    saved back to ``profile_data['insertion_template_path']``.

    Returns ``(insertion_path, placeholder_summary_dict)``.
    Raises on any storage or generation failure (caller decides how to handle).
    """
    from app.services.template_ai_service import generate_template_from_docx, summarize_placeholder_summary
    insertion_bytes, raw_summary = generate_template_from_docx(
        bytes(source_bytes),
        source_filename or 'source.docx',
        use_ai=False,
    )
    insertion_path = f"{user_id}/template_profiles/{profile_id}_insertion.docx"
    write_storage_bytes(STORAGE_BUCKET_NAME, insertion_path, insertion_bytes)
    return insertion_path, summarize_placeholder_summary(raw_summary)


def _set_insertion_template_metadata(profile_data, insertion_path, insertion_summary):
    profile_data['insertion_template_path'] = insertion_path
    profile_data['insertion_placeholder_count'] = insertion_summary.get('placeholder_count', 0)
    profile_data['insertion_algorithm_version'] = insertion_summary.get('algorithm_version', '')
    profile_data['insertion_placeholder_summary'] = insertion_summary
    profile_data['insertion_generated_at'] = datetime.now(timezone.utc).isoformat()


def _template_profile_generated_doc_key(profile_id, insertion_path, profile_data):
    version = profile_data.get('insertion_algorithm_version', '') if isinstance(profile_data, dict) else ''
    generated_at = profile_data.get('insertion_generated_at', '') if isinstance(profile_data, dict) else ''
    return hashlib.md5(
        f"template_profile_{profile_id}_{insertion_path}_{version}_{generated_at}".encode()
    ).hexdigest()


def _refresh_insertion_template_if_stale(profile_id, user_id, profile_row, profile_data, *, build_missing=False):
    """Regenerate stored placeholder DOCX when the generation algorithm changed."""
    if not isinstance(profile_data, dict):
        return False
    source_path = profile_row.get('source_storage_path')
    insertion_path = profile_data.get('insertion_template_path')
    if not source_path or (not insertion_path and not build_missing):
        return False

    from app.services.template_ai_service import TEMPLATE_INSERTION_ALGORITHM_VERSION

    if insertion_path and profile_data.get('insertion_algorithm_version') == TEMPLATE_INSERTION_ALGORITHM_VERSION:
        return False
    if not source_path:
        current_app.logger.warning(
            'Cannot refresh insertion template for profile %s: source_storage_path missing. '
            'Upload a new DOCX via the template profile page.',
            profile_id,
        )
        return False

    source_bytes = read_storage_bytes(STORAGE_BUCKET_NAME, source_path)
    refreshed_path, refreshed_summary = _build_insertion_template(
        profile_id,
        user_id,
        bytes(source_bytes),
        profile_row.get('source_filename', 'source.docx'),
    )
    _set_insertion_template_metadata(profile_data, refreshed_path, refreshed_summary)
    serialized_profile_data = json.dumps(profile_data)
    supabase.table('teacher_template_profiles').update({
        'profile_data': serialized_profile_data,
    }).eq('id', profile_id).eq('user_id', user_id).execute()
    profile_row['profile_data'] = serialized_profile_data
    return True


@teacher_bp.route('/template-profiles/upload', methods=['GET', 'POST'])
@login_required
@roles_required('teacher')
@limiter.limit("10 per hour")
def upload_template_profile():
    form = TemplateProfileUploadForm()
    user_profile = get_current_user_profile() or {}
    department = user_profile.get('assigned_department', '')

    if form.validate_on_submit():
        file = form.file.data
        if not file:
            flash('Please select a .docx file to upload.', 'warning')
            return redirect(url_for('teacher.upload_template_profile'))
        file_bytes = file.read()
        if not file_bytes:
            flash('Uploaded file is empty. Please choose a valid .docx file.', 'warning')
            return redirect(url_for('teacher.upload_template_profile'))
        name = form.name.data.strip()
        user_id = session.get('user_id')

        from app.services.template_profiler import profile_template, file_hash as compute_file_hash

        try:
            # Check for duplicate by file hash.
            fhash = compute_file_hash(file_bytes)
            existing = supabase.table('teacher_template_profiles').select('id,name').eq('user_id', user_id).eq('source_file_hash', fhash).execute()
            if existing.data:
                flash(f'This template file has already been profiled as "{existing.data[0]["name"]}".', 'warning')
                return redirect(url_for('teacher.list_template_profiles'))

            result = profile_template(
                file_bytes,
                department,
                user_id=user_id,
            )
        except ValueError as exc:
            flash(f'Template profiling failed: {exc}', 'danger')
            return redirect(url_for('teacher.upload_template_profile'))
        except Exception as exc:
            current_app.logger.exception("Template profile upload failed")
            flash(f'Upload failed: {exc}', 'danger')
            return redirect(url_for('teacher.upload_template_profile'))

        # Store the profile in draft status.
        profile_data = result['profile']
        validation_errors = result.get('validation_errors', [])
        doc_errors = result.get('doc_locator_errors', [])
        template_context = result.get('template_context', {})
        template_context_counts = result.get('template_context_counts', {})
        generation_spec = result.get('generation_spec', {})

        # Upload source template to storage for later rendering.
        safe_name = secure_filename(file.filename) or 'template.docx'
        storage_path = f"{user_id}/template_profiles/{fhash[:12]}_{safe_name}"
        write_storage_bytes(STORAGE_BUCKET_NAME, storage_path, file_bytes)

        row_data = {
            'user_id': user_id,
            'department': department,
            'name': name,
            'profile_data': json.dumps({
                'profile': profile_data,
                'template_context': template_context,
                'template_context_counts': template_context_counts,
                'generation_spec': generation_spec,
                'validation_errors': validation_errors,
                'doc_locator_errors': doc_errors,
                'structure_summary': {
                    'total_tables': result['structure'].get('total_tables', 0),
                    'total_paragraphs': result['structure'].get('total_paragraphs', 0),
                },
            }),
            'source_filename': safe_name,
            'source_file_hash': fhash,
            'source_storage_path': storage_path,
            'status': 'draft',
        }
        insert_result = supabase.table('teacher_template_profiles').insert(row_data).execute()
        profile_id = insert_result.data[0]['id']

        if validation_errors or doc_errors:
            flash(f'Profile created with {len(validation_errors)} schema warnings and {len(doc_errors)} locator warnings. Please review and confirm.', 'warning')
        else:
            extracted_total = sum(int(v or 0) for v in template_context_counts.values()) if isinstance(template_context_counts, dict) else 0
            flash(f'Template profiled successfully with {extracted_total} extracted context item(s). Review the profile below and confirm it.', 'success')

        return redirect(url_for('teacher.confirm_template_profile', profile_id=profile_id))

    return render_template('teacher_template_profile_upload.html', form=form, department=department)


@teacher_bp.route('/template-profiles/<int:profile_id>/confirm', methods=['GET', 'POST'])
@login_required
@roles_required('teacher')
def confirm_template_profile(profile_id):
    user_id = session.get('user_id')
    result = supabase.table('teacher_template_profiles').select('*').eq('id', profile_id).eq('user_id', user_id).single().execute()
    if not result.data:
        abort(404)
    profile_row = result.data
    profile_data = json.loads(profile_row.get('profile_data', '{}')) if isinstance(profile_row.get('profile_data'), str) else (profile_row.get('profile_data') or {})
    insertion_template_url = None
    try:
        if _refresh_insertion_template_if_stale(profile_id, user_id, profile_row, profile_data):
            flash('Generated insertion template refreshed to the latest placeholder layout.', 'info')
    except Exception as exc:
        current_app.logger.warning(
            'Could not refresh stale generated insertion template for profile %s: %s',
            profile_id,
            exc,
        )
    profile_obj = profile_data.get('profile') if isinstance(profile_data.get('profile'), dict) else None
    if profile_obj and any(
        isinstance(sec, dict) and isinstance(sec.get('locator'), dict) and sec.get('locator', {}).get('type') in {'table_index', 'table'}
        for sec in profile_obj.get('sections', [])
    ):
        source_path = profile_row.get('source_storage_path')
        if source_path:
            try:
                from app.services.template_profiler import normalize_profile_locators
                from app.services.template_profile import validate_profile, validate_profile_against_doc

                source_bytes = read_storage_bytes(STORAGE_BUCKET_NAME, source_path)
                source_doc = Document(BytesIO(bytes(source_bytes)))
                repaired_count = normalize_profile_locators(profile_obj, source_doc)
                if repaired_count:
                    profile_data['profile'] = profile_obj
                    profile_data['validation_errors'] = validate_profile(profile_obj)
                    profile_data['doc_locator_errors'] = validate_profile_against_doc(source_doc, profile_obj)
                    supabase.table('teacher_template_profiles').update({
                        'profile_data': json.dumps(profile_data),
                    }).eq('id', profile_id).eq('user_id', user_id).execute()
                    profile_row['profile_data'] = json.dumps(profile_data)
                    flash(f'Repaired {repaired_count} table-only locator(s) in this profile.', 'info')
            except Exception as exc:
                current_app.logger.warning('Could not repair table-only template profile locators for %s: %s', profile_id, exc)

    insertion_template_path = profile_data.get('insertion_template_path') if isinstance(profile_data, dict) else None
    if insertion_template_path:
        insertion_template_url = build_public_storage_url(STORAGE_BUCKET_NAME, insertion_template_path)

    if request.method == 'POST':
        action = request.form.get('action', 'confirm')
        if action == 'confirm':
            update_data = {
                'status': 'confirmed',
                'confirmed_at': datetime.utcnow().isoformat(),
            }
            source_path = profile_row.get('source_storage_path')
            if source_path:
                try:
                    src_bytes = read_storage_bytes(STORAGE_BUCKET_NAME, source_path)
                    insertion_path, ins_summary = _build_insertion_template(
                        profile_id, user_id, bytes(src_bytes),
                        profile_row.get('source_filename', 'source.docx'),
                    )
                    _set_insertion_template_metadata(profile_data, insertion_path, ins_summary)
                    update_data['profile_data'] = json.dumps(profile_data)
                    ph_count = ins_summary.get('placeholder_count', 0)
                    flash(
                        f'Template profile confirmed — {ph_count} insertion placeholder(s) generated '
                        f'from your template. CLP generation will now use your exact template layout.',
                        'success',
                    )
                except Exception as ins_exc:
                    current_app.logger.warning(
                        'Insertion template generation failed for profile %s: %s', profile_id, ins_exc,
                    )
                    flash(
                        'Template profile confirmed. (Insertion template could not be generated; '
                        'rendering will use locator-based fallback.)',
                        'warning',
                    )
            else:
                flash('Template profile confirmed! You can now use it when generating CLPs.', 'success')
            supabase.table('teacher_template_profiles').update(update_data).eq('id', profile_id).execute()
            return redirect(url_for('teacher.list_template_profiles'))
        elif action == 'delete':
            in_use = supabase.table('course_learning_plans').select('id').eq('template_profile_id', profile_id).limit(1).execute()
            if in_use.data:
                flash('This template profile is already used by one or more CLPs. Remove it from those plans before deleting.', 'warning')
                return redirect(url_for('teacher.confirm_template_profile', profile_id=profile_id))
            supabase.table('teacher_template_profiles').delete().eq('id', profile_id).eq('user_id', user_id).execute()
            flash('Profile draft discarded.', 'info')
            return redirect(url_for('teacher.list_template_profiles'))
        elif action == 'reprofile':
            return redirect(url_for('teacher.confirm_template_profile', profile_id=profile_id, show_reprofile=1))
        elif action == 'refresh_context':
            source_path = profile_row.get('source_storage_path')
            if not source_path:
                flash('No stored source DOCX is available. Please re-upload the template file to re-profile.', 'warning')
                return redirect(url_for('teacher.confirm_template_profile', profile_id=profile_id, show_reprofile=1))
            try:
                from app.services.template_profiler import profile_template
                file_bytes = read_storage_bytes(STORAGE_BUCKET_NAME, source_path)
                prof_result = profile_template(file_bytes, profile_row.get('department', ''), user_id=user_id)
                refreshed_data = {
                    'profile': prof_result.get('profile', {}),
                    'template_context': prof_result.get('template_context', {}),
                    'template_context_counts': prof_result.get('template_context_counts', {}),
                    'generation_spec': prof_result.get('generation_spec', {}),
                    'structure': prof_result.get('structure', {}),
                    'file_hash': prof_result.get('file_hash', profile_row.get('source_file_hash')),
                    'raw_ai_response': prof_result.get('raw_ai_response', ''),
                    'validation_errors': prof_result.get('validation_errors', []),
                    'doc_locator_errors': prof_result.get('doc_locator_errors', []),
                }
                try:
                    ins_path, ins_summary = _build_insertion_template(
                        profile_id,
                        user_id,
                        bytes(file_bytes),
                        profile_row.get('source_filename', 'source.docx'),
                    )
                    _set_insertion_template_metadata(refreshed_data, ins_path, ins_summary)
                except Exception as ins_exc:
                    current_app.logger.warning(
                        'Insertion template generation failed during stored-source refresh for %s: %s',
                        profile_id,
                        ins_exc,
                    )
                supabase.table('teacher_template_profiles').update({
                    'profile_data': json.dumps(refreshed_data),
                    'status': profile_row.get('status') or 'draft',
                    'confirmed_at': profile_row.get('confirmed_at'),
                }).eq('id', profile_id).eq('user_id', user_id).execute()
                flash('Template profile re-profiled from the stored DOCX and its insertion template was regenerated.', 'success')
            except Exception as exc:
                flash(f'Refresh failed: {exc}', 'danger')
            return redirect(url_for('teacher.confirm_template_profile', profile_id=profile_id))

    return render_template(
        'teacher_template_profile_confirm.html',
        profile_row=profile_row,
        profile_data=profile_data,
        profile_json=json.dumps(profile_data.get('profile', {}), indent=2),
        insertion_template_url=insertion_template_url,
    )


@teacher_bp.route('/template-profiles/<int:profile_id>/generated-template/view')
@login_required
@roles_required('teacher')
def view_template_profile_generated_template(profile_id):
    mode = request.args.get('mode', 'view')
    user_id = session.get('user_id')
    result = supabase.table('teacher_template_profiles').select('*').eq('id', profile_id).eq('user_id', user_id).single().execute()
    if not result.data:
        abort(404)
    profile_row = result.data
    profile_data = json.loads(profile_row.get('profile_data', '{}')) if isinstance(profile_row.get('profile_data'), str) else (profile_row.get('profile_data') or {})
    try:
        _refresh_insertion_template_if_stale(
            profile_id,
            user_id,
            profile_row,
            profile_data,
            build_missing=True,
        )
    except Exception as exc:
        current_app.logger.warning(
            'Could not refresh generated insertion template for profile %s before OnlyOffice view: %s',
            profile_id,
            exc,
        )
    insertion_path = profile_data.get('insertion_template_path') if isinstance(profile_data, dict) else None
    if not insertion_path or not str(insertion_path).lower().endswith('.docx'):
        flash('No generated insertion template is available for this profile yet.', 'warning')
        return redirect(url_for('teacher.confirm_template_profile', profile_id=profile_id))
    if not storage_file_exists(STORAGE_BUCKET_NAME, insertion_path):
        flash('The generated insertion template file could not be found. Refresh the profile to regenerate it.', 'warning')
        return redirect(url_for('teacher.confirm_template_profile', profile_id=profile_id))

    doc_key = _template_profile_generated_doc_key(profile_id, insertion_path, profile_data)
    base_url = get_onlyoffice_base_url(internal=True)
    doc_url = f"{base_url}/teacher/template-profiles/{profile_id}/generated-template/serve/{doc_key}"
    callback_url = f"{base_url}/teacher/template-profiles/{profile_id}/generated-template/callback/{doc_key}"
    doc_title = f"{profile_row.get('name') or 'Template Profile'} - Generated Template"

    is_editing = mode == 'edit'
    config = {
        "document": {
            "title": doc_title,
            "url": doc_url,
            "fileType": "docx",
            "key": doc_key,
            "permissions": {"edit": is_editing, "download": True, "review": is_editing},
        },
        "documentType": "word",
        "editorConfig": {
            "mode": "edit" if is_editing else "view",
            "callbackUrl": callback_url,
            "user": {"id": str(user_id), "name": session.get('username', 'Teacher')},
            "customization": {"autosave": True, "forcesave": True},
        },
        "width": "100%",
        "height": "100%",
    }
    token = generate_jwt_token(config) or ""
    return render_template(
        "teacher_view_document.html",
        config=config,
        doc_title=doc_title,
        doc_url=doc_url,
        callback_url=callback_url,
        doc_key=doc_key,
        token=token,
        back_url=url_for('teacher.confirm_template_profile', profile_id=profile_id),
    )


@teacher_bp.route('/template-profiles/<int:profile_id>/generated-template/serve/<doc_key>')
def serve_template_profile_generated_template(profile_id, doc_key):
    result = supabase.table('teacher_template_profiles').select('id,profile_data').eq('id', profile_id).single().execute()
    profile_row = result.data
    if not profile_row:
        abort(404)
    profile_data = json.loads(profile_row.get('profile_data', '{}')) if isinstance(profile_row.get('profile_data'), str) else (profile_row.get('profile_data') or {})
    insertion_path = profile_data.get('insertion_template_path') if isinstance(profile_data, dict) else None
    if not insertion_path:
        abort(404)
    expected_key = _template_profile_generated_doc_key(profile_id, insertion_path, profile_data)
    if doc_key != expected_key:
        abort(403)
    download_filename = os.path.basename(insertion_path).split('_', 1)[-1] or f"template_profile_{profile_id}.docx"
    return stream_storage_file(
        STORAGE_BUCKET_NAME,
        insertion_path,
        download_filename,
        mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        inline=True,
        cache_seconds=300,
    )


@teacher_bp.route('/template-profiles/<int:profile_id>/generated-template/callback/<doc_key>', methods=['POST'])
@csrf.exempt
def template_profile_generated_template_callback(profile_id, doc_key):
    return jsonify({"error": 0})


@teacher_bp.route('/template-profiles/<int:profile_id>/re-profile', methods=['POST'])
@login_required
@roles_required('teacher')
@limiter.limit("10 per hour")
def reprofile_template_profile(profile_id):
    """Re-upload a template to update an existing profile (re-runs AI profiling)."""
    user_id = session.get('user_id')
    result = supabase.table('teacher_template_profiles').select('*').eq('id', profile_id).eq('user_id', user_id).single().execute()
    if not result.data:
        abort(404)
    profile_row = result.data

    file = request.files.get('file')
    if not file or not file.filename.endswith('.docx'):
        flash('Please upload a .docx file.', 'danger')
        return redirect(url_for('teacher.confirm_template_profile', profile_id=profile_id))

    file_bytes = file.read()
    if len(file_bytes) < 1024:
        flash('File too small to be a valid DOCX template.', 'danger')
        return redirect(url_for('teacher.confirm_template_profile', profile_id=profile_id))

    import hashlib
    fhash = hashlib.sha256(file_bytes).hexdigest()

    try:
        # Save version snapshot of current profile before overwriting.
        try:
            ver_count_res = supabase.table('template_profile_versions').select('id').eq('profile_id', profile_id).execute()
            next_version = len(ver_count_res.data or []) + 1
            supabase.table('template_profile_versions').insert({
                'profile_id': profile_id,
                'version_number': next_version,
                'profile_data': profile_row.get('profile_data', '{}'),
                'source_filename': profile_row.get('source_filename'),
                'source_file_hash': profile_row.get('source_file_hash'),
                'note': 'Snapshot before re-profile',
            }).execute()
        except Exception:
            pass  # Non-critical if versioning fails.

        from app.services.template_profiler import profile_template
        department = profile_row.get('department', '')
        prof_result = profile_template(file_bytes, department, user_id=user_id)
        profile_data = {
            'profile': prof_result.get('profile', {}),
            'template_context': prof_result.get('template_context', {}),
            'template_context_counts': prof_result.get('template_context_counts', {}),
            'generation_spec': prof_result.get('generation_spec', {}),
            'structure': prof_result.get('structure', {}),
            'file_hash': prof_result.get('file_hash', fhash),
            'raw_ai_response': prof_result.get('raw_ai_response', ''),
            'validation_errors': prof_result.get('validation_errors', []),
            'doc_locator_errors': prof_result.get('doc_locator_errors', []),
        }
        # Upload new source template to storage.
        safe_name = secure_filename(file.filename) or 'template.docx'
        storage_path = f"{user_id}/template_profiles/{fhash[:12]}_{safe_name}"
        write_storage_bytes(STORAGE_BUCKET_NAME, storage_path, file_bytes)

        # Regenerate the insertion template from the new source DOCX.
        try:
            ins_path, ins_summary = _build_insertion_template(profile_id, user_id, file_bytes, safe_name)
            _set_insertion_template_metadata(profile_data, ins_path, ins_summary)
        except Exception as ins_exc:
            current_app.logger.warning(
                'Insertion template generation failed during re-profile for %s: %s', profile_id, ins_exc,
            )

        update_data = {
            'profile_data': json.dumps(profile_data),
            'source_filename': safe_name,
            'source_file_hash': fhash,
            'source_storage_path': storage_path,
            'status': 'draft',
            'confirmed_at': None,
        }
        supabase.table('teacher_template_profiles').update(update_data).eq('id', profile_id).execute()
        flash('Template re-profiled successfully. Please review and confirm the updated profile.', 'success')
    except Exception as exc:
        flash(f'Re-profiling failed: {exc}', 'danger')

    return redirect(url_for('teacher.confirm_template_profile', profile_id=profile_id))


@teacher_bp.route('/template-profiles/<int:profile_id>/versions')
@login_required
@roles_required('teacher')
def template_profile_versions(profile_id):
    """View version history of a template profile."""
    user_id = session.get('user_id')
    result = supabase.table('teacher_template_profiles').select('id,name,department').eq('id', profile_id).eq('user_id', user_id).single().execute()
    if not result.data:
        abort(404)
    profile_row = result.data
    versions = []
    try:
        ver_res = supabase.table('template_profile_versions').select('*').eq('profile_id', profile_id).order('version_number', desc=True).execute()
        versions = ver_res.data or []
    except Exception:
        pass
    return render_template('teacher_template_profile_versions.html', profile_row=profile_row, versions=versions)


@teacher_bp.route('/template-profiles/<int:profile_id>/export')
@login_required
@roles_required('teacher')
def export_template_profile(profile_id):
    """Export profile as downloadable JSON file."""
    user_id = session.get('user_id')
    result = supabase.table('teacher_template_profiles').select('*').eq('id', profile_id).eq('user_id', user_id).single().execute()
    if not result.data:
        abort(404)
    row = result.data
    profile_data = row.get('profile_data', '{}')
    if isinstance(profile_data, str):
        profile_data = json.loads(profile_data)

    export_payload = {
        'name': row.get('name'),
        'department': row.get('department'),
        'source_filename': row.get('source_filename'),
        'source_file_hash': row.get('source_file_hash'),
        'profile_data': profile_data,
        'exported_at': datetime.utcnow().isoformat(),
    }
    response = make_response(json.dumps(export_payload, indent=2))
    response.headers['Content-Type'] = 'application/json'
    safe_name = secure_filename(row.get('name', 'profile')) or 'profile'
    response.headers['Content-Disposition'] = f'attachment; filename="{safe_name}_profile.json"'
    return response


@teacher_bp.route('/template-profiles/import', methods=['POST'])
@login_required
@roles_required('teacher')
def import_template_profile():
    """Import a profile from an exported JSON file."""
    user_id = session.get('user_id')
    file = request.files.get('file')
    if not file or not file.filename.endswith('.json'):
        flash('Please upload a .json profile export file.', 'danger')
        return redirect(url_for('teacher.list_template_profiles'))

    try:
        data = json.loads(file.read().decode('utf-8'))
        if not isinstance(data, dict) or 'profile_data' not in data:
            flash('Invalid profile export file format.', 'danger')
            return redirect(url_for('teacher.list_template_profiles'))

        # Limit profile count.
        existing = supabase.table('teacher_template_profiles').select('id').eq('user_id', user_id).execute()
        if len(existing.data or []) >= 10:
            flash('You have reached the maximum number of profiles (10).', 'warning')
            return redirect(url_for('teacher.list_template_profiles'))

        user_profile = get_current_user_profile() or {}
        department = data.get('department') or user_profile.get('assigned_department', '')

        profile_data = data['profile_data']
        if isinstance(profile_data, dict):
            profile_data = json.dumps(profile_data)

        new_row = {
            'user_id': user_id,
            'department': department,
            'name': f"{data.get('name', 'Imported Profile')} (imported)",
            'profile_data': profile_data,
            'source_filename': data.get('source_filename'),
            'source_file_hash': data.get('source_file_hash'),
            'status': 'confirmed',
            'confirmed_at': datetime.utcnow().isoformat(),
        }
        supabase.table('teacher_template_profiles').insert(new_row).execute()
        flash('Profile imported successfully.', 'success')
    except json.JSONDecodeError:
        flash('File is not valid JSON.', 'danger')
    except Exception as exc:
        flash(f'Import failed: {exc}', 'danger')

    return redirect(url_for('teacher.list_template_profiles'))


@teacher_bp.route('/template-profiles/<int:profile_id>/delete', methods=['POST'])
@login_required
@roles_required('teacher')
def delete_template_profile(profile_id):
    user_id = session.get('user_id')
    try:
        in_use = supabase.table('course_learning_plans').select('id').eq('template_profile_id', profile_id).limit(1).execute()
        if in_use.data:
            flash('This template profile is attached to one or more CLPs. Remove it from those plans before deleting.', 'warning')
            return redirect(url_for('teacher.list_template_profiles'))
        supabase.table('teacher_template_profiles').delete().eq('id', profile_id).eq('user_id', user_id).execute()
        flash('Template profile deleted.', 'success')
    except Exception as exc:
        flash(f'Could not delete profile: {exc}', 'danger')
    return redirect(url_for('teacher.list_template_profiles'))


@teacher_bp.route('/copilot/check-profile-compatibility', methods=['POST'])
@login_required
@roles_required('teacher')
@limiter.limit("30 per hour")
def check_profile_compatibility():
    """AJAX endpoint — returns a compatibility assessment between a subject and a template profile."""
    data = request.get_json(silent=True) or {}
    subject_id = str(data.get('subject_id', '')).strip()
    profile_id = str(data.get('profile_id', '')).strip()

    if not profile_id:
        return jsonify({'status': 'ok', 'title': '', 'message': ''})

    # Resolve course details: from subject_id if available, else from direct fields
    course_dept = ''
    course_title = ''
    course_code = ''
    course_type = ''
    subject = None
    user_id = session.get('user_id')

    if subject_id:
        active_semester, active_academic_year = _get_active_term_settings()
        visible_subjects = _get_visible_teacher_subjects(get_current_user_profile() or {}, active_semester, active_academic_year)
        subject_lookup = {str(row.get('id')): row for row in visible_subjects}
        subject = subject_lookup.get(subject_id)
        if subject:
            course_dept = (subject.get('department') or '').strip()
            course_title = (subject.get('course_title') or '').strip()
            course_code = (subject.get('course_code') or '').strip()
            course_type = (subject.get('type_of_course') or '').strip().lower()
    else:
        course_dept = (data.get('department') or '').strip()
        course_code = (data.get('course_code') or '').strip()
        course_title = (data.get('course_title') or '').strip()

    if not course_code or not course_title:
        return jsonify({'status': 'ok', 'title': '', 'message': ''})

    profile = None
    try:
        pr = supabase.table('teacher_template_profiles').select('id,name,department,source_filename,profile_data').eq('id', profile_id).eq('user_id', user_id).eq('status', 'confirmed').limit(1).execute()
        profile = (pr.data or [None])[0]
    except Exception:
        pass
    if not profile:
        return jsonify({'status': 'ok', 'title': '', 'message': ''})

    # Re-extract from subject if available (prefer it over form data)
    if subject:
        course_dept = (subject.get('department') or '').strip()
        course_title = (subject.get('course_title') or '').strip()
        course_code = (subject.get('course_code') or '').strip()
        course_type = (subject.get('type_of_course') or '').strip().lower()
    profile_dept = (profile.get('department') or '').strip()
    profile_name = (profile.get('name') or '').strip()
    source_file = (profile.get('source_filename') or '').strip().upper()
    profile_data = {}
    try:
        profile_data = profile.get('profile_data') or {}
        if isinstance(profile_data, str):
            profile_data = json.loads(profile_data)
    except Exception:
        profile_data = {}

    def _text_tokens(*values):
        text = ' '.join(str(value or '') for value in values).lower()
        return re.sub(r'[^a-z0-9]+', ' ', text).strip()

    def _profile_generation_spec():
        if isinstance(profile_data.get('generation_spec'), dict):
            return profile_data.get('generation_spec') or {}
        pdata_profile = profile_data.get('profile') if isinstance(profile_data.get('profile'), dict) else {}
        if isinstance(pdata_profile.get('generation_spec'), dict):
            return pdata_profile.get('generation_spec') or {}
        return {}

    def _profile_metadata_values():
        values = []
        spec = _profile_generation_spec()
        for field in spec.get('metadata_fields', []) if isinstance(spec.get('metadata_fields'), list) else []:
            if not isinstance(field, dict):
                continue
            for key in ('value', 'detected_value', 'sample_value', 'label', 'field_key', 'aliases'):
                value = field.get(key)
                if isinstance(value, list):
                    values.extend(value)
                elif value:
                    values.append(value)
        template_context = profile_data.get('template_context') if isinstance(profile_data.get('template_context'), dict) else {}
        for key in ('course_metadata', 'metadata', 'course_info'):
            block = template_context.get(key)
            if isinstance(block, dict):
                values.extend(block.values())
            elif isinstance(block, list):
                values.extend(block)
        return values

    def _profile_clo_groups():
        spec = _profile_generation_spec()
        groups = spec.get('clo_groups') if isinstance(spec, dict) else []
        return groups if isinstance(groups, list) else []

    general_education_titles = {
        'the contemporary world',
        'understanding the self',
        'readings in philippine history',
        'mathematics in the modern world',
        'science technology and society',
        'art appreciation',
        'ethics',
        'purposive communication',
        'life and works of rizal',
    }
    course_title_key = _text_tokens(course_title)
    profile_metadata_text = _text_tokens(profile_name, profile.get('source_filename', ''), *_profile_metadata_values())
    profile_course_match = bool(course_title_key and course_title_key in profile_metadata_text)

    # --- Fast rule-based checks ---
    issues = []
    profile_groups = _profile_clo_groups()
    has_multi_program_groups = len(profile_groups) > 1 and any(
        isinstance(group, dict) and isinstance(group.get('program_scope'), dict) and group.get('program_scope')
        for group in profile_groups
    )
    profile_text = _text_tokens(profile_name, source_file, profile_metadata_text)
    is_gec_profile = (
        any(k in source_file for k in ['GEC', 'NSTP', 'CWTS', 'PATHFIT', 'PATH'])
        or has_multi_program_groups
        or any(title in profile_text for title in general_education_titles)
    )
    is_gec_course = (
        any(k in course_code.upper() for k in ['GEC', 'NSTP', 'CWTS', 'PATH'])
        or course_title_key in general_education_titles
        or any(title in course_title_key for title in general_education_titles)
        or profile_course_match
    )
    is_major_course = any(k in course_type for k in ['major', 'professional', 'core', 'elective']) and not is_gec_course

    if is_gec_profile and is_major_course:
        issues.append(f"The profile '{profile_name}' appears to be a GEC/general education template, but '{course_code} {course_title}' is a {course_type} course. GEC templates contain multi-program CLO alignment groups (e.g., Nursing, Accountancy) that are not applicable to a professional major course.")

    if profile_dept and course_dept and profile_dept.lower() != course_dept.lower() and not (is_gec_profile and is_gec_course):
        issues.append(f"The profile is tagged to '{profile_dept}' but the subject belongs to '{course_dept}'. The PLO alignment columns may reference the wrong program outcomes.")

    if is_gec_course and profile_dept and not is_gec_profile:
        issues.append(f"'{course_code}' looks like a GEC course but the selected profile is a department-specific template ({profile_dept}). The alignment structure may not match.")

    if issues:
        return jsonify({
            'status': 'warning',
            'title': 'Potential Template Mismatch Detected',
            'message': ' '.join(issues),
            'can_proceed': True,
        })

    # --- AI check for ambiguous cases (only if no rule-based issues found) ---
    profile_sections = []
    try:
        pd_profile = profile_data.get('profile') if isinstance(profile_data.get('profile'), dict) else profile_data
        profile_sections = [s.get('label') or s.get('section') for s in (pd_profile.get('sections') or []) if isinstance(s, dict)]
    except Exception:
        pass

    try:
        ai_client = current_app.config.get('AI_CLIENT') or AIClient()
        prompt = f"""You are a curriculum quality advisor. Assess whether this template profile is compatible with the given course.

COURSE:
- Code: {course_code}
- Title: {course_title}
- Department: {course_dept}
- Type: {course_type}

TEMPLATE PROFILE:
- Name: {profile_name}
- Department tag: {profile_dept or 'unset'}
- Source file: {profile.get('source_filename', '')}
- Detected sections: {json.dumps(profile_sections[:12])}

Respond with a JSON object only, no markdown:
{{
  "compatible": true or false,
  "confidence": "high" or "medium" or "low",
  "issue": "one concise sentence if not compatible, otherwise empty string"
}}
Only flag incompatibility if there is a clear academic mismatch (wrong program, wrong course level, etc.). When in doubt, return compatible=true."""
        resp = ai_client.generate_content(prompt)
        raw = (resp.text if hasattr(resp, 'text') else str(resp)).strip()
        raw = re.sub(r'^```[a-z]*\n?|\n?```$', '', raw).strip()
        ai_result = json.loads(raw)
        if not ai_result.get('compatible') and ai_result.get('confidence') in ('high', 'medium') and ai_result.get('issue'):
            return jsonify({
                'status': 'warning',
                'title': 'Potential Template Mismatch Detected',
                'message': ai_result['issue'],
                'can_proceed': True,
            })
    except Exception:
        pass

    return jsonify({'status': 'ok', 'title': 'Looks compatible', 'message': f"'{profile_name}' appears suitable for this course."})


@teacher_bp.route('/copilot/create', methods=['GET', 'POST'])
@teacher_bp.route('/copilot/beta/create', methods=['GET', 'POST'])
@login_required
@roles_required('teacher')
@limiter.limit("5 per hour")
def create_clp_ai_beta():
    form = AICopilotBetaForm()
    user_profile = get_current_user_profile() or {}
    active_semester, active_academic_year = _get_active_term_settings()
    visible_subjects = _get_visible_teacher_subjects(user_profile, active_semester, active_academic_year)
    form.subject_id.choices = [
        (str(row.get('id')), f"{row.get('course_code', '')} — {row.get('course_title', '')}{' (Department)' if row.get('user_id') is None else ''}")
        for row in visible_subjects
    ]

    if not _user_has_consultation_hours(user_profile):
        consultation_needed = True
    else:
        consultation_needed = False

    if not visible_subjects:
        flash('No subjects found for the active semester yet. Add your subjects first.', 'warning')
        return redirect(url_for('teacher.manage_subjects'))

    # Load confirmed template profiles for the teacher.
    user_id = session.get('user_id')
    confirmed_profiles = []
    try:
        tp_result = supabase.table('teacher_template_profiles').select('id,name,department,source_filename').eq('user_id', user_id).eq('status', 'confirmed').order('created_at', desc=True).execute()
        confirmed_profiles = tp_result.data or []
    except Exception:
        pass

    # Signatory data for pre-filling
    subject_signatories = {}
    for subj in visible_subjects:
        subject_signatories[str(subj['id'])] = {
            'prepared_by_name': subj.get('prepared_by_name') or '',
            'prepared_by_position': subj.get('prepared_by_position') or '',
            'reviewed_by_name': subj.get('reviewed_by_name') or '',
            'reviewed_by_position': subj.get('reviewed_by_position') or '',
            'endorsed_by_name': subj.get('endorsed_by_name') or '',
            'endorsed_by_position': subj.get('endorsed_by_position') or '',
            'approved_by_name': subj.get('approved_by_name') or '',
            'approved_by_position': subj.get('approved_by_position') or '',
        }

    if form.validate_on_submit():
        subject_lookup = {str(row.get('id')): row for row in visible_subjects}
        selected_subject = subject_lookup.get(str(form.subject_id.data))
        if not selected_subject:
            flash('Selected subject is not available for this term.', 'danger')
            return redirect(url_for('teacher.create_clp_ai_beta'))

        department_value = selected_subject.get('department')
        course_title = selected_subject.get('course_title')

        # Auto-detect template profile: first check subject's linked profile
        template_profile_id = None
        subject_profile_id = selected_subject.get('template_profile_id')
        if subject_profile_id:
            try:
                tp_check = supabase.table('teacher_template_profiles').select('id').eq('id', subject_profile_id).eq('status', 'confirmed').limit(1).execute()
                if tp_check.data:
                    template_profile_id = int(subject_profile_id)
            except Exception:
                pass

        # If not auto-detected, check form-selected profile
        if not template_profile_id:
            selected_profile_id = request.form.get('template_profile_id', '')
            if selected_profile_id and selected_profile_id.isdigit():
                profile_ids = {str(p['id']) for p in confirmed_profiles}
                if selected_profile_id in profile_ids:
                    template_profile_id = int(selected_profile_id)

        # Fallback to department default profile if none resolved
        if not template_profile_id and department_value:
            try:
                dept_default = supabase.table('teacher_template_profiles').select('id').eq('department', department_value).eq('status', 'confirmed').eq('is_department_default', True).limit(1).execute()
                if dept_default.data:
                    template_profile_id = dept_default.data[0]['id']
            except Exception:
                pass

        generation_spec = _load_template_profile_generation_spec(template_profile_id, user_id=None) if template_profile_id else None
        course_data = {
            'department': department_value,
            'course_code': selected_subject.get('course_code'),
            'course_title': course_title,
            'course_description': selected_subject.get('course_description'),
            'type_of_course': selected_subject.get('type_of_course'),
            'unit': selected_subject.get('units'),
            'contact_hours_per_week': selected_subject.get('contact_hours'),
            'pre_requisites': selected_subject.get('pre_requisites'),
            'co_requisites': selected_subject.get('co_requisites'),
            'class_schedule': selected_subject.get('class_schedule') or '',
            'room_assignment': selected_subject.get('room_assignment') or '',
            'service_learning_component': selected_subject.get('service_learning_component') or form.service_learning_component.data,
            'target_sdgs_display': selected_subject.get('target_sdg') or '',
            'source_context': form.source_context.data,
            'template_generation_spec': generation_spec,
        }
        content = build_initial_beta_content(course_data, user_profile=user_profile, generation_spec=generation_spec)

        # Override signatories with user-provided form values
        signatories = content.setdefault("signatories", {})
        for skey in ["prepared_by_name", "prepared_by_position", "reviewed_by_name", "reviewed_by_position",
                      "endorsed_by_name", "endorsed_by_position", "approved_by_name", "approved_by_position"]:
            val = request.form.get(skey, "").strip()
            if val:
                signatories[skey] = val

        metadata_errors = _beta_subject_metadata_errors(selected_subject)
        if metadata_errors:
            form.subject_id.errors.extend(metadata_errors)
            dept_signatories = {}
            if confirmed_profiles:
                dept_name = confirmed_profiles[0].get("department", "")
                if dept_name:
                    dept_signatories = get_department_signatory_settings(department_name=dept_name) or {}
            return render_template(
                'create_clp_ai_beta.html',
                form=form,
                active_semester=active_semester,
                active_academic_year=active_academic_year,
                alpha_mode=False,
                form_action=url_for('teacher.create_clp_ai_beta'),
                confirmed_profiles=confirmed_profiles,
                subject_signatories=subject_signatories,
                user_profile=user_profile,
                dept_signatories=dept_signatories,
                consultation_needed=consultation_needed,
            )

        insert_data = {
            "user_id": user_id,
            "subject": course_title,
            "department": department_value,
            "status": "generating",
            "upload_type": "ai_copilot_beta",
            "filename": "",
            "content": _dump_beta_content(content),
        }
        if template_profile_id:
            insert_data["template_profile_id"] = template_profile_id
        new_plan = supabase.table('course_learning_plans').insert(insert_data).execute()
        plan_id = new_plan.data[0]['id']

        supabase.table('course_learning_plans').update({
            "status": "beta_review",
            "content": _dump_beta_content(content),
        }).eq('id', plan_id).execute()
        try:
            task_id = TaskQueue.enqueue(
                "beta_action",
                {"action": "generate_all", "form_data": {}},
                user_id=user_id,
                plan_id=plan_id,
            )
            if task_id:
                flash('Draft created. Full AI generation has started in the background.', 'success')
            else:
                flash('Draft created, but AI generation could not be queued. You can start generation from the review page.', 'warning')
        except Exception as exc:
            current_app.logger.warning("Failed to queue fresh beta generation for plan %s: %s", plan_id, exc)
            flash('Draft created, but AI generation could not be queued. You can start generation from the review page.', 'warning')

        return redirect(url_for('teacher.review_clp_ai_beta', plan_id=plan_id))

    # Load department signatories for pre-filling the form
    dept_signatories = {}
    if confirmed_profiles:
        dept_name = confirmed_profiles[0].get("department", "")
        if dept_name:
            dept_signatories = get_department_signatory_settings(department_name=dept_name) or {}

    return render_template(
        'create_clp_ai_beta.html',
        form=form,
        active_semester=active_semester,
        active_academic_year=active_academic_year,
        alpha_mode=False,
        form_action=url_for('teacher.create_clp_ai_beta'),
        confirmed_profiles=confirmed_profiles,
        subject_signatories=subject_signatories,
        user_profile=user_profile,
        dept_signatories=dept_signatories,
        consultation_needed=consultation_needed,
    )


@teacher_bp.route('/copilot/review/<int:plan_id>', methods=['GET', 'POST'])
@teacher_bp.route('/copilot/beta/review/<int:plan_id>', methods=['GET', 'POST'])
@login_required
@roles_required('teacher')
def review_clp_ai_beta(plan_id):
    plan = _get_beta_plan_or_404(plan_id)
    raw_content = _coerce_json_object(plan.get('content')) or {}
    if plan.get('template_profile_id') and not isinstance(raw_content.get('template_generation_spec'), dict):
        generation_spec = _load_template_profile_generation_spec(plan.get('template_profile_id'))
        if generation_spec:
            raw_content['template_generation_spec'] = generation_spec
            raw_content['template_profile_warnings'] = generation_spec.get('warnings', [])
    content = normalize_beta_content(raw_content)

    if request.method == 'POST':
        updated_content = update_beta_content_from_form(content, request.form)
        _, _, updated_content = validate_beta_content(ensure_beta_shape(updated_content))
        supabase.table('course_learning_plans').update({
            'subject': updated_content['metadata'].get('course_title') or plan.get('subject'),
            'department': updated_content['metadata'].get('department') or plan.get('department'),
            'content': _dump_beta_content(updated_content),
        }).eq('id', plan_id).execute()
        flash('Draft saved.', 'success')
        return redirect(url_for('teacher.review_clp_ai_beta', plan_id=plan_id))

    cleaned_content, _ = enforce_beta_weekly_quality(content)
    _, _, cleaned_content = validate_beta_content(cleaned_content)
    if json.dumps(cleaned_content, sort_keys=True) != json.dumps(content, sort_keys=True):
        content = cleaned_content
        supabase.table('course_learning_plans').update({
            'content': _dump_beta_content(content),
        }).eq('id', plan_id).execute()
    else:
        content = cleaned_content

    display = beta_display_texts(content)
    stage_labels = get_review_stage_labels()
    # Ensure validation is fresh
    validation_version = content.get("_validation_version", 0)
    computed_at = content.get("_validation_computed_at", -1)
    if computed_at != validation_version:
        _, _, content = compute_validation(content)
    row_regen_controls_enabled = (
        'teacher.regenerate_beta_copilot_clo_row' in current_app.view_functions and
        'teacher.regenerate_beta_copilot_week_row' in current_app.view_functions
    )
    async_beta_controls_enabled = (
        'teacher.start_beta_background_task' in current_app.view_functions and
        'teacher.beta_background_task_status' in current_app.view_functions
    )
    # Load confirmed profiles for template-profile selector on review page.
    user_id = session.get('user_id')
    confirmed_profiles = []
    try:
        tp_result = supabase.table('teacher_template_profiles').select('id,name,department,source_filename').eq('user_id', user_id).eq('status', 'confirmed').order('created_at', desc=True).execute()
        confirmed_profiles = tp_result.data or []
    except Exception:
        pass

    return render_template(
        'edit_ai_clp_beta.html',
        plan=plan,
        content_data=content,
        display=display,
        alignment_context=get_static_alignment_context(),
        stage_labels=stage_labels,
        inserted_doc_ready=bool(content.get('beta_document_inserted') and plan.get('filename')),
        beta_document_error=content.get('beta_document_error'),
        context_source_log=get_context_source_log(
            _get_plan_template_context(plan),
            department_name=content.get('metadata', {}).get('department') or plan.get('department'),
        ),
        row_regen_controls_enabled=row_regen_controls_enabled,
        async_beta_controls_enabled=async_beta_controls_enabled,
        confirmed_profiles=confirmed_profiles,
        current_profile_id=plan.get('template_profile_id'),
    )


@teacher_bp.route('/copilot/review/<int:plan_id>/set-profile', methods=['POST'])
@login_required
@roles_required('teacher')
def set_plan_template_profile(plan_id):
    """Attach, change, or remove a template profile on an existing beta CLP."""
    plan = _get_beta_plan_or_404(plan_id)
    user_id = session.get('user_id')
    selected_id = request.form.get('template_profile_id', '').strip()

    if selected_id and selected_id.isdigit():
        # Verify ownership and confirmed status.
        check = supabase.table('teacher_template_profiles').select('id').eq('id', int(selected_id)).eq('user_id', user_id).eq('status', 'confirmed').execute()
        if not check.data:
            flash('Invalid template profile selection.', 'danger')
            return redirect(url_for('teacher.review_clp_ai_beta', plan_id=plan_id))
        content = _coerce_json_object(plan.get('content')) or {}
        generation_spec = _load_template_profile_generation_spec(int(selected_id), user_id=user_id)
        if generation_spec:
            content['template_generation_spec'] = generation_spec
            content['template_profile_warnings'] = generation_spec.get('warnings', [])
            content = ensure_beta_shape(normalize_beta_content(content))
        supabase.table('course_learning_plans').update({
            'template_profile_id': int(selected_id),
            'content': _dump_beta_content(content) if content else plan.get('content'),
        }).eq('id', plan_id).execute()
        flash('Template profile updated. It will be used on next document finalization.', 'success')
    else:
        # Remove profile association.
        content = _coerce_json_object(plan.get('content')) or {}
        content.pop('template_generation_spec', None)
        content.pop('template_profile_warnings', None)
        content = ensure_beta_shape(normalize_beta_content(content))
        supabase.table('course_learning_plans').update({
            'template_profile_id': None,
            'content': _dump_beta_content(content),
        }).eq('id', plan_id).execute()
        flash('Template profile removed. Default placeholder template will be used.', 'info')

    return redirect(url_for('teacher.review_clp_ai_beta', plan_id=plan_id))


def _load_template_profile_generation_spec(profile_id, user_id=None):
    if not profile_id:
        return None
    query = supabase.table('teacher_template_profiles').select('profile_data,status,source_storage_path').eq('id', profile_id).eq('status', 'confirmed')
    if user_id:
        query = query.eq('user_id', user_id)
    result = query.single().execute()
    if not result.data:
        return None
    profile_data = result.data.get('profile_data')
    if isinstance(profile_data, str):
        profile_data = json.loads(profile_data)
    source_doc = None
    if result.data.get('source_storage_path'):
        try:
            file_bytes = read_storage_bytes(STORAGE_BUCKET_NAME, result.data['source_storage_path'])
            source_doc = Document(BytesIO(bytes(file_bytes)))
        except Exception:
            source_doc = None
    return build_generation_spec(profile_data or {}, source_doc=source_doc)


def _get_plan_template_context(plan):
    profile_id = plan.get('template_profile_id') if isinstance(plan, dict) else None
    if not profile_id:
        return None
    try:
        result = supabase.table('teacher_template_profiles').select('profile_data,status').eq('id', profile_id).eq('status', 'confirmed').single().execute()
        if not result.data:
            return None
        profile_data = result.data.get('profile_data')
        if isinstance(profile_data, str):
            profile_data = json.loads(profile_data)
        template_context = (profile_data or {}).get('template_context')
        return template_context if isinstance(template_context, dict) else None
    except Exception:
        return None


def _get_plan_profile_hints(plan):
    """Load template profile hints for a plan (same as copilot_beta_tasks._get_plan_profile_hints)."""
    profile_id = plan.get('template_profile_id') if isinstance(plan, dict) else None
    if not profile_id:
        return None
    try:
        from app.services.copilot_beta_service import _build_template_profile_hints
        result = supabase.table('teacher_template_profiles').select('profile_data,status').eq('id', profile_id).eq('status', 'confirmed').single().execute()
        if not result.data:
            return None
        profile_data = result.data.get('profile_data')
        if isinstance(profile_data, str):
            profile_data = json.loads(profile_data)
        tp_profile = (profile_data or {}).get('profile', {})
        if isinstance(tp_profile, dict) and isinstance((profile_data or {}).get('generation_spec'), dict):
            tp_profile = {**tp_profile, 'generation_spec': (profile_data or {}).get('generation_spec')}
        return _build_template_profile_hints(tp_profile)
    except Exception:
        return None


@teacher_bp.route('/copilot/start-task/<int:plan_id>', methods=['POST'])
@login_required
@roles_required('teacher')
@limiter.limit("10 per hour")
def start_beta_background_task(plan_id):
    try:
        plan = _get_beta_plan_or_404(plan_id)
        action = (request.form.get('beta_async_action') or '').strip()
        allowed_actions = {
            'generate_all',
            'generate_clos',
            'generate_alignment',
            'generate_weekly',
            'regenerate_clo_row',
            'regenerate_week_row',
            'fix_final_review',
            'finalize_clp',
            'finalize_document',
        }
        if action not in allowed_actions:
            return jsonify({'ok': False, 'error': 'Unsupported beta action.'}), 400

        data, status = _start_beta_action_task(
            plan_id,
            action,
            request.form,
            clo_index=request.form.get('clo_index'),
            week_index=request.form.get('week_index'),
        )
        data['plan_id'] = plan.get('id')
        return jsonify(data), status
    except HTTPException:
        raise
    except Exception as exc:
        return _json_error_response('Could not start the beta task right now.', exc=exc)


@teacher_bp.route('/copilot/task-status/<int:task_id>', methods=['GET'])
@login_required
@roles_required('teacher')
def beta_background_task_status(task_id):
    try:
        task = TaskQueue.get_status(task_id)
        if not task:
            return jsonify({'ok': False, 'error': 'Task not found.'}), 404

        # Check in-memory cache (avoids DB hit for rapid consecutive polls)
        cache_key = f'beta_task_status:{task_id}'
        cached = cache.get(cache_key)
        if cached:
            return jsonify(cached)

        task_row = (current_app.config.get('SUPABASE_SERVICE') or supabase).table('background_tasks').select('user_id, payload, status, error_message, progress_percent, progress_label').eq('id', task_id).single().execute()
        data = task_row.data or {}
        if data.get('user_id') != session.get('user_id'):
            return jsonify({'ok': False, 'error': 'Unauthorized'}), 403

        payload = data.get('payload') if isinstance(data.get('payload'), dict) else {}
        result = {
            'ok': True,
            'status': data.get('status'),
            'error': data.get('error_message'),
            'percent': data.get('progress_percent') or 0,
            'label': data.get('progress_label') or 'Processing',
            'redirect_url': payload.get('redirect_url'),
            'ai_live_preview': payload.get('ai_live_preview') or '',
            'ai_live_task': payload.get('ai_live_task') or '',
        }
        # Cache for 2 seconds to absorb rapid polling
        try:
            cache.set(cache_key, result, timeout=2)
        except Exception:
            pass
        return jsonify(result)
    except HTTPException:
        raise
    except Exception as exc:
        return _json_error_response('Could not read beta task progress right now.', exc=exc)


@teacher_bp.route('/copilot/generate-clo/<int:plan_id>', methods=['POST'])
@teacher_bp.route('/copilot/beta/generate-clo/<int:plan_id>', methods=['POST'])
@login_required
@roles_required('teacher')
@limiter.limit("10 per hour")
def generate_beta_copilot_clos(plan_id):
    plan = _get_beta_plan_or_404(plan_id)
    return _redirect_after_beta_task_start(
        plan['id'],
        'generate_clos',
        'Draft CLO generation started in the background.',
    )


@teacher_bp.route('/copilot/generate-alignment/<int:plan_id>', methods=['POST'])
@teacher_bp.route('/copilot/beta/generate-alignment/<int:plan_id>', methods=['POST'])
@login_required
@roles_required('teacher')
@limiter.limit("10 per hour")
def generate_beta_copilot_alignment(plan_id):
    plan = _get_beta_plan_or_404(plan_id)
    return _redirect_after_beta_task_start(
        plan['id'],
        'generate_alignment',
        'Alignment generation started in the background.',
    )


@teacher_bp.route('/copilot/generate-weekly/<int:plan_id>', methods=['POST'])
@teacher_bp.route('/copilot/beta/generate-weekly/<int:plan_id>', methods=['POST'])
@login_required
@roles_required('teacher')
@limiter.limit("10 per hour")
def generate_beta_copilot_weekly(plan_id):
    plan = _get_beta_plan_or_404(plan_id)
    return _redirect_after_beta_task_start(
        plan['id'],
        'generate_weekly',
        'Weekly outline generation started in the background.',
    )


@teacher_bp.route('/copilot/regenerate-clo-row/<int:plan_id>/<int:clo_index>', methods=['POST'])
@teacher_bp.route('/copilot/beta/regenerate-clo-row/<int:plan_id>/<int:clo_index>', methods=['POST'])
@login_required
@roles_required('teacher')
@limiter.limit("10 per hour")
def regenerate_beta_copilot_clo_row(plan_id, clo_index):
    plan = _get_beta_plan_or_404(plan_id)
    current_content = ensure_beta_shape(normalize_beta_content(_coerce_json_object(plan.get('content')) or {}))
    max_clo_rows = len(current_content.get('clo_alignment_table') or [])
    if clo_index < 1 or clo_index > max_clo_rows:
        flash('Invalid CLO row selection.', 'danger')
        return redirect(url_for('teacher.review_clp_ai_beta', plan_id=plan_id))

    return _redirect_after_beta_task_start(
        plan['id'],
        'regenerate_clo_row',
        f'CLO {clo_index} regeneration started in the background.',
        clo_index=clo_index,
    )


@teacher_bp.route('/copilot/regenerate-week-row/<int:plan_id>/<int:week_index>', methods=['POST'])
@teacher_bp.route('/copilot/beta/regenerate-week-row/<int:plan_id>/<int:week_index>', methods=['POST'])
@login_required
@roles_required('teacher')
@limiter.limit("10 per hour")
def regenerate_beta_copilot_week_row(plan_id, week_index):
    plan = _get_beta_plan_or_404(plan_id)
    content = update_beta_content_from_form(_coerce_json_object(plan.get('content')) or {}, request.form)
    weekly_rows = content.get('weekly_course_outline') if isinstance(content.get('weekly_course_outline'), list) else []
    if week_index < 1 or week_index > len(weekly_rows):
        flash('Invalid weekly row selection.', 'danger')
        return redirect(url_for('teacher.review_clp_ai_beta', plan_id=plan_id))

    label = weekly_rows[week_index - 1].get('time_frame_label', 'Selected week')
    return _redirect_after_beta_task_start(
        plan['id'],
        'regenerate_week_row',
        f'{label} regeneration started in the background.',
        week_index=week_index,
    )


@teacher_bp.route('/copilot/finalize/<int:plan_id>', methods=['POST'])
@teacher_bp.route('/copilot/beta/finalize/<int:plan_id>', methods=['POST'])
@login_required
@roles_required('teacher')
@limiter.limit("10 per hour")
def finalize_beta_copilot(plan_id):
    plan = _get_beta_plan_or_404(plan_id)
    content = update_beta_content_from_form(_coerce_json_object(plan.get('content')) or {}, request.form)
    errors, warnings, validated_content = validate_beta_content(content, require_final_ready=True)
    validated_content.setdefault('validation', {})
    validated_content['validation']['final_review_issues'] = []
    validated_content['validation']['final_review_warnings'] = []
    validated_content['validation']['final_review_summary'] = ''
    if errors:
        for item in errors[:6]:
            flash(item, 'warning')
        flash('Draft is not ready yet. Fix the missing sections before finalizing.', 'danger')
        supabase.table('course_learning_plans').update({'content': _dump_beta_content(validated_content)}).eq('id', plan_id).execute()
        return redirect(url_for('teacher.review_clp_ai_beta', plan_id=plan_id))

    validated_content['beta_ready_for_template'] = True
    # Only clear the insert flag if the user hasn't already inserted a document;
    # re-finalizing after an insert should let them re-insert (not start over).
    if not content.get('beta_document_inserted'):
        validated_content['beta_document_inserted'] = False
        validated_content['beta_document_filename'] = ''
    validated_content['review_stage'] = 'beta_ready'
    supabase.table('course_learning_plans').update({
        'content': _dump_beta_content(validated_content),
        'status': 'beta_ready',
    }).eq('id', plan_id).execute()
    if warnings:
        flash('CLP finalized, but review the remaining warnings before inserting it into the template.', 'warning')
    else:
        flash('CLP finalized. You can now use "Insert Final Data Into Template" to generate a new populated DOCX file.', 'success')
    return redirect(url_for('teacher.review_clp_ai_beta', plan_id=plan_id))


@teacher_bp.route('/copilot/finalize-document/<int:plan_id>', methods=['POST'])
@teacher_bp.route('/copilot/beta/finalize-document/<int:plan_id>', methods=['POST'])
@login_required
@roles_required('teacher')
@limiter.limit("10 per hour")
def finalize_beta_copilot_document(plan_id):
    plan = _get_beta_plan_or_404(plan_id)
    content = update_beta_content_from_form(_coerce_json_object(plan.get('content')) or {}, request.form)
    errors, warnings, validated_content = validate_beta_content(content, require_final_ready=True)
    validated_content.setdefault('validation', {})
    validated_content['validation']['final_review_issues'] = []
    validated_content['validation']['final_review_warnings'] = []
    validated_content['validation']['final_review_summary'] = ''
    if errors:
        try:
            with beta_ai_call_context(plan_id=plan_id, user_id=session.get('user_id')):
                repaired_content = run_beta_final_fix(content, issues=errors, warnings=warnings)
            repaired_errors, repaired_warnings, repaired_content = validate_beta_content(repaired_content, require_final_ready=True)
            repaired_content.setdefault('validation', {})
            repaired_content['validation']['final_review_issues'] = []
            repaired_content['validation']['final_review_warnings'] = []
            repaired_content['validation']['final_review_summary'] = 'AI repair pass ran automatically before template insertion.'
            if not repaired_errors:
                validated_content = repaired_content
                warnings = repaired_warnings
                flash('AI repair pass fixed the blocked draft issues before template insertion. Continuing with document generation now.', 'info')
            else:
                supabase.table('course_learning_plans').update({'content': _dump_beta_content(repaired_content)}).eq('id', plan_id).execute()
                flash(f'Finalize document is still blocked after an AI repair pass. {len(repaired_errors)} issue(s) remain in the validation panel below.', 'danger')
                return redirect(url_for('teacher.review_clp_ai_beta', plan_id=plan_id))
        except Exception as repair_exc:
            current_app.logger.error(f"Beta pre-insert auto-fix failed for {plan_id}: {repair_exc}")
            flash(f'Finalize document is blocked until the draft is fully complete. {len(errors)} issue(s) are listed in the validation panel below.', 'danger')
            supabase.table('course_learning_plans').update({'content': _dump_beta_content(validated_content)}).eq('id', plan_id).execute()
            return redirect(url_for('teacher.review_clp_ai_beta', plan_id=plan_id))

    try:
        cached_review = get_cached_beta_final_review(validated_content)
        if cached_review:
            final_review = cached_review
        else:
            with beta_ai_call_context(plan_id=plan_id, user_id=session.get('user_id')):
                final_review = run_beta_final_review(validated_content)
        validated_content.setdefault('validation', {})
        validated_content['validation']['final_review_signature'] = get_beta_final_review_signature(validated_content)
        validated_content['validation']['final_review_approved'] = bool(final_review.get('approved'))
        validated_content['validation']['final_review_summary'] = final_review.get('summary', '')
        validated_content['validation']['final_review_warnings'] = final_review.get('warnings', [])
        review_issues = []
        if not final_review.get('approved'):
            review_issues = final_review.get('issues', []) or ['Final AI review did not approve this draft yet.']
            validated_content['validation']['final_review_issues'] = []
            validated_content['validation']['errors'] = []
            validated_content['validation']['final_review_warnings'] = list(dict.fromkeys(
                (validated_content['validation'].get('final_review_warnings') or []) +
                review_issues
            ))
            validated_content['validation']['warnings'] = list(dict.fromkeys(
                (validated_content['validation'].get('warnings') or []) +
                final_review.get('warnings', []) +
                review_issues
            ))

        user_profile = get_current_user_profile() or {}
        template_label = ''
        template_source = ''
        profile_used = False
        doc = None
        template_profile_id = plan.get('template_profile_id')
        enforce_profile_template = bool(template_profile_id)
        if template_profile_id:
            try:
                profile_row = supabase.table('teacher_template_profiles').select('*').eq('id', template_profile_id).eq('status', 'confirmed').single().execute()
                if profile_row.data:
                    template_label = profile_row.data.get('name', 'Custom Template Profile')
                    template_source = 'template_profile'
                    profile_data_raw = profile_row.data.get('profile_data')
                    if isinstance(profile_data_raw, str):
                        profile_data_raw = json.loads(profile_data_raw)
                    tp_profile = (profile_data_raw or {}).get('profile', {})
                    source_path = profile_row.data.get('source_storage_path')
                    if tp_profile and source_path:
                        source_bytes = read_storage_bytes(STORAGE_BUCKET_NAME, source_path)
                        source_doc_for_spec = Document(BytesIO(bytes(source_bytes)))
                        generation_spec = build_generation_spec(profile_data_raw or {}, source_doc=source_doc_for_spec)
                        tp_profile = merge_spec_into_profile(tp_profile, generation_spec)
                        validated_content['template_generation_spec'] = generation_spec
                        validated_content['template_profile_warnings'] = generation_spec.get('warnings', [])
                        validated_content = ensure_beta_shape(validated_content)
                        from app.services.template_content_adapter import build_profile_content, validate_profile_coverage
                        from app.services.template_writer import apply_profile

                        coverage_gaps = validate_profile_coverage(validated_content, tp_profile, user_profile=user_profile)
                        if coverage_gaps:
                            existing_warnings = validated_content.get('validation', {}).get('warnings', [])
                            for gap in coverage_gaps:
                                prefixed = f'[Template Profile] {gap}'
                                if prefixed not in existing_warnings:
                                    existing_warnings.append(prefixed)
                            validated_content.setdefault('validation', {})['warnings'] = existing_warnings
                        profile_content = build_profile_content(validated_content, tp_profile, user_profile=user_profile)
                        doc = Document(BytesIO(bytes(source_bytes)))
                        doc, write_results = apply_profile(doc, tp_profile, profile_content)
                        ok_count = sum(1 for v in write_results.values() if v == 'ok')
                        fail_count = sum(1 for v in write_results.values() if v != 'ok')
                        template_label = profile_row.data.get('name', 'Custom Template Profile')
                        template_source = 'template_profile'
                        profile_used = True
                        validated_content['beta_profile_render_ok'] = ok_count
                        validated_content['beta_profile_render_fail'] = fail_count
                        validated_content['beta_profile_coverage_gaps'] = len(coverage_gaps) if coverage_gaps else 0
            except Exception as profile_exc:
                current_app.logger.warning('Profile-based beta document rendering failed for %s; falling back to placeholder template: %s', plan_id, profile_exc)
                profile_used = False

        if enforce_profile_template and not profile_used:
            raise ValueError(
                'Selected template profile could not be rendered. '
                'Please re-confirm or re-profile the template profile, then try again.'
            )

        if not profile_used:
            template_bytes, template_label, template_source = _resolve_beta_output_template(plan)
            doc = Document(BytesIO(bytes(template_bytes)))

        replacements = build_beta_docx_replacements(validated_content, user_profile=user_profile)
        doc = replace_placeholders(doc, replacements)

        output = BytesIO()
        doc.save(output)
        output.seek(0)

        subject_name = secure_filename(
            validated_content.get('metadata', {}).get('course_title') or plan.get('subject') or f"beta_clp_{plan_id}"
        ) or f"beta_clp_{plan_id}"
        new_filename = f"{session.get('user_id')}/{subject_name}_{int(time.time())}.docx"
        old_filename = plan.get('filename')

        write_storage_bytes(STORAGE_BUCKET_NAME, new_filename, output.read())
        if not storage_file_exists(STORAGE_BUCKET_NAME, new_filename):
            raise FileNotFoundError(f'Generated template file was not found after upload: {new_filename}')

        validated_content['beta_ready_for_template'] = True
        validated_content['beta_document_inserted'] = True
        validated_content['beta_document_filename'] = new_filename
        validated_content['beta_template_source_label'] = template_label
        validated_content['beta_template_source_type'] = template_source
        validated_content['beta_template_copy_mode'] = True
        validated_content['review_stage'] = 'beta_inserted'
        validated_content['validation']['warnings'] = list(dict.fromkeys((validated_content['validation'].get('warnings') or []) + final_review.get('warnings', [])))

        # Clear any previous error marker before finalising the DB update.
        validated_content.pop('beta_document_error', None)

        try:
            supabase.table('course_learning_plans').update({
                'subject': validated_content['metadata'].get('course_title') or plan.get('subject'),
                'department': validated_content['metadata'].get('department') or plan.get('department'),
                'content': _dump_beta_content(validated_content),
                'filename': new_filename,
                'status': 'beta_ready',
            }).eq('id', plan_id).execute()
        except Exception as db_exc:
            # DB update failed — roll back the storage write so the plan is
            # not left with a dangling file and a missing filename.
            current_app.logger.error("Beta finalize document DB update failed for %s; rolling back storage write: %s", plan_id, db_exc)
            try:
                delete_storage_paths(STORAGE_BUCKET_NAME, [new_filename])
            except Exception as cleanup_exc:
                current_app.logger.warning("Beta finalize document storage cleanup failed for %s: %s", plan_id, cleanup_exc)
            validated_content['beta_document_inserted'] = False
            validated_content['beta_document_filename'] = ''
            validated_content['beta_document_error'] = {
                'time': int(time.time()),
                'message': str(db_exc),
            }
            supabase.table('course_learning_plans').update({
                'content': _dump_beta_content(validated_content),
            }).eq('id', plan_id).execute()
            raise

        if old_filename and old_filename != new_filename:
            cleanup_storage_after_commit(
                old_filename,
                context='Beta finalize document',
                user_id=session.get('user_id'),
                plan_id=plan_id,
            )

        log_system_event(
            'workflow',
            'info',
            'Beta copilot draft rendered into DOCX template',
            details={'plan_id': plan_id, 'template_label': template_label, 'template_source': template_source, 'filename': new_filename},
            user_id=session.get('user_id'),
            plan_id=plan_id,
        )

        if review_issues:
            flash(f'A fresh copy of the template was generated and opened in OnlyOffice, but the final AI review flagged {len(review_issues)} issue(s). Review the inserted document and the warning panel carefully.', 'warning')
        elif warnings or final_review.get('warnings'):
            flash('A new populated copy of the template was generated. Review the generated document carefully before final submission.', 'warning')
        else:
            flash('A new copy of the template was generated with the final CLP data inserted. You can review and edit that file now.', 'success')
        return redirect(url_for('teacher.edit_clp_document', plan_id=plan_id))
    except Exception as exc:
        current_app.logger.error(f"Beta document finalization failed for {plan_id}: {exc}")
        flash(f'Could not finalize the draft into the DOCX template: {exc}', 'danger')
        return redirect(url_for('teacher.review_clp_ai_beta', plan_id=plan_id))


@teacher_bp.route('/copilot/fix-final-review/<int:plan_id>', methods=['POST'])
@teacher_bp.route('/copilot/beta/fix-final-review/<int:plan_id>', methods=['POST'])
@login_required
@roles_required('teacher')
@limiter.limit("10 per hour")
def fix_beta_copilot_final_review(plan_id):
    plan = _get_beta_plan_or_404(plan_id)
    content = update_beta_content_from_form(_coerce_json_object(plan.get('content')) or {}, request.form)
    validation = content.get('validation') if isinstance(content.get('validation'), dict) else {}
    issues = validation.get('final_review_issues', []) if isinstance(validation.get('final_review_issues'), list) else []
    if not issues:
        issues = validation.get('errors', []) if isinstance(validation.get('errors'), list) else []
    review_warnings = validation.get('final_review_warnings', []) if isinstance(validation.get('final_review_warnings'), list) else []
    if not issues:
        flash('There are no blocked final-review issues to auto-fix right now.', 'info')
        return redirect(url_for('teacher.review_clp_ai_beta', plan_id=plan_id))

    try:
        with beta_ai_call_context(plan_id=plan_id, user_id=session.get('user_id')):
            fixed_content = run_beta_final_fix(content, issues=issues, warnings=review_warnings)
        fixed_content.setdefault('validation', {})
        fixed_content['validation']['errors'] = []
        fixed_content['validation']['warnings'] = []
        fixed_content['validation']['final_review_issues'] = []
        fixed_content['validation']['final_review_warnings'] = []
        fixed_content['validation']['final_review_summary'] = ''

        base_errors, base_warnings, fixed_content = validate_beta_content(fixed_content, require_final_ready=False)
        with beta_ai_call_context(plan_id=plan_id, user_id=session.get('user_id')):
            final_review = run_beta_final_review(fixed_content)
        fixed_content['validation']['final_review_summary'] = final_review.get('summary', '')
        fixed_content['validation']['final_review_warnings'] = final_review.get('warnings', [])
        if final_review.get('approved'):
            fixed_content['validation']['errors'] = base_errors
            fixed_content['validation']['warnings'] = list(dict.fromkeys((base_warnings or []) + final_review.get('warnings', [])))
        else:
            review_issues = final_review.get('issues', []) or ['AI repair improved the draft, but final review still found issues to fix.']
            fixed_content['validation']['final_review_issues'] = review_issues
            fixed_content['validation']['errors'] = base_errors
            fixed_content['validation']['warnings'] = list(dict.fromkeys((base_warnings or []) + final_review.get('warnings', [])))
        # Preserve the current review stage instead of dropping back to weekly_generated.
        # The fix only repairs content; it doesn't invalidate subsequent steps.
        fixed_content['review_stage'] = fixed_content.get('review_stage') or content.get('review_stage', 'weekly_generated')
        supabase.table('course_learning_plans').update({
            'subject': fixed_content['metadata'].get('course_title') or plan.get('subject'),
            'department': fixed_content['metadata'].get('department') or plan.get('department'),
            'content': _dump_beta_content(fixed_content),
            'status': 'beta_review',
        }).eq('id', plan_id).execute()
        if final_review.get('approved'):
            flash('AI repair pass cleared the blocked final-review issues. You can try finalizing again.', 'success')
        else:
            flash('AI repair pass updated the draft, but the final review still found issues. Please review the refreshed validation list.', 'warning')
    except Exception as exc:
        current_app.logger.error(f"Beta final-review auto-fix failed for {plan_id}: {exc}")
        flash(f'AI repair could not complete the suggested fixes: {exc}', 'danger')
    return redirect(url_for('teacher.review_clp_ai_beta', plan_id=plan_id))


@teacher_bp.route('/copilot/discard/<int:plan_id>', methods=['POST'])
@teacher_bp.route('/copilot/beta/discard/<int:plan_id>', methods=['POST'])
@login_required
@roles_required('teacher')
def discard_beta_copilot(plan_id):
    _get_beta_plan_or_404(plan_id)
    delete_clp_with_dependencies(plan_id)
    flash('AI copilot draft discarded.', 'success')
    return redirect(url_for('teacher.teacher_my_clps'))

# --- AI CREATION FLOW (legacy, replaced by copilot beta workflow) ---
from app.services.clp_service import CLPService
from app.services.validator_service import ValidatorService

# ... existing imports ...

@teacher_bp.route('/clp/<int:plan_id>/validate', methods=['POST'])
@login_required
def validate_clp(plan_id):
    """
    Validates the CLP against Bloom's Taxonomy.
    """
    try:
        # Fetch the CLP content
        res = supabase.table('course_learning_plans').select('content, user_id').eq('id', plan_id).single().execute()
        if not res.data:
            return jsonify({'error': 'CLP not found.'}), 404

        if res.data.get('user_id') != session.get('user_id'):
            return jsonify({'error': 'Unauthorized'}), 403

        if not res.data.get('content'):
            return jsonify({'error': 'CLP content not found.'}), 404

        clp_content = _coerce_json_object(res.data.get('content'))
        if clp_content is None:
            return jsonify({'error': 'Invalid JSON content.'}), 400

        # Run Validation
        report = ValidatorService.validate_clp_outcomes(clp_content)

        return jsonify(report)
    except Exception as e:
        current_app.logger.error(f"Validation Failed: {e}")
        return jsonify({'error': str(e)}), 500

@teacher_bp.route('/clp/<int:plan_id>/apply_fixes', methods=['POST'])
@login_required
def apply_validation_fixes(plan_id):
    """
    Triggers an AI task to refine the CLP based on validation suggestions.
    """
    try:
        data = request.get_json()
        validation_data = data.get('validation_data') if data else None

        if not validation_data:
            return jsonify({'error': 'No validation data provided'}), 400

        res = supabase.table('course_learning_plans').select('content, user_id').eq('id', plan_id).single().execute()
        if not res.data: return jsonify({'error': 'Not found'}), 404
        if res.data.get('user_id') != session.get('user_id'): return jsonify({'error': 'Unauthorized'}), 403

        current_content = _coerce_json_object(res.data.get('content'))
        if current_content is None:
            return jsonify({'error': 'Invalid JSON content.'}), 400

        # Update status to generating
        supabase.table('course_learning_plans').update({'status': 'generating'}).eq('id', plan_id).execute()

        from app.utils import start_apply_validation_fixes
        start_apply_validation_fixes(plan_id, session.get('user_id'), current_content, validation_data)

        return jsonify({'success': True})
    except Exception as e:
        current_app.logger.error(f"Apply Fixes Init Failed: {e}")
        return jsonify({'error': str(e)}), 500

@teacher_bp.route('/clp/<int:plan_id>/assessments', methods=['POST'])
@login_required
def generate_assessments(plan_id):
    """
    Generates a bank of assessment questions based on the CLP content.
    """
    try:
        res = supabase.table('course_learning_plans').select('content, subject, user_id').eq('id', plan_id).single().execute()
        if not res.data: return jsonify({'error': 'Not found'}), 404

        if res.data.get('user_id') != session.get('user_id'):
            return jsonify({'error': 'Unauthorized'}), 403

        content = _coerce_json_object(res.data.get('content'))
        if content is None:
            return jsonify({'error': 'Invalid JSON content.'}), 400
        subject = res.data['subject']

        model_instance = AIClient.get_model()
        prompt = f"""
        Based on the following Course Learning Plan for "{subject}", generate a bank of 10 assessment questions.

        REQUIREMENTS:
        - Include 5 Multiple Choice Questions (with options A, B, C, D and the Correct Answer).
        - Include 3 True/False Questions.
        - Include 2 Essay/Discussion Questions.
        - Ensure questions align with the Learning Outcomes in the CLP.

        CLP CONTENT:
        {json.dumps(content)}

        Return the result as a clean, professionally formatted JSON object.
        """

        resp = AIClient.generate_with_retry(model_instance, [prompt], {"response_mime_type": "application/json"})
        text = AIClient.clean_ai_json(resp.text)

        return jsonify(json.loads(text))
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@teacher_bp.route('/clp/<int:plan_id>/lms-export')
@login_required
def lms_export(plan_id):
    """
    Exports the CLP weekly schedule as a CSV for LMS import (Canvas/Moodle).
    """
    # Fetch and verify ownership first
    try:
        res = supabase.table('course_learning_plans').select('content, subject, user_id').eq('id', plan_id).single().execute()
    except Exception as e:
        current_app.logger.error(f"DB Error in export: {e}")
        abort(500)

    if not res.data: abort(404)

    if res.data.get('user_id') != session.get('user_id'):
        abort(403)

    try:
        content = _coerce_json_object(res.data.get('content'))
        if content is None:
            flash('Invalid CLP content format.', 'danger')
            return redirect(url_for('teacher.view_clp', plan_id=plan_id))
        subject = res.data['subject']

        import csv
        from io import StringIO

        si = StringIO()
        cw = csv.writer(si)
        cw.writerow(['Title', 'Description', 'Type']) # Simplified LMS format

        # 1. Course Metadata
        code = content.get('course_code') or content.get('course_number', 'N/A')
        units = content.get('units', 'N/A')
        ctype = content.get('type_of_course', 'Lecture')
        cw.writerow(['Course Metadata', f"Code: {code} | Units: {units} | Type: {ctype}", 'Page'])

        # 2. Course Description
        desc = content.get('course_description', '')
        if desc:
            cw.writerow(['Course Description', desc, 'Page'])

        # 2. Outcomes Rationale
        rationale = content.get('po_io_rationale', '')
        if rationale:
            cw.writerow(['Alignment Rationale', rationale, 'Page'])

        # 3. Weekly Schedule
        for i in range(1, 19):
            prefix = f'W{i}'
            if i in [10, 11]: prefix = 'W1011'
            elif i in [14, 15]: prefix = 'W1415'
            elif i in [16, 17]: prefix = 'W1617'

            # Skip duplicate merged weeks to avoid duplicate LMS rows
            if i in [11, 15, 17]:
                continue

            topic = content.get(f'{prefix}_TO', '')
            lo = content.get(f'{prefix}_LO', '')
            assessment = content.get(f'{prefix}_Assesment', '')
            method = content.get(f'{prefix}_Method', '')
            mapped_cos = content.get(f'{prefix}_Mapped_COs', '')
            assessed_cos = content.get(f'{prefix}_Assessment_COs', '')

            if topic:
                week_label = str(i)
                if i == 10:
                    week_label = "10-11"
                elif i == 14:
                    week_label = "14-15"
                elif i == 16:
                    week_label = "16-17"

                cw.writerow([
                    f"Week {week_label}: {topic}",
                    (
                        f"Learning Outcomes: {lo}\n"
                        f"Method: {method}\n"
                        f"Assessment: {assessment}\n"
                        f"Mapped COs: {mapped_cos or '-'}\n"
                        f"Assessed COs: {assessed_cos or '-'}"
                    ),
                    "Assignment"
                ])

        output = si.getvalue()
        return Response(
            output,
            mimetype="text/csv",
            headers={"Content-disposition": f"attachment; filename={subject.replace(' ', '_')}_LMS_Export.csv"}
        )
    except Exception as e:
        flash(f"Export failed: {e}", "danger")
        return redirect(url_for('teacher.view_clp', plan_id=plan_id))

@teacher_bp.route('/clp/<int:plan_id>/finalize', methods=['POST'])
@login_required
@roles_required('teacher')
def finalize_clp(plan_id):
    """Triggers the document finalization process."""
    try:
        # Verify ownership
        res = supabase.table('course_learning_plans').select('user_id, status').eq('id', plan_id).single().execute()
        if not res.data or res.data['user_id'] != session.get('user_id'):
            return jsonify({'success': False, 'error': 'Unauthorized'}), 403

        if res.data['status'] not in ['draft_review', 'draft']:
            flash("This plan cannot be finalized right now.", "warning")
            return redirect(url_for('teacher.view_clp', plan_id=plan_id))

        # Start background finalization
        supabase.table('course_learning_plans').update({'status': 'generating_doc'}).eq('id', plan_id).execute()
        start_clp_finalization(plan_id, session['user_id'])
        flash("CLP finalization started. You'll be notified when the Word document is ready.", "info")
        return redirect(url_for('teacher.teacher_my_clps'))
    except Exception as e:
        flash(f"Error starting finalization: {e}", "danger")
        return redirect(url_for('teacher.view_clp', plan_id=plan_id))

# --- Legacy create_clp_ai route: redirects to the production AI Copilot ---
@teacher_bp.route('/create_clp_ai', methods=['GET'])
@login_required
@roles_required('teacher')
def create_clp_ai():
    return redirect(url_for('teacher.create_clp_ai_beta'))

# @teacher_bp.route('/create_clp_ai', methods=['GET', 'POST'])
# @login_required
# @roles_required('teacher')
# @limiter.limit("5 per hour")
# def create_clp_ai():
#     form = AIClpForm()
#     user_id = session.get('user_id')
#
#     if request.method == 'GET':
#         user_profile = get_current_user_profile()
#         if user_profile and user_profile.get('assigned_department'):
#             form.department.data = user_profile['assigned_department']
#
#     if form.validate_on_submit():
#         if user_id:
#             try:
#                 form_data = {
#                     'department': form.department.data,
#                     'course_title': form.course_title.data,
#                     'course_code': form.course_code.data,
#                     'course_description': form.course_description.data,
#                     'type_of_course': form.type_of_course.data,
#                     'unit': form.unit.data,
#                     'pre_requisites': form.pre_requisites.data,
#                     'co_requisites': form.co_requisites.data,
#                     'contact_hours_per_week': form.contact_hours_per_week.data,
#                     'class_schedule': form.class_schedule.data,
#                     'room_assignment': form.room_assignment.data
#                 }
#
#                 if form.department.choices:
#                     choices_dict = dict(form.department.choices)
#                     form_data['department'] = choices_dict.get(form.department.data, form.department.data)
#
#                 CLPService.create_ai_clp(user_id, form_data, session.get('username', 'Instructor'))
#
#                 flash("Draft generation started. Check back soon to review.", "info")
#                 return redirect(url_for('teacher.teacher_my_clps'))
#
#             except Exception as e:
#                 current_app.logger.error(f"CLP Creation Error: {e}")
#                 flash(f"Error initiating AI generation: {e}", "danger")
#         else:
#             return redirect(url_for('auth.login'))
#
#     return render_template('create_clp_ai.html', form=form)

@teacher_bp.route('/clp/<int:plan_id>/clone', methods=['POST'])
@login_required
@roles_required('teacher')
def clone_clp(plan_id):
    try:
        new_plan_id, is_legacy = CLPService.clone_clp(plan_id, session['user_id'])
        if is_legacy:
            flash(
                "This CLP was authored with the legacy editor. Only the metadata "
                "(course code, title, description, schedule, etc.) was cloned into a "
                "new AI Copilot draft — CLOs, weekly outline, and other sections "
                "need to be generated again.",
                "warning",
            )
        else:
            flash("Plan cloned into a new AI Copilot draft. You can now edit your copy.", "success")
        return redirect(url_for('teacher.review_clp_ai_beta', plan_id=new_plan_id))
    except Exception as e:
        flash(f"Error cloning plan: {e}", "danger")
        return redirect(url_for('teacher.teacher_all_clps'))

@teacher_bp.route('/my_clps', methods=['GET', 'POST'])
@login_required
@roles_required('teacher')
def teacher_my_clps():
    upload_form = CLPUploadForm()

    if request.method == 'GET':
        user_profile = get_current_user_profile()
        if user_profile and user_profile.get('assigned_department'):
            upload_form.department.data = user_profile['assigned_department']

    try:
        # Load all plans for this user to allow client-side instant filtering
        # Adding simple retry logic for transient connection issues
        attempts = 0
        plans_res = None
        while attempts < 3:
            try:
                plans_res = (
                    supabase.table('course_learning_plans')
                    .select('id,subject,status,date_posted,user_id,department,filename,template_profile_id,upload_type,dean_comments')
                    .eq('user_id', session['user_id'])
                    .order('date_posted', desc=True)
                    .execute()
                )
                break
            except Exception as conn_e:
                attempts += 1
                if attempts >= 3: raise conn_e
                time.sleep(1) # Wait before retry

        plans = parse_supabase_timestamp(plans_res.data, 'date_posted')

        # Keep this list page light: full CLP content can be several MB per row.
        for plan in plans:
            plan['data'] = {}
            plan['course_code'] = ''
            plan['course_title'] = plan.get('subject') or ''
            plan['workflow_type'] = 'copilot_beta' if plan.get('upload_type') == 'ai_copilot_beta' else ''
            plan['progress'] = {
                'label': 'Working',
                'percent': 5,
            }
            plan['is_legacy_clp'] = not (
                plan.get('upload_type') == 'ai_copilot_beta'
            )

    except Exception as e:
        plans = []
        current_app.logger.error(f"Error fetching plans: {e}")

    return render_template(
        'teacher_courses.html',
        plans=plans,
        upload_form=upload_form,
    )

@teacher_bp.route('/courses/generate', methods=['POST'])
@login_required
@roles_required('teacher')
def generate_clp():
    form = CLPGenerateForm()
    if form.validate_on_submit():
        if not current_app.config.get('GEMINI_API_KEY'):
            flash('AI service not configured.', 'danger')
            return redirect(url_for('teacher.teacher_my_clps'))
        flash("Please use the full AI Generator page.", "info")
        return redirect(url_for('teacher.create_clp_ai_beta'))
    return redirect(url_for('teacher.teacher_my_clps'))

@teacher_bp.route('/clp/<int:plan_id>/edit', methods=['GET', 'POST'])
@login_required
@roles_required('teacher')
def edit_clp(plan_id):
    plan_res = supabase.table('course_learning_plans').select('*').eq('id', plan_id).single().execute()
    plan = plan_res.data

    if not plan: abort(404)
    if plan['user_id'] != session['user_id']: abort(403)

    # --- DRAFT REVIEW EDITOR ---
    if plan.get('upload_type') == 'ai_generated' and plan.get('status') == 'draft_review':
        content_data = _coerce_json_object(plan.get('content')) or {}

        form = FlaskForm() # For CSRF
        if request.method == 'POST' and form.validate_on_submit():
            action = request.form.get('action')

            # 1. Prepare "Updated" Data from Form
            updated_data = content_data.copy()
            # We iterate over the raw form to catch all dynamic fields (Weeks, Matrices)
            for key in request.form:
                if key not in ['csrf_token', 'submit', 'action']:
                    updated_data[key] = request.form.get(key)

            # 2. Capture "Original" Data (Snapshot before this update)
            original_data = content_data

            try:
                if action == 'save_draft':
                    # Pure manual save (no AI refinement)
                    supabase.table('course_learning_plans').update({
                        'content': json.dumps(updated_data)
                    }).eq('id', plan_id).execute()

                    # Save version
                    VersionService.save_version(plan_id, updated_data, "Manual Update (Draft)", session['user_id'])
                    flash('Draft saved successfully.', 'success')
                    return redirect(url_for('teacher.view_clp', plan_id=plan_id))
                else:
                    # AI Refinement path
                    updated_data['progress'] = {'step': 1, 'label': 'Initializing Refinement...', 'percent': 5}
                    supabase.table('course_learning_plans').update({
                        'status': 'generating',
                        'content': json.dumps(updated_data)
                    }).eq('id', plan_id).execute()

                    start_clp_refinement(plan_id, session['user_id'], original_data, updated_data)
                    flash('Updates saved! AI is comparing versions and refining...', 'info')
                    return redirect(url_for('teacher.view_clp', plan_id=plan_id))
            except Exception as e: flash(f"Error: {e}", "danger")

        outcomes_bundle = get_department_outcomes_bundle(plan.get('department'))

        return render_template('edit_ai_clp.html', plan=plan, content_data=content_data,
                               program_outcomes=outcomes_bundle['program_outcomes'],
                               course_outcomes=outcomes_bundle['course_outcomes'],
                               institutional_headers=outcomes_bundle['institutional_headers'], form=form)

    # --- STANDARD METADATA EDITOR ---
    form = CLPUpdateForm()
    if form.validate_on_submit():
        update_data = {'department': form.department.data, 'subject': form.subject.data}
        if form.file.data:
            new_filename = secure_filename(form.file.data.filename)
            file_path = f"{session['user_id']}/{datetime.now().timestamp()}_{new_filename}"
            file_bytes = form.file.data.read()
            write_storage_bytes(STORAGE_BUCKET_NAME, file_path, file_bytes)
            update_data.update({'filename': file_path, 'content': None, 'upload_type': 'file_upload'})
        elif form.content.data:
            update_data.update({'content': form.content.data, 'filename': None, 'upload_type': 'manual_text'})

        previous_filename = plan.get('filename')
        try:
            supabase.table('course_learning_plans').update(update_data).eq('id', plan_id).execute()
            cleanup_ok = True
            if form.file.data and previous_filename and previous_filename != update_data.get('filename'):
                cleanup_ok = cleanup_storage_after_commit(
                    previous_filename,
                    context='Teacher plan update',
                    user_id=session.get('user_id'),
                    plan_id=plan_id,
                )
            feedback = make_workflow_feedback(
                'Plan update',
                detail='Plan updated successfully.' if cleanup_ok else 'Plan updated, but the previous file could not be removed from storage.',
                partial=not cleanup_ok,
            )
            _flash_workflow_feedback(feedback)
            return redirect(url_for('teacher.view_clp', plan_id=plan_id))
        except Exception as e:
            if form.file.data:
                cleanup_storage_after_commit(
                    update_data.get('filename'),
                    context='Teacher failed plan update',
                    user_id=session.get('user_id'),
                    plan_id=plan_id,
                )
            _flash_workflow_feedback(make_workflow_feedback('Plan update', e))
            return redirect(url_for('teacher.view_clp', plan_id=plan_id))

    elif request.method == 'GET':
        form.department.data = plan['department']
        form.subject.data = plan['subject']
        if plan['upload_type'] in ['manual_text']:
            form.content.data = plan.get('content') or ''

    return render_template('edit_clp.html', title='Edit Plan', form=form, plan=plan)

@teacher_bp.route('/clp/<int:plan_id>/edit_document')
@login_required
@roles_required('teacher')
def edit_clp_document(plan_id):
    started_at = time.perf_counter()
    try:
        plan_res = supabase.table('course_learning_plans').select('id, user_id, filename, upload_type, content, department, subject').eq('id', plan_id).single().execute()
        plan = plan_res.data
    except httpx.HTTPError as e:
        return _handle_supabase_route_failure(e, 'teacher.view_clp', plan_id=plan_id)

    if not plan: abort(404)
    if plan['user_id'] != session['user_id']: abort(403)

    if not plan['filename'] or not plan['filename'].lower().endswith('.docx'):
        flash('This plan is not a .docx document.', 'warning')
        return redirect(url_for('teacher.view_clp', plan_id=plan_id))

    try:
        import hashlib
        key_string = f"clp_{plan['id']}_{plan['filename']}"
        doc_key = hashlib.md5(key_string.encode()).hexdigest()

        base_url = get_onlyoffice_base_url(internal=True)

        doc_url = f"{base_url}/teacher/serve_clp/{plan_id}/{doc_key}"
        callback_url = f"{base_url}/teacher/onlyoffice_clp_callback/{plan_id}/{doc_key}"
        client_log_url = url_for('teacher.onlyoffice_client_log', plan_id=plan_id)
        public_base_url = get_onlyoffice_base_url(internal=False)

        config = {
            "document": {
                "title": os.path.basename(plan['filename']).split('_', 1)[-1],
                "url": doc_url,
                "fileType": "docx",
                "key": doc_key,
                "permissions": {"edit": True, "download": True, "review": True}
            },
            "documentType": "word",
            "editorConfig": {
                "mode": "edit", "callbackUrl": callback_url,
                "user": {"id": str(session['user_id']), "name": session.get('username', 'Teacher')},
                "customization": {
                    "autosave": True,
                    "forcesave": True,
                    "hideRightMenu": False
                }
            },
            "width": "100%", "height": "100%"
        }

        attach_onlyoffice_ai_plugin(config)
        _attach_lpms_onlyoffice_rewrite_plugin(config, plan_id, doc_key, session['user_id'])
        onlyoffice_ai_rewrite_url = url_for('teacher.onlyoffice_ai_rewrite_selection', plan_id=plan_id)

        # Generate JWT Token for ONLYOFFICE after all editor config customizations are in place.
        token = generate_jwt_token(config) or ""
        _log_teacher_onlyoffice(
            logging.INFO,
            'edit_document_config_built',
            plan_id=plan_id,
            session_user_id=session.get('user_id'),
            filename=plan.get('filename'),
            doc_key=doc_key,
            doc_url=doc_url,
            callback_url=callback_url,
            client_log_url=client_log_url,
            token_present=bool(token),
            token_length=len(token),
            base_url=base_url,
            remote_addr=request.headers.get('X-Forwarded-For', request.remote_addr),
        )
        log_document_timing('teacher_edit_document_config', started_at, plan_id=plan_id, user_id=session.get('user_id'))

        return render_template("teacher_edit_document.html",
                               config=config,
                               doc_title=config["document"]["title"],
                               doc_url=doc_url,
                               callback_url=callback_url,
                               doc_key=doc_key,
                               token=token,
                               plan_id=plan_id,
                               onlyoffice_client_log_url=client_log_url,
                               onlyoffice_ai_rewrite_url=onlyoffice_ai_rewrite_url,

                               inserted_success=request.args.get('inserted') == '1',
                               inserted_source=request.args.get('source', ''),
                               inserted_flagged=request.args.get('flagged') == '1')

    except HTTPException:
        raise
    except Exception as e:
        _log_teacher_onlyoffice(logging.ERROR, 'edit_document_init_failed', plan_id=plan_id, error=e, traceback=traceback.format_exc())
        flash(f"Error loading editor: {str(e)}", "danger")
        return redirect(url_for('teacher.view_clp', plan_id=plan_id))


@teacher_bp.route('/serve_clp/<int:plan_id>/<doc_key>')
def serve_clp_document(plan_id, doc_key):
    started_at = time.perf_counter()
    try:
        _log_teacher_onlyoffice(
            logging.INFO,
            'serve_document_requested',
            plan_id=plan_id,
            doc_key=doc_key,
            remote_addr=request.headers.get('X-Forwarded-For', request.remote_addr),
            user_agent=request.headers.get('User-Agent'),
            authorization_present=bool(request.headers.get('Authorization')),
        )
        plan_res = supabase.table('course_learning_plans').select('filename, user_id, subject, department, content, upload_type').eq('id', plan_id).single().execute()
        plan = plan_res.data
        if not plan or not plan.get('filename'): abort(404)
        expected_edit_key = hashlib.md5(f"clp_{plan_id}_{plan['filename']}".encode()).hexdigest()
        expected_view_key = hashlib.md5(f"view_clp_{plan_id}_{plan['filename']}".encode()).hexdigest()
        if doc_key not in {expected_edit_key, expected_view_key}:
            _log_teacher_onlyoffice(logging.WARNING, 'serve_document_key_mismatch', plan_id=plan_id, provided_doc_key=doc_key, expected_doc_key=expected_edit_key, expected_view_key=expected_view_key, filename=plan.get('filename'))
            abort(403)
        recovered = False
        if not storage_file_exists(STORAGE_BUCKET_NAME, plan['filename']):
            recovered_filename = ensure_plan_document_file({**plan, 'id': plan_id})
            plan['filename'] = recovered_filename
            recovered = True
            if doc_key == expected_edit_key:
                expected_edit_key = hashlib.md5(f"clp_{plan_id}_{plan['filename']}".encode()).hexdigest()
                if doc_key != expected_edit_key:
                    _log_teacher_onlyoffice(logging.WARNING, 'serve_document_key_changed_after_recovery', plan_id=plan_id, provided_doc_key=doc_key, recovered_doc_key=expected_edit_key, filename=plan.get('filename'))
                    abort(409)
        download_filename = os.path.basename(plan['filename']).split('_', 1)[-1]
        response = stream_storage_file(
            STORAGE_BUCKET_NAME,
            plan['filename'],
            download_filename,
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            inline=True,
            cache_seconds=300,
        )
        _log_teacher_onlyoffice(logging.INFO, 'serve_document_success', plan_id=plan_id, filename=plan.get('filename'), download_filename=download_filename, recovered=recovered)
        log_document_timing('teacher_serve_clp', started_at, plan_id=plan_id, recovered=recovered, filename=plan.get('filename'))
        return response
    except FileNotFoundError:
        _log_teacher_onlyoffice(logging.WARNING, 'serve_document_missing_file', plan_id=plan_id, doc_key=doc_key)
        abort(404)
    except Exception as exc:
        _log_teacher_onlyoffice(logging.ERROR, 'serve_document_failed', plan_id=plan_id, doc_key=doc_key, error=exc, traceback=traceback.format_exc())
        abort(500)

@teacher_bp.route('/onlyoffice_clp_callback/<int:plan_id>/<doc_key>', methods=['POST'])
@csrf.exempt
def onlyoffice_clp_callback(plan_id, doc_key):
    try:
        data = request.get_json(silent=True) or {}
        _log_teacher_onlyoffice(
            logging.INFO,
            'callback_received',
            plan_id=plan_id,
            doc_key=doc_key,
            status=data.get('status'),
            callback_keys=sorted(list(data.keys())),
            has_download_url=bool(data.get('url')),
            remote_addr=request.headers.get('X-Forwarded-For', request.remote_addr),
            user_agent=request.headers.get('User-Agent'),
            authorization_present=bool(request.headers.get('Authorization')),
        )
        if not _validate_onlyoffice_callback(data, expected_key=doc_key):
            _log_teacher_onlyoffice(logging.WARNING, 'callback_rejected', plan_id=plan_id, doc_key=doc_key, status=data.get('status'))
            return jsonify({"error": 1})
        if data.get("status") in [2, 6]:
            download_url = data.get("url")
            if not download_url:
                _log_teacher_onlyoffice(logging.WARNING, 'callback_missing_download_url', plan_id=plan_id, doc_key=doc_key, status=data.get('status'))
                return jsonify({"error": 1})

            res = supabase.table('course_learning_plans').select('filename, user_id').eq('id', plan_id).single().execute()
            if not res.data:
                _log_teacher_onlyoffice(logging.WARNING, 'callback_plan_missing', plan_id=plan_id, doc_key=doc_key)
                return jsonify({"error": 1})

            # Download file from ONLYOFFICE
            resp = requests.get(download_url, timeout=60)
            resp.raise_for_status()
            _log_teacher_onlyoffice(logging.INFO, 'callback_download_success', plan_id=plan_id, doc_key=doc_key, status=data.get('status'), response_status=resp.status_code, download_url=download_url)

            # --- DATE REPLACEMENT LOGIC ---
            # Process the returned DOCX to update the date_today placeholder
            try:
                processed_content = resp.content # Default to original if processing fails
                doc = Document(BytesIO(resp.content))

                # Create replacement dict with current month and year
                replacements = {'date_today': datetime.now().strftime("%B %Y")}

                # Use shared utility to replace text in all paragraphs/tables/headers/footers
                doc = replace_placeholders(doc, replacements)

                # Save modified doc to bytes
                output = BytesIO()
                doc.save(output)
                output.seek(0)
                processed_content = output.read()

            except Exception as e:
                _log_teacher_onlyoffice(logging.ERROR, 'callback_date_replace_failed', plan_id=plan_id, doc_key=doc_key, error=e, traceback=traceback.format_exc())
                # Fallback to original content if processing fails

            # Upload processed content (or original on error)
            clean_name = os.path.basename(res.data['filename']).split('_', 1)[-1] or "document.docx"
            new_filename = f"{res.data['user_id']}/{int(time.time())}_{clean_name}"

            write_storage_bytes(STORAGE_BUCKET_NAME, new_filename, processed_content)

            supabase.table('course_learning_plans').update({'filename': new_filename}).eq('id', plan_id).execute()
            cleanup_storage_after_commit(
                res.data['filename'],
                context='Teacher OnlyOffice save',
                user_id=res.data.get('user_id'),
                plan_id=plan_id,
            )
            _log_teacher_onlyoffice(logging.INFO, 'callback_save_success', plan_id=plan_id, doc_key=doc_key, old_filename=res.data.get('filename'), new_filename=new_filename, byte_count=len(processed_content))

            return jsonify({"error": 0})
        _log_teacher_onlyoffice(logging.INFO, 'callback_ignored_status', plan_id=plan_id, doc_key=doc_key, status=data.get('status'))
        return jsonify({"error": 0})
    except Exception as exc:
        _log_teacher_onlyoffice(logging.ERROR, 'callback_failed', plan_id=plan_id, doc_key=doc_key, error=exc, traceback=traceback.format_exc())
        return jsonify({"error": 1})

@teacher_bp.route('/clp/<int:plan_id>')
@login_required
@roles_required('teacher', 'dean', 'admin')
def view_clp(plan_id):
    try:
        plan_res = supabase.table('course_learning_plans').select('*, author:users(*)').eq('id', plan_id).single().execute()
        plan = plan_res.data
    except httpx.HTTPError as e:
        return _handle_supabase_route_failure(e, 'teacher.teacher_my_clps')
    if not plan: abort(404)
    parsed_plan_list = parse_supabase_timestamp([plan], 'date_posted')
    plan = parsed_plan_list[0]

    if not user_can_access_clp(plan):
        abort(403)

    if plan['upload_type'] == 'file_upload' and plan['status'] != 'draft_review':
        return render_template('view_uploaded_clp.html', plan=plan)

    try:
        content_data = _coerce_json_object(plan.get('content'))
        if content_data is None:
            return render_template('view_clp.html', plan=plan)

        outcomes_bundle = get_department_outcomes_bundle(plan.get('department'))

        return render_template('view_ai_clp.html', plan=plan, content_data=content_data,
                               program_outcomes=outcomes_bundle['program_outcomes'],
                               course_outcomes=outcomes_bundle['course_outcomes'],
                               institutional_headers=outcomes_bundle['institutional_headers'],
                               program_headers=outcomes_bundle['program_headers'],
                               form=FlaskForm())
    except (json.JSONDecodeError, TypeError):
        return render_template('view_clp.html', plan=plan)

@teacher_bp.route('/clp/<int:plan_id>/download')
@login_required
def download_clp(plan_id):
    plan_res = supabase.table('course_learning_plans').select('filename, user_id, department, status, subject, content, upload_type').eq('id', plan_id).single().execute()
    plan = plan_res.data
    if not plan or not plan.get('filename'):
        flash('File not ready.', 'danger')
        return redirect(url_for('teacher.view_clp', plan_id=plan_id))
    if not user_can_access_clp(plan):
        flash('Unauthorized access to this file.', 'danger')
        return redirect(url_for('dean.dean_courses' if session.get('role') == 'dean' else 'teacher.teacher_my_clps'))
    try:
        try:
            file_bytes = read_storage_bytes(STORAGE_BUCKET_NAME, plan['filename'])
        except FileNotFoundError:
            recovered_filename = ensure_plan_document_file({**plan, 'id': plan_id})
            plan['filename'] = recovered_filename
            file_bytes = read_storage_bytes(STORAGE_BUCKET_NAME, plan['filename'])
        download_name = os.path.basename(plan['filename']).split('_', 1)[-1]
        return Response(file_bytes, mimetype='application/octet-stream', headers={"Content-disposition": f"attachment; filename=\"{download_name}\""})
    except Exception as e:
        current_app.logger.warning(f"Error downloading plan {plan_id}: {e}")
        flash('The file could not be downloaded right now.', 'danger')
        return redirect(url_for('teacher.view_clp', plan_id=plan_id))

@teacher_bp.route('/check_generation_status')
@login_required
@limiter.exempt
def check_generation_status():
    try:
        res = supabase.table('course_learning_plans').select('id, subject, content').eq('user_id', session['user_id']).in_('status', ['generating', 'generating_doc']).execute()
        generating = len(res.data) > 0
        return jsonify({'generating': generating, 'plans': res.data})
    except Exception as e:
        current_app.logger.warning(f"Generation status polling failed for user {session.get('user_id')}: {e}")
        return jsonify({'generating': False, 'error': 'Unable to refresh generation status right now.'})


@teacher_bp.route('/courses/upload', methods=['POST'])
@login_required
@roles_required('teacher')
def upload_clp():
    form = CLPUploadForm()
    if form.validate_on_submit():
        user_id = session['user_id']
        file = form.file.data
        if file and allowed_file(file.filename):
            filename = secure_filename(file.filename)
            file_path = f"{user_id}/{datetime.now().timestamp()}_{filename}"
            supabase.storage.from_(STORAGE_BUCKET_NAME).upload(file=file.read(), path=file_path, file_options={"content-type": file.mimetype})
            supabase.table('course_learning_plans').insert({'department': form.department.data, 'subject': form.subject.data, 'filename': file_path, 'upload_type': 'file_upload', 'status': 'draft', 'user_id': user_id}).execute()
            flash('CLP uploaded as draft!', 'success')
    return redirect(url_for('teacher.teacher_my_clps'))

@teacher_bp.route('/submit_to_dean/<int:plan_id>', methods=['POST'])
@login_required
def submit_to_dean(plan_id):
    try:
        plan_res = supabase.table('course_learning_plans').select('user_id, status, upload_type, content, filename, department').eq('id', plan_id).single().execute()
        plan = plan_res.data
        if not plan or (plan.get('user_id') != session.get('user_id') and session.get('role') != 'admin'):
            flash('Unauthorized action.', 'danger')
            return redirect(url_for('teacher.teacher_my_clps'))

        # Check if department has a dean
        dept = plan.get('department', '')
        if dept:
            dept_signatories = get_department_signatory_settings(department_name=dept) or {}
            dean_name = dept_signatories.get('dean_name', '')
            if not dean_name:
                flash(f'No dean assigned for {dept}. Please notify your dean to create an account.', 'warning')
                return redirect(url_for('teacher.teacher_my_clps'))

        allowed_statuses = {'draft', 'draft_review', 'returned_for_revision', 'beta_ready'}
        can_submit = plan.get('status') in allowed_statuses
        if not can_submit and plan.get('upload_type') == 'ai_copilot_beta':
            plan_content = _coerce_json_object(plan.get('content')) or {}
            review_stage = str(plan_content.get('review_stage') or '').strip().lower()
            beta_ready_for_template = bool(plan_content.get('beta_ready_for_template'))
            has_beta_document = bool(plan.get('filename')) and str(plan.get('filename')).lower().endswith('.docx')
            can_submit = review_stage in {'beta_ready', 'beta_inserted'} or (beta_ready_for_template and has_beta_document)
        if not can_submit:
            flash('This plan cannot be submitted right now.', 'warning')
            return redirect(url_for('teacher.teacher_my_clps'))
        supabase.table('course_learning_plans').update({'status': 'pending'}).eq('id', plan_id).execute()
        flash('Submitted to Dean.', 'success')
        create_notification(session['user_id'], 'CLP submitted to Dean.', reference_type='clp', reference_id=plan_id)
    except Exception as e:
        current_app.logger.warning(f"Failed to submit plan {plan_id} to dean: {e}")
        _flash_workflow_feedback(make_workflow_feedback('Submit to dean', e))
    return redirect(url_for('teacher.teacher_my_clps'))

@teacher_bp.route('/clp/<int:plan_id>/delete', methods=['POST'])
@login_required
def delete_clp(plan_id):
    try:
        plan = supabase.table('course_learning_plans').select('filename, user_id, department, status').eq('id', plan_id).single().execute().data
        if not plan or not user_can_delete_clp(plan):
            flash('Unauthorized action.', 'danger')
            return redirect(url_for('teacher.teacher_my_clps'))
        deleted = delete_clp_with_dependencies(plan_id)
        if not deleted:
            raise ValueError('The plan could not be deleted.')
        cleanup_ok = cleanup_storage_after_commit(
            plan.get('filename'),
            context='Teacher plan delete',
            user_id=session.get('user_id'),
            plan_id=plan_id,
        )
        _flash_workflow_feedback(
            make_workflow_feedback(
                'Plan delete',
                detail='Deleted.' if cleanup_ok else 'Plan deleted, but the file could not be removed from storage.',
                partial=not cleanup_ok,
            )
        )
    except Exception as e:
        current_app.logger.warning(f"Failed to delete plan {plan_id}: {e}")
        _flash_workflow_feedback(make_workflow_feedback('Plan delete', e))
    return redirect(url_for('teacher.teacher_my_clps'))

# ... existing imports ...

@teacher_bp.route('/teacher_profile', methods=['GET', 'POST'])
@login_required
@roles_required('teacher')
def teacher_profile():
    user = get_current_user_profile()
    if not user:
        flash('We could not load your profile. Please sign in again.', 'warning')
        session.clear()
        return redirect(url_for('auth.login'))
    info_form = UserProfileForm(obj=user)
    pwd_form = ChangePasswordForm()

    # --- GET: Populate Form from JSON ---
    if request.method == 'GET':
        info_form.first_name.data = user.get('first_name')
        info_form.last_name.data = user.get('last_name')
        info_form.title.data = user.get('title')

        # Load Consultation Hours
        cons = user.get('consultation_hours') or {}
        info_form.cons_mon_time.data = cons.get('monday', {}).get('time', '')
        info_form.cons_mon_room.data = cons.get('monday', {}).get('room', '')
        info_form.cons_tue_time.data = cons.get('tuesday', {}).get('time', '')
        info_form.cons_tue_room.data = cons.get('tuesday', {}).get('room', '')
        info_form.cons_wed_time.data = cons.get('wednesday', {}).get('time', '')
        info_form.cons_wed_room.data = cons.get('wednesday', {}).get('room', '')
        info_form.cons_thu_time.data = cons.get('thursday', {}).get('time', '')
        info_form.cons_thu_room.data = cons.get('thursday', {}).get('room', '')
        info_form.cons_fri_time.data = cons.get('friday', {}).get('time', '')
        info_form.cons_fri_room.data = cons.get('friday', {}).get('room', '')

    # --- POST: Save Form to JSON ---
    if info_form.validate_on_submit() and info_form.submit_info.data:

        # Construct the JSON object
        consultation_data = {
            'monday': {'time': info_form.cons_mon_time.data, 'room': info_form.cons_mon_room.data},
            'tuesday': {'time': info_form.cons_tue_time.data, 'room': info_form.cons_tue_room.data},
            'wednesday': {'time': info_form.cons_wed_time.data, 'room': info_form.cons_wed_room.data},
            'thursday': {'time': info_form.cons_thu_time.data, 'room': info_form.cons_thu_room.data},
            'friday': {'time': info_form.cons_fri_time.data, 'room': info_form.cons_fri_room.data}
        }

        update_data = {
            'first_name': info_form.first_name.data,
            'last_name': info_form.last_name.data,
            'title': info_form.title.data,
            'consultation_hours': consultation_data # Save the JSON
        }

        # ... (Existing Signature Upload Logic remains here) ...
        if info_form.signature.data:
            file = info_form.signature.data
            filename = secure_filename(file.filename)
            file_path = f"signatures/{session['user_id']}_{int(time.time())}_{filename}"
            try:
                supabase.storage.from_(STORAGE_BUCKET_NAME).upload(
                    path=file_path, file=file.read(),
                    file_options={"content-type": file.mimetype}
                )
                update_data['signature_url'] = build_public_storage_url(STORAGE_BUCKET_NAME, file_path)
            except Exception as e:
                flash(f"Error uploading signature: {e}", "danger")

        if info_form.profile_photo.data:
            file = info_form.profile_photo.data
            filename = secure_filename(file.filename)
            file_path = f"profiles/{session['user_id']}_{int(time.time())}_{filename}"
            try:
                supabase.storage.from_(STORAGE_BUCKET_NAME).upload(
                    path=file_path, file=file.read(),
                    file_options={"content-type": file.mimetype}
                )
                update_data['profile_photo_url'] = build_public_storage_url(STORAGE_BUCKET_NAME, file_path)
            except Exception as e:
                flash(f"Error uploading profile photo: {e}", "danger")

        try:
            supabase.table('users').update(update_data).eq('id', session['user_id']).execute()
            invalidate_user_profile_cache(session['user_id'])
            session['name'] = f"{update_data.get('first_name', session.get('name', ''))} {update_data.get('last_name', '')}"
            if 'profile_photo_url' in update_data:
                session['profile_photo_url'] = update_data['profile_photo_url']
            flash('Profile updated successfully.', 'success')
            next_url = request.args.get('next')
            if next_url:
                return redirect(next_url)
            return redirect(url_for('teacher.teacher_profile'))
        except Exception as e:
            flash(f"Error updating profile: {e}", 'danger')

    # ... (Password form logic remains here) ...
    if pwd_form.validate_on_submit() and pwd_form.submit.data:
        try:
            password_row = (
                supabase.table('users')
                .select('password_hash')
                .eq('id', session['user_id'])
                .single()
                .execute()
            )
            current_password_hash = (password_row.data or {}).get('password_hash')
            if not verify_password_hash(current_password_hash, pwd_form.current_password.data):
                flash('Current password is incorrect.', 'danger')
                return render_template('teacher_profile.html', info_form=info_form, pwd_form=pwd_form, user=user)
            supabase.auth.update_user({"password": pwd_form.new_password.data})
            flash('Password changed successfully.', 'success')
            return redirect(url_for('teacher.teacher_profile'))
        except Exception as e:
            flash(f"Error changing password: {e}", 'danger')

    return render_template('teacher_profile.html', info_form=info_form, pwd_form=pwd_form, user=user)


@teacher_bp.route('/all_clps')
@login_required
def teacher_all_clps():
    user = get_current_user_profile()
    teacher_department = user.get('assigned_department')

    if not teacher_department:
        flash('No department is assigned to your account yet.', 'warning')
        return render_template('all_courses.html', plans=[], department_name=None)

    plans_res = (
        supabase.table('course_learning_plans')
        .select('id,subject,status,date_posted,user_id,department,filename,upload_type,template_profile_id,author:users(username)')
        .eq('status', 'approved')
        .eq('department', teacher_department)
        .order('date_posted', desc=True)
        .execute()
    )
    plans = parse_supabase_timestamp(plans_res.data, 'date_posted')

    # Annotate each plan so the Clone button can show a tailored confirmation
    # ("new AI Copilot data" vs "legacy — only metadata will be copied").
    for plan in plans:
        plan['is_legacy_clp'] = not (
            plan.get('upload_type') == 'ai_copilot_beta'
        )

    return render_template('all_courses.html', plans=plans, department_name=teacher_department)

@teacher_bp.route('/clp/<int:plan_id>/delete_approved', methods=['POST'])
@login_required
@roles_required('teacher', 'dean')
def delete_approved_clp(plan_id):
    plan_res = supabase.table('course_learning_plans').select('user_id, filename, subject, department, status').eq('id', plan_id).single().execute()
    plan = plan_res.data

    if not plan: abort(404)

    # Security check
    if not user_can_delete_clp(plan):
        flash('You do not have permission to delete this plan.', 'danger')
        return redirect(url_for('dean.dean_courses' if session.get('role') == 'dean' else 'teacher.teacher_all_clps'))

    try:
        deleted = delete_clp_with_dependencies(plan_id)
        if not deleted:
            raise ValueError('The approved plan could not be deleted.')
        cleanup_ok = cleanup_storage_after_commit(
            plan.get('filename'),
            context='Approved plan delete',
            user_id=session.get('user_id'),
            plan_id=plan_id,
        )
        _flash_workflow_feedback(
            make_workflow_feedback(
                'Approved plan delete',
                detail=f"Approved CLP '{plan['subject']}' deleted." if cleanup_ok else f"Approved CLP '{plan['subject']}' deleted, but the file cleanup failed.",
                partial=not cleanup_ok,
            )
        )
    except Exception as e:
        current_app.logger.warning(f"Failed to delete approved plan {plan_id}: {e}")
        _flash_workflow_feedback(make_workflow_feedback('Approved plan delete', e))

    if session['role'] == 'dean':
        return redirect(url_for('dean.dean_courses'))
    else:
        return redirect(url_for('teacher.teacher_my_clps'))

@teacher_bp.route('/clp/<int:plan_id>/view_document')
@login_required
def view_approved_clp_document(plan_id):
    """Open the CLP in ONLYOFFICE editor for viewing (Read-Only)."""
    started_at = time.perf_counter()
    try:
        # Fetch plan details
        res = supabase.table('course_learning_plans').select('id, user_id, status, department, subject, filename').eq('id', plan_id).single().execute()
        plan = res.data

        if not plan: abort(404)

        if not user_can_access_clp(plan):
            abort(403)

        # Ensure it is a valid DOCX file
        if not plan.get('filename') or not plan['filename'].lower().endswith('.docx'):
            flash('This plan is not a .docx document.', 'warning')
            return redirect(url_for('teacher.teacher_all_clps'))

        # Generate unique key for viewing session
        key_string = f"view_clp_{plan['id']}_{plan['filename']}"
        doc_key = hashlib.md5(key_string.encode()).hexdigest()
        doc_title = f"{plan['subject']} - Read Only"

        base_url = get_onlyoffice_base_url(internal=True)

        # Use the internal serve_clp route for the document URL
        # This is more reliable than direct Supabase URLs in local/docker environments
        doc_url = f"{base_url}/teacher/serve_clp/{plan_id}/{doc_key}"
        callback_url = f"{base_url}/teacher/onlyoffice_clp_callback/{plan_id}/{doc_key}"

        # Setup Configuration
        config = {
            "document": {
                "title": doc_title,
                "url": doc_url,
                "fileType": "docx",
                "key": doc_key,
                "permissions": {
                    "edit": False,
                    "download": True,
                    "review": False
                }
            },
            "documentType": "word",
            "editorConfig": {
                "mode": "view",
                "callbackUrl": callback_url,
                "user": {
                    "id": str(session['user_id']),
                    "name": session.get('username', 'Teacher')
                },
                "customization": {
                    "autosave": False,
                    "forcesave": False
                }
            },
            "width": "100%", "height": "100%"
        }

        # Generate JWT Token for ONLYOFFICE
        token = generate_jwt_token(config) or ""
        log_document_timing('teacher_view_document_config', started_at, plan_id=plan_id, user_id=session.get('user_id'))

        return render_template("teacher_view_document.html",
                               config=config,
                               doc_title=doc_title,
                               doc_url=doc_url,
                               callback_url=callback_url,
                               doc_key=doc_key,
                               token=token)

    except HTTPException:
        raise
    except Exception as e:
        current_app.logger.error(f"Error opening Viewer: {e}")
        flash(f"Error loading document viewer: {str(e)}", "danger")
        return redirect(url_for('teacher.teacher_all_clps'))
