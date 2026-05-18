import os
import re
from io import BytesIO
from zipfile import ZipFile
from docx import Document
from docx.oxml.ns import qn

def get_template_filepath():
    path = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'clp_templates', 'nursing.docx')
    if not os.path.exists(path):
        path = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'clp_templates', 'PBSIT-001-LP-20242 copy.docx')
    return path

def _apply_run_font(run, font_name):
    run.font.name = font_name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn('w:ascii'), font_name)
    r_fonts.set(qn('w:hAnsi'), font_name)
    r_fonts.set(qn('w:eastAsia'), font_name)
    r_fonts.set(qn('w:cs'), font_name)


def _apply_style_font(style, font_name):
    if not style:
        return
    style.font.name = font_name
    style_element = getattr(style, "_element", None)
    if style_element is None:
        return
    r_pr = style_element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn('w:ascii'), font_name)
    r_fonts.set(qn('w:hAnsi'), font_name)
    r_fonts.set(qn('w:eastAsia'), font_name)
    r_fonts.set(qn('w:cs'), font_name)

def replace_placeholders(doc, replacements, font_name='Arial Narrow'):
    """
    Replaces placeholders in a docx file, preserving the original formatting
    of the placeholder run (font, size, bold, italic, color, etc.).

    Falls back to *font_name* only when the placeholder paragraph has no
    existing run formatting to clone.
    """
    # 1. Prepare replacements by including braced versions if they don't exist
    final_replacements = {}
    for k, v in replacements.items():
        val = str(v if v is not None else "").replace('\\n', '\n')
        final_replacements[k] = val
        if not (k.startswith('{{') and k.endswith('}}')):
            final_replacements[f"{{{{{k}}}}}"] = val
    sorted_keys = sorted(final_replacements.keys(), key=len, reverse=True)

    # Import the robust cloning helpers from template_writer
    from app.services.template_writer import _clone_run_props, _apply_run_props

    def process_paragraph(paragraph):
        original_text = paragraph.text
        if not original_text.strip():
            return

        new_text = original_text
        found = False
        for key in sorted_keys:
            if key in new_text:
                new_text = new_text.replace(key, final_replacements[key])
                found = True

        if found and new_text != original_text:
            # Strategy: do run-level text replacement when possible.
            replaced_in_runs = False
            for run in paragraph.runs:
                run_orig = run.text
                run_new = run_orig
                for key in sorted_keys:
                    if key in run_new:
                        run_new = run_new.replace(key, final_replacements[key])
                        found = True
                if run_new != run_orig:
                    run.text = run_new
                    replaced_in_runs = True

            if not replaced_in_runs:
                # Fallback: replacement spans multiple runs.
                # Save paragraph properties (tab stops, alignment) before clearing.
                from lxml import etree
                
                pPr = paragraph._element.find(qn('w:pPr'))
                saved_pPr = etree.tostring(pPr) if pPr is not None else None
                
                # Capture run formatting from the first non-empty run
                donor_rPr = None
                for r in paragraph.runs:
                    if r.text.strip():
                        donor_rPr = _clone_run_props(r)
                        break
                if donor_rPr is None and paragraph.runs:
                    donor_rPr = _clone_run_props(paragraph.runs[0])
                if donor_rPr is None:
                    rPr = etree.SubElement(etree.Element(qn('w:rPr')), qn('w:rFonts'))
                    rPr.set(qn('w:ascii'), font_name)
                    rPr.set(qn('w:hAnsi'), font_name)
                    rPr.set(qn('w:eastAsia'), font_name)
                    rPr.set(qn('w:cs'), font_name)
                    donor_rPr = rPr
                
                paragraph.clear()
                run = paragraph.add_run(new_text)
                _apply_run_props(run, donor_rPr)
                
                # Restore paragraph properties (tab stops, alignment, indentation)
                if saved_pPr is not None:
                    new_pPr = paragraph._element.find(qn('w:pPr'))
                    if new_pPr is not None:
                        paragraph._element.remove(new_pPr)
                    restored = etree.fromstring(saved_pPr)
                    paragraph._element.insert(0, restored)

            # Strip list-numbering so bullet styles do not conflict.
            pPr = paragraph._element.find(qn('w:pPr'))
            if pPr is not None:
                numPr = pPr.find(qn('w:numPr'))
                if numPr is not None:
                    pPr.remove(numPr)

    def process_tables(tables):
        for table in tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs: 
                        process_paragraph(para)
                    if cell.tables:
                        process_tables(cell.tables)

    # 2. Process Body
    for para in doc.paragraphs: 
        process_paragraph(para)
    process_tables(doc.tables)

    # 3. Process Sections (Headers/Footers)
    for section in doc.sections:
        for para in section.header.paragraphs: process_paragraph(para)
        process_tables(section.header.tables)
        for para in section.footer.paragraphs: process_paragraph(para)
        process_tables(section.footer.tables)
    return doc


def _iter_docx_xml_text(file_bytes):
    with ZipFile(BytesIO(file_bytes)) as archive:
        for name in archive.namelist():
            if not name.endswith(".xml"):
                continue
            xml_text = archive.read(name).decode("utf-8", "ignore")
            # Strip XML tags so placeholders split across runs still become visible.
            visible_text = re.sub(r"<[^>]+>", "", xml_text)
            if visible_text:
                yield visible_text


def _placeholder_identifier_variants(value):
    value = str(value or "").strip()
    if not value:
        return set()
    if value.startswith("{{") and value.endswith("}}"):
        bare = value[2:-2].strip()
        return {value, bare} if bare else {value}
    return {value, f"{{{{{value}}}}}"}


def extract_docx_placeholders(file_bytes):
    pattern = re.compile(r"\{\{[a-z0-9_]+\}\}", re.IGNORECASE)
    placeholders = set()
    for visible_text in _iter_docx_xml_text(file_bytes):
        placeholders.update(match.strip() for match in pattern.findall(visible_text))
    return placeholders


def has_docx_placeholders(file_bytes, required_placeholders=None):
    placeholders = extract_docx_placeholders(file_bytes)
    required = set(required_placeholders or [])
    if not placeholders and not required:
        return False, placeholders
    if required:
        visible_text = "\n".join(_iter_docx_xml_text(file_bytes))
        missing = []
        normalized_found = set(placeholders)
        for required_placeholder in required:
            variants = _placeholder_identifier_variants(required_placeholder)
            if any(variant in placeholders for variant in variants if variant.startswith("{{")):
                continue
            bare_variants = [variant for variant in variants if not variant.startswith("{{")]
            matched_bare = None
            for bare in bare_variants:
                if re.search(rf"(?<![A-Za-z0-9_]){re.escape(bare)}(?![A-Za-z0-9_])", visible_text):
                    matched_bare = bare
                    break
            if matched_bare:
                normalized_found.add(required_placeholder if required_placeholder.startswith("{{") else f"{{{{{matched_bare}}}}}")
                continue
            missing.append(required_placeholder)
        if missing:
            return False, normalized_found
        return True, normalized_found
    if not placeholders:
        return False, placeholders
    return True, placeholders

def flatten_json(data):
    out = {}
    def flatten(x, name=''):
        if isinstance(x, dict):
            for a in x: flatten(x[a], name + a)
        elif isinstance(x, list): out[name] = "\n".join(map(str, x))
        else: out[name] = x
    flatten(data)
    return out
