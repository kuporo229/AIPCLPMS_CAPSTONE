import os
import json
import time
import traceback
from io import BytesIO
from datetime import datetime
from flask import current_app
from app.compat_supabase import create_client
from docx import Document

from app.services.ai_client import AIClient
from app.services.docx_service import replace_placeholders, flatten_json, get_template_filepath
# Import utilities that remain in utils.py
from app.utils import (
    get_system_prompt, update_progress, handle_system_error, 
    create_notification, shred_clp_mappings, get_department_outcomes_bundle,
    get_department_id_by_name
)
from app import (
    STORAGE_BUCKET_NAME
)

from app.services.observability import StructuredLogger, timed

def _stringify_change_value(value, limit=160):
    if isinstance(value, (dict, list)):
        try:
            text = json.dumps(value, ensure_ascii=False)
        except Exception:
            text = str(value)
    else:
        text = str(value)
    if len(text) > limit:
        return f"{text[:limit]}…"
    return text

def summarize_clp_changes(original_data, updated_data, max_items=12, max_depth=3):
    changes = []

    def walk(original, updated, path, depth):
        if len(changes) >= max_items or depth > max_depth:
            return
        if isinstance(original, dict) and isinstance(updated, dict):
            keys = set(original.keys()) | set(updated.keys())
            for key in sorted(keys):
                if key in {'progress', 'last_change_summary'}:
                    continue
                next_path = f"{path}.{key}" if path else key
                walk(original.get(key), updated.get(key), next_path, depth + 1)
        elif original != updated:
            changes.append({
                "field": path or "root",
                "before": _stringify_change_value(original),
                "after": _stringify_change_value(updated)
            })

    walk(original_data, updated_data, "", 0)
    return {
        "change_id": str(int(time.time() * 1000)),
        "generated_at": datetime.utcnow().isoformat(),
        "items": changes,
        "count": len(changes)
    }

def _ensure_dict_content(value):
    """Safely normalizes CLP content payloads that may be JSON strings or dicts."""
    if isinstance(value, dict):
        return value
    if isinstance(value, str):
        try:
            parsed = json.loads(value)
            return parsed if isinstance(parsed, dict) else {}
        except Exception:
            return {}
    return {}

@timed("CLP Data Generation")
def generate_clp_data_task(app_context, plan_id, user_id, course_data):
    subject_name = course_data['subject']
    department_name = course_data['department']
    with app_context:
        try:
            StructuredLogger.info(f"--- [AI PHASE 1] Starting Task for CLP {plan_id} ---", plan_id=plan_id, user_id=user_id)
            url = current_app.config.get("SUPABASE_URL")
            key = current_app.config.get("SUPABASE_KEY")
            local_supabase = create_client(url, key)
            
            heartbeat_data = {'progress': {'step': 0, 'label': 'Connecting to AI Engine...', 'percent': 1}}
            update_progress(local_supabase, plan_id, 0, "Connecting to AI Engine...", 1, heartbeat_data)
            
            clp_data = {'progress': {'step': 0, 'label': 'Initializing AI Engine...', 'percent': 5}}
            update_progress(local_supabase, plan_id, 0, "Initializing AI Engine...", 5, clp_data)
            
            StructuredLogger.info(f"--- [AI PHASE 1] Fetching outcomes for department: {department_name} ---")
            
            program_outcomes_list = []
            course_outcomes_list = []
            program_headers = []
            institutional_headers = []

            dept_id = None
            try:
                outcomes_bundle = get_department_outcomes_bundle(department_name)
                dept_id = outcomes_bundle.get('department_id')
                program_outcomes_list = outcomes_bundle.get('program_outcomes', [])
                course_outcomes_list = outcomes_bundle.get('course_outcomes', [])
                program_headers = outcomes_bundle.get('program_headers', [])
                institutional_headers = outcomes_bundle.get('institutional_headers', [])

                if not dept_id:
                    raise ValueError(f"Department '{department_name}' not found.")
                if not program_outcomes_list:
                    raise ValueError(f"No Program Outcomes configured for department '{department_name}'.")
                if not course_outcomes_list:
                    raise ValueError(f"No Course Outcomes configured for department '{department_name}'.")
                if not institutional_headers:
                    raise ValueError("No Institutional Outcomes configured.")
            except Exception as e:
                raise ValueError(f"Outcome configuration error: {e}")

            model_instance = AIClient.get_model()
            clp_data = {'progress': {'step': 1, 'label': 'Designing Course Blueprint', 'percent': 10}}
            
            rag_context = ""
            try:
                from app.services.rag_service import RAGService
                if dept_id:
                    rag_results = RAGService.query_knowledge_base(f"Course policy and syllabus guidelines for {subject_name}", dept_id)
                    if rag_results:
                        rag_context = "\nINSTITUTIONAL GUIDELINES (STRICTLY FOLLOW):\n"
                        for doc in rag_results:
                            rag_context += f"- {doc['content']}\n"
            except Exception as rag_e:
                StructuredLogger.error(f"RAG Retrieval failed: {rag_e}")

            update_progress(local_supabase, plan_id, 1, "Analyzing Department Outcomes...", 20, clp_data)

            # Step 2: PO-IO
            update_progress(local_supabase, plan_id, 2, "Mapping Program to Institutional Outcomes...", 40, clp_data)
            po_io_schema = {f"{po}_{io}": {"type": "STRING", "enum": ["✔", " "]} for po in program_headers for io in institutional_headers}
            po_io_schema["po_io_rationale"] = {"type": "STRING"}
            
            step2_config = {
                "response_mime_type": "application/json", 
                "response_schema": {"type": "OBJECT", "properties": po_io_schema, "required": list(po_io_schema.keys())}
            }
            po_str = "\n".join([f"{po['code']}: {po['description']}" for po in program_outcomes_list])
            
            po_io_prompt = f"""
            {get_system_prompt(local_supabase, 'prompt_po_io', 'Analyze alignment between Program Outcomes (PO) and Institutional Outcomes (IO). Map ✔ for aligned.')}
            
            CONTEXT:
            {rag_context}
            
            PROGRAM OUTCOMES:
            {po_str}
            
            INSTITUTIONAL HEADERS:
            {', '.join(institutional_headers)}
            
            Return JSON with keys like 'PO_IO'. Use '✔' or ' '. ALSO include a 'po_io_rationale' explaining the alignment logic.
            """
            resp_2 = AIClient.generate_with_retry(model_instance, [po_io_prompt], step2_config, retries=5, 
                                                task_type="po_io_mapping", plan_id=plan_id, user_id=user_id)
            clp_data.update(json.loads(AIClient.clean_ai_json(resp_2.text)))

            # Step 3: CO-PO
            update_progress(local_supabase, plan_id, 3, "Aligning Course Outcomes to Program Goals...", 65, clp_data)
            co_codes = [c['code'] for c in course_outcomes_list]
            co_po_schema = {f"{co}_{po}": {"type": "STRING", "enum": ["I", "R", "M", " "]} for co in co_codes for po in program_headers}
            co_po_schema["co_po_rationale"] = {"type": "STRING"}
            
            step3_config = {
                "response_mime_type": "application/json", 
                "response_schema": {"type": "OBJECT", "properties": co_po_schema, "required": list(co_po_schema.keys())}
            }
            co_str = "\n".join([f"{co['code']}: {co['description']}" for co in course_outcomes_list])
            
            co_po_prompt = f"""
            {get_system_prompt(local_supabase, 'prompt_co_po', 'Map Course Outcomes (CO) to Program Outcomes (PO).')}
            
            CONTEXT:
            {rag_context}
            
            COURSE OUTCOMES:
            {co_str}
            
            PROGRAM OUTCOMES:
            {po_str}
            
            Return JSON with keys like 'CO_PO'. Use 'I' (Introduced), 'R' (Reinforced), 'M' (Mastered), or ' '. ALSO include a 'co_po_rationale' explaining the mapping.
            """
            resp_3 = AIClient.generate_with_retry(model_instance, [co_po_prompt], step3_config, retries=5,
                                                task_type="co_po_mapping", plan_id=plan_id, user_id=user_id)
            clp_data.update(json.loads(AIClient.clean_ai_json(resp_3.text)))

            # Step 4: Weekly
            update_progress(local_supabase, plan_id, 4, "Building Comprehensive Weekly Outline...", 85, clp_data)
            week_keys = {}
            for i in range(1, 19):
                prefix = f"W{i}"
                if i in [10, 11]: prefix = "W1011"
                elif i in [14, 15]: prefix = "W1415"
                elif i in [16, 17]: prefix = "W1617"
                for suffix in ["_LO", "_TO", "_Method", "_Assesment", "_LR"]: week_keys[f"{prefix}{suffix}"] = {"type": "STRING"}
                week_keys[f"{prefix}_Mapped_COs"] = {"type": "STRING"}
                week_keys[f"{prefix}_Assessment_COs"] = {"type": "STRING"}
            week_keys['references'] = {"type": "STRING"}
            step4_config = {
                "response_mime_type": "application/json", 
                "response_schema": {"type": "OBJECT", "properties": week_keys, "required": list(week_keys.keys())}
            }
            
            weekly_prompt = f"""
            {get_system_prompt(local_supabase, 'prompt_weekly', 'Generate comprehensive weekly outline with Learning Outcomes (LO), Topics (TO), Methods, and Assessments.')}
            
            CONTEXT:
            {rag_context}
            
            COURSE INFO:
            <subject>{subject_name}</subject>
            <description>{course_data.get('course_description', '')}</description>
            COURSE OUTCOME CODES:
            {', '.join(co_codes)}
            
            INSTRUCTIONS:
            1. Generate the weekly outline based ONLY on the COURSE INFO provided above.
            2. For every week, populate `<WEEK_PREFIX>_Mapped_COs` using comma-separated course outcome codes (example: "CO1, CO3").
            3. For every week, populate `<WEEK_PREFIX>_Assessment_COs` using comma-separated course outcome codes explicitly assessed that week.
            4. Use only valid codes from COURSE OUTCOME CODES.
            """
            resp_4 = AIClient.generate_with_retry(model_instance, [weekly_prompt], step4_config, retries=5,
                                                task_type="weekly_gen", plan_id=plan_id, user_id=user_id)
            clp_data.update(json.loads(AIClient.clean_ai_json(resp_4.text)))

            clp_data.update(course_data)
            
            user_res = local_supabase.table('users').select('first_name, last_name, title, consultation_hours').eq('id', user_id).single().execute()
            
            if user_res.data:
                clp_data['NAME'] = f"{user_res.data.get('first_name', '')} {user_res.data.get('last_name', '')}".strip()
                clp_data['TITLE'] = user_res.data.get('title', '')
                cons = user_res.data.get('consultation_hours') or {}
                clp_data['cons_mon_time'] = cons.get('monday', {}).get('time', 'N/A')
                clp_data['cons_mon_room'] = cons.get('monday', {}).get('room', 'N/A')
                clp_data['cons_tue_time'] = cons.get('tuesday', {}).get('time', 'N/A')
                clp_data['cons_tue_room'] = cons.get('tuesday', {}).get('room', 'N/A')
                clp_data['cons_wed_time'] = cons.get('wednesday', {}).get('time', 'N/A')
                clp_data['cons_wed_room'] = cons.get('wednesday', {}).get('room', 'N/A')
                clp_data['cons_thu_time'] = cons.get('thursday', {}).get('time', 'N/A')
                clp_data['cons_thu_room'] = cons.get('thursday', {}).get('room', 'N/A')
                clp_data['cons_fri_time'] = cons.get('friday', {}).get('time', 'N/A')
                clp_data['cons_fri_room'] = cons.get('friday', {}).get('room', 'N/A')

            local_supabase.table('course_learning_plans').update({
                'content': json.dumps(clp_data), 
                'status': 'draft_review',
                'upload_type': 'ai_generated'
            }).eq('id', plan_id).execute()
            
            shred_clp_mappings(local_supabase, plan_id, clp_data)
            
            # Save Initial Version
            try:
                from app.services.version_service import VersionService
                VersionService.save_version(plan_id, clp_data, "AI Generated Draft", user_id)
            except Exception as v_e:
                StructuredLogger.warning(f"Failed to save version for {plan_id}: {v_e}")

            create_notification(user_id, f'Draft for "{subject_name}" is ready for review.', reference_type='clp', reference_id=plan_id)

        except Exception as e:
            error_msg = handle_system_error(e, "CLP Generation", user_id)
            try: local_supabase.table('course_learning_plans').update({'status': 'failed', 'content': json.dumps({'error': error_msg})}).eq('id', plan_id).execute()
            except: pass

@timed("CLP Data Refinement")
def refine_clp_data_task(app_context, plan_id, user_id, original_data, updated_data):
    """Refines data by comparing original vs edited."""
    with app_context:
        try:
            url: str = os.environ.get("SUPABASE_URL")
            key: str = os.environ.get("SUPABASE_KEY")
            local_supabase = create_client(url, key)
            StructuredLogger.info(f"--- [AI PHASE 2] Refining Data for CLP {plan_id} ---", plan_id=plan_id)

            update_progress(local_supabase, plan_id, 1, "Analyzing User Edits...", 10, updated_data)
            
            model_instance = AIClient.get_model()
            
            update_progress(local_supabase, plan_id, 2, "Ensuring Academic Consistency...", 50, updated_data)
            
            consistency_prompt = f"""
            You are an academic expert specializing in Course Learning Plan (CLP) validation and consistency checking.
            Your task is to analyze the user's edited CLP and update it while preserving their edits exactly.

            ORIGINAL_DATA (before user edits):
            {json.dumps(original_data)}

            UPDATED_DATA (after user edits):
            {json.dumps(updated_data)}

            INSTRUCTIONS:
            1. Identify all fields the user changed. Compare ORIGINAL_DATA and UPDATED_DATA field-by-field.
            2. Automatically update any related or dependent fields to maintain academic consistency.
            3. Improve grammar, phrasing, and formatting where needed.
            4. Output MUST be a valid JSON object.
            """
            
            try:
                resp = AIClient.generate_with_retry(model_instance, [consistency_prompt], 
                                                  {"response_mime_type": "application/json"},
                                                  task_type="refinement", plan_id=plan_id, user_id=user_id)
                refined = json.loads(AIClient.clean_ai_json(resp.text))
                updated_data.update(refined)
            except Exception as ai_e:
                StructuredLogger.warning(f"AI Refinement failed, using user data: {ai_e}")

            try:
                updated_data['last_change_summary'] = summarize_clp_changes(original_data, updated_data)
            except Exception as summary_error:
                StructuredLogger.warning(f"Failed to summarize changes for {plan_id}: {summary_error}")

            # Check for Abort BEFORE final update
            task_check = local_supabase.table('background_tasks').select('status, error_message').eq('plan_id', plan_id).in_('status', ['failed']).order('id', desc=True).limit(1).execute()
            if task_check.data and 'Aborted' in (task_check.data[0].get('error_message') or ''):
                StructuredLogger.info(f"Refinement for {plan_id} skipped final update due to user abort.")
                return

            local_supabase.table('course_learning_plans').update({
                'content': json.dumps(updated_data), 
                'status': 'draft_review'
            }).eq('id', plan_id).execute()
            
            # Save Version
            from app.services.version_service import VersionService
            VersionService.save_version(plan_id, updated_data, "AI Refinement", user_id)
            
            create_notification(user_id, 'CLP draft updated.', reference_type='clp', reference_id=plan_id)

        except Exception as e:
            StructuredLogger.error(f"Refinement Critical Failure: {e}")
            try: local_supabase.table('course_learning_plans').update({'status': 'draft_review'}).eq('id', plan_id).execute()
            except: pass

@timed("CLP Finalization")
def finalize_clp_task(app_context, plan_id, user_id):
    with app_context:
        try:
            url: str = os.environ.get("SUPABASE_URL")
            key: str = os.environ.get("SUPABASE_KEY")
            local_supabase = create_client(url, key)
            StructuredLogger.info(f"--- [AI PHASE 3] Finalizing CLP {plan_id} ---", plan_id=plan_id)

            plan_res = local_supabase.table('course_learning_plans').select('*').eq('id', plan_id).single().execute()
            plan = plan_res.data
            clp_data = _ensure_dict_content(plan.get('content'))
            subject_name = plan['subject']
            department_name = plan['department']

            update_progress(local_supabase, plan_id, 5, "Finalizing Document Alignment...", 92, clp_data)

            model_instance = AIClient.get_model()
            consistency_prompt = f"""
            You are an academic CLP (Course Learning Plan) validator.
            Refine the provided CLP JSON while strictly preserving all user-authored content and structure.
            INPUT_JSON: {json.dumps(clp_data)}
            Return ONLY the final JSON object.
            """

            try:
                resp = AIClient.generate_with_retry(model_instance, [consistency_prompt], 
                                                  {"response_mime_type": "application/json"},
                                                  task_type="finalization", plan_id=plan_id, user_id=user_id)
                clp_data.update(json.loads(AIClient.clean_ai_json(resp.text)))
            except Exception as e:
                StructuredLogger.warning(f"AI Finalization refinement failed: {e}")

            dept_id = get_department_id_by_name(department_name, local_supabase=local_supabase)
            
            template_key = None
            if dept_id:
                tmpl_res = local_supabase.table('templates').select('filename').eq('department_id', dept_id).limit(1).execute()
                if tmpl_res.data: template_key = tmpl_res.data[0]['filename']
            if not template_key:
                def_res = local_supabase.table('templates').select('filename').eq('is_default', True).limit(1).execute()
                if def_res.data: template_key = def_res.data[0]['filename']

            if template_key: template_bytes = local_supabase.storage.from_(STORAGE_BUCKET_NAME).download(template_key)
            else:
                with open(get_template_filepath(), 'rb') as f: template_bytes = f.read()

            doc = Document(BytesIO(template_bytes)) 
            flat = flatten_json(clp_data)
            flat['date_today'] = datetime.now().strftime("%B %Y")
            
            doc = replace_placeholders(doc, flat)
            output = BytesIO()
            doc.save(output)
            output.seek(0)
            
            docx_filename = f"{subject_name.replace(' ', '_')}_{int(time.time())}.docx"
            file_path = f"{user_id}/{docx_filename}"
            local_supabase.storage.from_(STORAGE_BUCKET_NAME).upload(path=file_path, file=output.read(), file_options={"content-type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"})
            
            local_supabase.table('course_learning_plans').update({
                'filename': file_path, 'upload_type': 'file_upload', 'content': json.dumps(clp_data), 'status': 'draft'
            }).eq('id', plan_id).execute()
            
            # Save Final Version
            from app.services.version_service import VersionService
            VersionService.save_version(plan_id, clp_data, "Finalized Document", user_id)
            
            create_notification(user_id, f'CLP "{subject_name}" has been finalized.', reference_type='clp', reference_id=plan_id)

        except Exception as e:
            StructuredLogger.error(f"Finalization FAILED: {traceback.format_exc()}")
            try: local_supabase.table('notifications').insert({'user_id': user_id, 'message': f'Finalization Failed: {str(e)}'}).execute()
            except: pass

@timed("Magic Syllabus Extraction")
def magic_syllabus_extraction_task(app_context, plan_id, user_id, temp_path, department_name):
    with app_context:
        try:
            url: str = os.environ.get("SUPABASE_URL")
            key: str = os.environ.get("SUPABASE_KEY")
            local_supabase = create_client(url, key)
            
            doc = Document(temp_path)
            full_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            try: os.unlink(temp_path) 
            except: pass
            
            update_progress(local_supabase, plan_id, 1, "Reading Syllabus Content...", 15, {})
            
            model_instance = AIClient.get_model()
            
            update_progress(local_supabase, plan_id, 2, "Extracting Curriculum Structure...", 45, {})
            
            magic_prompt = f"""
            Extract content from SYLLABUS TEXT into structured CLP JSON.
            SYLLABUS TEXT: {full_text}
            DEPARTMENT: {department_name}
            Return ONLY the JSON object.
            """
            
            resp = AIClient.generate_with_retry(model_instance, [magic_prompt], 
                                              {"response_mime_type": "application/json"},
                                              task_type="magic_sync", plan_id=plan_id, user_id=user_id)
            clp_data = json.loads(AIClient.clean_ai_json(resp.text))
            clp_data['upload_type'] = 'ai_generated'
            clp_data['department'] = department_name

            local_supabase.table('course_learning_plans').update({
                'content': json.dumps(clp_data), 
                'subject': clp_data.get('subject', f"Magic Sync: {department_name}"),
                'status': 'draft_review'
            }).eq('id', plan_id).execute()
            
            # Save Initial Magic Version
            from app.services.version_service import VersionService
            VersionService.save_version(plan_id, clp_data, "Magic Syllabus Import", user_id)
            
            shred_clp_mappings(local_supabase, plan_id, clp_data)
            create_notification(user_id, f"Magic Sync complete for {clp_data.get('subject')}!", reference_type='clp', reference_id=plan_id)

        except Exception as e:
            StructuredLogger.error(f"Magic Extraction FAILED: {e}")
            try: local_supabase.table('course_learning_plans').update({'status': 'failed'}).eq('id', plan_id).execute()
            except: pass

@timed("Validation Fix Application")
def apply_validation_fixes_task(app_context, plan_id, user_id, current_content, validation_data):
    """
    Refines the CLP content based on validation feedback (Audit results).
    """
    with app_context:
        try:
            url: str = os.environ.get("SUPABASE_URL")
            key: str = os.environ.get("SUPABASE_KEY")
            local_supabase = create_client(url, key)
            StructuredLogger.info(f"--- [AI REFINEMENT] Applying Validation Fixes for CLP {plan_id} ---", plan_id=plan_id)

            update_progress(local_supabase, plan_id, 1, "Analyzing Pedagogical Issues...", 10, current_content)

            model_instance = AIClient.get_model()
            
            update_progress(local_supabase, plan_id, 2, "Refactoring Curriculum Content...", 50, current_content)

            refine_prompt = f"""
            You are an educational curriculum expert. Your task is to refactor the following Course Learning Plan (CLP) JSON based on the provided VALIDATION AUDIT.

            CURRENT_CLP_JSON:
            {json.dumps(current_content)}

            VALIDATION_AUDIT:
            {json.dumps(validation_data)}

            INSTRUCTIONS:
            1. Address every 'issue' mentioned in the VALIDATION_AUDIT by applying the 'suggestion' to the relevant fields in the JSON.
            2. Ensure all Learning Outcomes (LO), Topics (TO), and Assessments are consistent and meet the required Bloom's Taxonomy levels.
            3. Preserve all other fields and the overall structure.
            4. Return ONLY the updated JSON object.
            """

            original_snapshot = json.loads(json.dumps(current_content))

            try:
                resp = AIClient.generate_with_retry(model_instance, [refine_prompt],
                                                  {"response_mime_type": "application/json"},
                                                  task_type="validation_fix", plan_id=plan_id, user_id=user_id)

                update_progress(local_supabase, plan_id, 3, "Applying Structural Improvements...", 85, current_content)

                refined_content = json.loads(AIClient.clean_ai_json(resp.text))
                # Merge refined content back (keeping system fields etc)
                current_content.update(refined_content)

                try:
                    current_content['last_change_summary'] = summarize_clp_changes(original_snapshot, current_content)
                except Exception as summary_error:
                    StructuredLogger.warning(f"Failed to summarize validation fixes for {plan_id}: {summary_error}")
                
                local_supabase.table('course_learning_plans').update({
                    'content': json.dumps(current_content), 
                    'status': 'draft_review'
                }).eq('id', plan_id).execute()
                
                # Save Version
                from app.services.version_service import VersionService
                VersionService.save_version(plan_id, current_content, "AI Validation Improvements", user_id)
                
                shred_clp_mappings(local_supabase, plan_id, current_content)
                create_notification(user_id, f"AI has successfully applied improvements to your CLP based on the audit.", reference_type='clp', reference_id=plan_id)
                
            except Exception as ai_e:
                StructuredLogger.error(f"AI Fix Application failed: {ai_e}")
                local_supabase.table('course_learning_plans').update({'status': 'draft_review'}).eq('id', plan_id).execute()

        except Exception as e:
            StructuredLogger.error(f"Apply Validation Fixes Critical Failure: {e}")
            try: local_supabase.table('course_learning_plans').update({'status': 'draft_review'}).eq('id', plan_id).execute()
            except: pass
